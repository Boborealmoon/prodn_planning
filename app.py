from __future__ import annotations

import json
import os
import re
import sqlite3
import threading
import time
from datetime import date, datetime, timedelta
from pathlib import Path
from uuid import uuid4

from dotenv import load_dotenv
from flask import Flask, jsonify, redirect, render_template, request

load_dotenv()

try:
    from openpyxl import load_workbook
except ImportError:  # pragma: no cover - optional until Excel sync is used
    load_workbook = None

try:
    import psycopg2
    import psycopg2.extras
except ImportError:  # pragma: no cover - optional dependency until ERP sync is enabled
    psycopg2 = None


ROOT = Path(__file__).resolve().parent
DB_PATH = Path(os.environ.get("DB_PATH", str(ROOT / "planner.db")))
ERP_PG_HOST = os.environ.get("ERP_PG_HOST", "localhost")
ERP_PG_PORT = int(os.environ.get("ERP_PG_PORT", "5432"))
ERP_PG_DBNAME = os.environ.get("ERP_PG_DBNAME", "")
ERP_PG_USER = os.environ.get("ERP_PG_USER", "")
ERP_PG_PASSWORD = os.environ.get("ERP_PG_PASSWORD", "")
STANDARD_START = 510
STANDARD_END = 1200
STANDARD_WINDOWS = [(510, 720), (765, 960), (975, 1200)]
STANDARD_BREAKS = [(720, 765), (960, 975)]
MATERIAL_LEAD_DAYS = 2
PROGRAM_LEAD_DAYS = 1
TOOLLIST_LEAD_DAYS = 1

app = Flask(__name__)

DEV_RELOAD_EXIT_CODE = 3
DEV_RELOAD_POLL_SECONDS = 1.0
DEV_RELOAD_EXTENSIONS = {".py", ".html", ".htm", ".js", ".css", ".sql", ".bat"}


def _dev_reload_paths():
    paths = []
    for root in (ROOT, ROOT / "templates", ROOT / "static", ROOT / "sql"):
        if not root.exists():
            continue
        for path in root.rglob("*"):
            if path.is_file() and path.suffix.lower() in DEV_RELOAD_EXTENSIONS:
                paths.append(path)
    return paths


def _dev_reload_snapshot():
    snapshot = {}
    for path in _dev_reload_paths():
        try:
            stat = path.stat()
        except OSError:
            continue
        snapshot[str(path)] = (stat.st_mtime_ns, stat.st_size)
    return snapshot


def _start_dev_reload_watcher():
    baseline = _dev_reload_snapshot()

    def watch():
        while True:
            time.sleep(DEV_RELOAD_POLL_SECONDS)
            current = _dev_reload_snapshot()
            if current != baseline:
                print("Code change detected, restarting...")
                os._exit(DEV_RELOAD_EXIT_CODE)

    thread = threading.Thread(target=watch, name="dev-reload-watcher", daemon=True)
    thread.start()
    return thread

ERP_REQUIRED_SHEETS = {
    "ProcessSheets": "erp_process_sheets_staging",
    "Materials (Per PS)": "erp_materials_per_ps_staging",
    "Material (Per BOM)": "erp_materials_per_bom_staging",
    "BOM_op_stage": "erp_bom_op_stage_staging",
    "Workorder Tracker": "erp_workorder_tracker_staging",
}

ERP_OPTIONAL_SHEET_ALIASES = {
    "Active Orders": {
        "activeorders",
        "activeorder",
        "openorders",
        "openorder",
        "activeps",
        "activeprocesssheets",
    }
}


def db():
    con = sqlite3.connect(DB_PATH)
    con.row_factory = sqlite3.Row
    con.execute("PRAGMA foreign_keys = ON")
    return con


def rows(cur):
    return [dict(row) for row in cur.fetchall()]


def one(cur):
    row = cur.fetchone()
    return dict(row) if row else None


def api_error(message, status=400, details=None):
    payload = {"error": message}
    if details is not None:
        payload["details"] = details
    return jsonify(payload), status


def erp_pg_connect(host=None, port=None, dbname=None, user=None, password=None):
    if psycopg2 is None:
        raise RuntimeError("psycopg2 is not installed. Run: pip install psycopg2-binary")
    return psycopg2.connect(
        host=host or ERP_PG_HOST,
        port=port or ERP_PG_PORT,
        dbname=dbname or ERP_PG_DBNAME,
        user=user or ERP_PG_USER,
        password=password or ERP_PG_PASSWORD,
        connect_timeout=5,
    )


def erp_test_connection(host=None, port=None, dbname=None, user=None, password=None):
    with erp_pg_connect(host=host, port=port, dbname=dbname, user=user, password=password) as con:
        with con.cursor() as cur:
            cur.execute("SELECT current_database(), current_user")
            db_name, current_user = cur.fetchone()
        return {
            "host": host or ERP_PG_HOST,
            "database": db_name,
            "current_user": current_user,
        }


def normalize_column_name(value):
    value = str(value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "_", value)
    return value.strip("_")


def normalize_sheet_name(value):
    value = str(value or "").strip().lower()
    value = re.sub(r"[^a-z0-9]+", "", value)
    return value


def compact_text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value != value:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def process_sheet_key(ps_id, pp_partial_no):
    base = compact_text(ps_id)
    partial = compact_text(pp_partial_no) or "1"
    return f"{base}::{partial}" if base else ""


def split_process_sheet_key(ps_key):
    text = compact_text(ps_key)
    if "::" not in text:
        return text, "1"
    base, partial = text.rsplit("::", 1)
    return base, partial or "1"


def process_sheet_source_id(ps):
    if not ps:
        return ""
    return compact_text(ps.get("source_ps_id")) or split_process_sheet_key(ps.get("ps_id"))[0]


def process_sheet_partial_no(ps):
    if not ps:
        return "1"
    return compact_text(ps.get("pp_partial_no")) or split_process_sheet_key(ps.get("ps_id"))[1]


def first_nonempty(row, *keys):
    for key in keys:
        value = compact_text(row.get(key))
        if value:
            return value
    return ""


def workbook_partial_qty(row, total_key="total_qty", partial_key="partial_qty"):
    """
    Prefer the workbook's partial quantity for split orders.
    Fall back to the total quantity only when the partial quantity is missing or zero.
    """
    partial_qty = parse_number(row.get(partial_key))
    if partial_qty > 0:
        return partial_qty
    return parse_number(row.get(total_key))


def parse_date_text(value):
    if value in (None, ""):
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if text.lower() in {"nan", "nat", "none", "null"}:
        return ""
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y", "%d-%m-%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue
    return text


def parse_number(value, default=0.0):
    if value in (None, ""):
        return default
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    try:
        return float(text)
    except ValueError:
        return default


def parse_bool_flag(value):
    if isinstance(value, bool):
        return value
    text = compact_text(value).lower()
    return text in {"1", "true", "yes", "y", "ready", "done", "complete", "completed"}


def map_ps_status(value):
    text = compact_text(value).upper()
    if text in {"", "OPEN", "RELEASED", "IN_PROGRESS", "WIP", "PLANNING", "PARTIAL", "OUTSTANDING", "OUTSTANDING PARTIAL"}:
        return "ACTIVE"
    if text == "HISTORY":
        return "COMPLETED"
    if "HOLD" in text:
        return "ON_HOLD"
    if text in {"COMPLETE", "COMPLETED", "CLOSED", "DONE"}:
        return "COMPLETED"
    if text in {"CANCELLED", "CANCELED"}:
        return "ON_HOLD"
    return text or "ACTIVE"


def excel_sheet_records(ws):
    rows_iter = ws.iter_rows(values_only=True)
    header_row = next(rows_iter, None)
    if not header_row:
        return []
    headers = [normalize_column_name(col) for col in header_row]
    records = []
    for row in rows_iter:
        if row is None or not any(cell not in (None, "") for cell in row):
            continue
        record = {}
        for idx, header in enumerate(headers):
            if not header:
                continue
            record[header] = row[idx] if idx < len(row) else None
        records.append(record)
    return records


def load_erp_workbook(file_storage):
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed. Run pip install -r requirements.txt first.")
    workbook = load_workbook(file_storage, data_only=True, read_only=True)
    sheet_map = {normalize_sheet_name(name): name for name in workbook.sheetnames}
    missing = [sheet for sheet in ERP_REQUIRED_SHEETS if normalize_sheet_name(sheet) not in sheet_map]
    if missing:
        raise ValueError(f"Workbook is missing required sheets: {', '.join(missing)}")
    workbook_data = {
        sheet: excel_sheet_records(workbook[sheet_map[normalize_sheet_name(sheet)]])
        for sheet in ERP_REQUIRED_SHEETS
    }
    for canonical, aliases in ERP_OPTIONAL_SHEET_ALIASES.items():
        matched_sheet = next((sheet_map[key] for key in aliases if key in sheet_map), None)
        workbook_data[canonical] = excel_sheet_records(workbook[matched_sheet]) if matched_sheet else []
    return workbook_data


def load_active_orders_workbook(file_storage):
    if load_workbook is None:
        raise RuntimeError("openpyxl is not installed. Run pip install -r requirements.txt first.")
    workbook = load_workbook(file_storage, data_only=True, read_only=True)
    sheet_map = {normalize_sheet_name(name): name for name in workbook.sheetnames}
    aliases = ERP_OPTIONAL_SHEET_ALIASES["Active Orders"]
    matched_sheet = next((sheet_map[key] for key in aliases if key in sheet_map), None)
    if not matched_sheet:
        raise ValueError("Active orders workbook is missing a sheet named Active Orders / ActiveOrders.")
    return excel_sheet_records(workbook[matched_sheet])


def clear_erp_staging(con):
    for table in (
        "erp_process_sheets_staging",
        "erp_materials_per_ps_staging",
        "erp_materials_per_bom_staging",
        "erp_bom_op_stage_staging",
        "erp_workorder_tracker_staging",
        "erp_active_orders_staging",
    ):
        con.execute(f"DELETE FROM {table}")


def insert_erp_staging_rows(con, sync_batch_id, workbook_data):
    for row in workbook_data["ProcessSheets"]:
        con.execute(
            """
            INSERT INTO erp_process_sheets_staging
            (sync_batch_id, ps_id, pp_partial_no, part_no, description, total_qty, partial_qty, due_date, order_date, bom_code, status, raw_payload)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                compact_text(row.get("ps_id")),
                compact_text(row.get("pp_partial_no")),
                compact_text(row.get("part_no")),
                compact_text(row.get("description")),
                parse_number(row.get("total_qty")),
                parse_number(row.get("partial_qty")),
                parse_date_text(row.get("due_date")),
                parse_date_text(row.get("order_date")),
                compact_text(row.get("bom_code")),
                compact_text(row.get("status")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )
    for row in workbook_data["Materials (Per PS)"]:
        con.execute(
            """
            INSERT INTO erp_materials_per_ps_staging
            (sync_batch_id, pp_voucher, pp_partial_no, inventory_code, material_inventory_code, raw_payload)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                compact_text(row.get("pp_voucher")),
                compact_text(row.get("pp_partial_no")),
                compact_text(row.get("inventory_code")),
                compact_text(row.get("material_inventory_code")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )
    for row in workbook_data["Material (Per BOM)"]:
        con.execute(
            """
            INSERT INTO erp_materials_per_bom_staging
            (sync_batch_id, bom_code, source_inventory_code, material_inventory_code, description, raw_payload)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                compact_text(row.get("bom_code")),
                compact_text(row.get("source_inventory_code")),
                compact_text(row.get("material_inventory_code")),
                compact_text(row.get("description")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )
    for row in workbook_data["BOM_op_stage"]:
        inventory_code = first_nonempty(row, "inventory_code", "part_no", "item_code", "source_inventory_code")
        preferred_machine = first_nonempty(row, "machine_no", "preferred_machine", "machine_code", "machine")
        con.execute(
            """
            INSERT INTO erp_bom_op_stage_staging
            (sync_batch_id, bom_code, inventory_code, seq_index, stage_no, stage_desc, machine_no, machine_category, cycle_time, setup_time, raw_payload)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                compact_text(row.get("bom_code")),
                inventory_code,
                int(parse_number(row.get("index"), 0)),
                first_nonempty(row, "op_no", "stage_no"),
                compact_text(row.get("stage_desc")),
                preferred_machine,
                "",
                parse_number(row.get("cycle_time")),
                parse_number(row.get("setup_time")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )
    for row in workbook_data["Workorder Tracker"]:
        con.execute(
            """
            INSERT INTO erp_workorder_tracker_staging
            (sync_batch_id, voucher_no, source_pp_no, inventory_code, machine_no, partial_seq_no, stage_no, stage_desc,
             acc_completion_qty, rej_completion_qty, total_acc_qty_produced, total_rej_qty_produced, employee_name, bom_code, raw_payload)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                compact_text(row.get("voucher_no")),
                compact_text(row.get("source_pp_no")),
                compact_text(row.get("inventory_code")),
                compact_text(row.get("machine_no")),
                compact_text(row.get("partial_seq_no")),
                compact_text(row.get("stage_no")),
                compact_text(row.get("stage_desc")),
                parse_number(row.get("acc_completion_qty")),
                parse_number(row.get("rej_completion_qty")),
                parse_number(row.get("total_acc_qty_produced")),
                parse_number(row.get("total_rej_qty_produced")),
                compact_text(row.get("employee_name")),
                compact_text(row.get("bom_code")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )
    for row in workbook_data.get("Active Orders", []):
        con.execute(
            """
            INSERT INTO erp_active_orders_staging
            (sync_batch_id, ps_id, part_no, bom_code, raw_payload)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                first_nonempty(row, "ps_id", "source_pp_no", "pp_no", "pp_partial_no", "voucher_no"),
                first_nonempty(row, "part_no", "inventory_code", "item_code", "source_inventory_code"),
                compact_text(row.get("bom_code")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )


def insert_active_orders_staging_rows(con, sync_batch_id, rows_data):
    con.execute("DELETE FROM erp_active_orders_staging")
    for row in rows_data:
        con.execute(
            """
            INSERT INTO erp_active_orders_staging
            (sync_batch_id, ps_id, part_no, bom_code, raw_payload)
            VALUES (?, ?, ?, ?, ?)
            """,
            (
                sync_batch_id,
                first_nonempty(row, "ps_id", "source_pp_no", "pp_no", "pp_partial_no", "voucher_no"),
                first_nonempty(row, "part_no", "inventory_code", "item_code", "source_inventory_code"),
                compact_text(row.get("bom_code")),
                json.dumps({k: compact_text(v) for k, v in row.items()}),
            ),
        )


def apply_active_orders_statuses(con, sync_batch_id):
    updated = skipped = 0
    active_rows = rows(con.execute("SELECT ps_id, part_no, bom_code FROM erp_active_orders_staging WHERE sync_batch_id = ?", (sync_batch_id,)))
    active_ps_ids = {compact_text(row["ps_id"]) for row in active_rows if compact_text(row["ps_id"])}
    existing_rows = rows(
        con.execute(
            """
            SELECT ps.ps_id, ps.source_ps_id, ps.inv_code AS part_no, l.bom_code, ps.status
            FROM process_sheet ps
            LEFT JOIN process_sheet_erp_link l ON l.ps_id = ps.ps_id
            """
        )
    )
    for ps in existing_rows:
        source_ps_id = compact_text(ps.get("source_ps_id")) or split_process_sheet_key(ps["ps_id"])[0]
        new_status = "ACTIVE" if source_ps_id in active_ps_ids else "COMPLETED"
        if compact_text(ps["status"]) != new_status:
            con.execute(
                "UPDATE process_sheet SET status = ?, updated_at = CURRENT_TIMESTAMP WHERE ps_id = ?",
                (new_status, ps["ps_id"]),
            )
            updated += 1
        else:
            skipped += 1
    return updated, skipped


def find_or_create_part(con, part_no, description=""):
    part = one(con.execute("SELECT * FROM parts WHERE part_name = ?", (part_no,)))
    if part:
        if description and description != (part["part_desc"] or ""):
            con.execute("UPDATE parts SET part_desc = ?, updated_at = CURRENT_TIMESTAMP WHERE part_id = ?", (description, part["part_id"]))
        return part["part_id"]
    cur = con.execute("INSERT INTO parts (part_name, part_desc) VALUES (?, ?)", (part_no, description))
    return cur.lastrowid


def machine_category_for(con, machine_code):
    machine_code = compact_text(machine_code)
    if not machine_code:
        return "UNKNOWN"
    row = one(con.execute("SELECT machine_category FROM machines WHERE machine_code = ?", (machine_code,)))
    return compact_text(row["machine_category"]).upper() if row and row.get("machine_category") else "UNKNOWN"


def ensure_part_flow_step_columns(con):
    cols = [row[1] for row in con.execute("PRAGMA table_info(part_flow_steps)").fetchall()]
    if "share_setup_with_prev" in cols:
        try:
            con.execute("ALTER TABLE part_flow_steps DROP COLUMN share_setup_with_prev")
        except sqlite3.OperationalError:
            pass


def infer_machine_categories_from_step(con, step):
    categories = []

    preferred_machine = compact_text(step.get("preferred_machine"))
    if preferred_machine:
        preferred_category = machine_category_for(con, preferred_machine)
        if preferred_category != "UNKNOWN":
            categories.append(preferred_category)

    existing_category = compact_text(step.get("machine_category")).upper()
    if existing_category and existing_category != "UNKNOWN" and existing_category not in categories:
        categories.append(existing_category)

    op_text = " ".join(
        [
            compact_text(step.get("op_type")).upper(),
            compact_text(step.get("op_no")).upper(),
        ]
    )
    keyword_map = [
        ("MILL", ["MILLING", "MILL"]),
        ("TURN", ["TURNING", "LATHE"]),
        ("LATHE", ["LATHE", "TURNING"]),
        ("GRIND", ["GRINDER", "GRINDING"]),
        ("DRILL", ["DRILLING", "DRILL"]),
        ("FURNACE", ["FURNACE", "HEAT TREAT"]),
        ("HEAT", ["HEAT TREAT", "FURNACE"]),
        ("POLISH", ["POLISH"]),
    ]
    for keyword, candidates in keyword_map:
        if keyword in op_text:
            for category in candidates:
                if category not in categories:
                    categories.append(category)

    if "UNKNOWN" not in categories:
        categories.append("UNKNOWN")
    return categories


def pick_machine_for_step(con, step, require_preferred_machine=False):
    preferred_machine = compact_text(step.get("preferred_machine"))
    if require_preferred_machine and not preferred_machine:
        return None
    for category in infer_machine_categories_from_step(con, step):
        machine = one(
            con.execute(
                """
                SELECT * FROM machines WHERE active = 1 AND machine_category = ?
                ORDER BY CASE WHEN machine_code = ? THEN 0 ELSE 1 END, machine_id LIMIT 1
                """,
                (category, preferred_machine),
            )
        )
        if machine:
            if compact_text(step.get("machine_category")).upper() != category and category != "UNKNOWN":
                con.execute(
                    "UPDATE part_flow_steps SET machine_category = ? WHERE step_id = ?",
                    (category, step["step_id"]),
                )
                step["machine_category"] = category
            return machine
    return None


def refresh_flow_step_machine_categories(con, machine_code=None):
    if machine_code:
        machine_code = compact_text(machine_code)
        new_category = machine_category_for(con, machine_code)
        con.execute(
            """
            UPDATE part_flow_steps
            SET machine_category = ?
            WHERE COALESCE(preferred_machine, '') = ?
            """,
            (new_category, machine_code),
        )
        return

    steps = rows(con.execute("SELECT step_id, preferred_machine FROM part_flow_steps WHERE COALESCE(preferred_machine, '') <> ''"))
    for step in steps:
        con.execute(
            "UPDATE part_flow_steps SET machine_category = ? WHERE step_id = ?",
            (machine_category_for(con, step["preferred_machine"]), step["step_id"]),
        )


def bom_material_map(con):
    data = {}
    for row in con.execute("SELECT bom_code, source_inventory_code, material_inventory_code, description FROM erp_materials_per_bom_staging"):
        key = (row["bom_code"] or "", row["source_inventory_code"] or "", row["material_inventory_code"] or "")
        data[key] = row["description"] or row["material_inventory_code"] or ""
    return data


def ps_lookup_from_staging(con):
    lookup = {}
    rows_data = rows(con.execute("SELECT ps_id, pp_partial_no, part_no, bom_code FROM erp_process_sheets_staging"))
    partial_counts = {}
    for row in rows_data:
        base = compact_text(row["ps_id"])
        if base:
            partial_counts[base] = partial_counts.get(base, 0) + 1
    for row in rows_data:
        base = compact_text(row["ps_id"])
        partial = compact_text(row["pp_partial_no"]) or "1"
        key = process_sheet_key(base, partial)
        if not base:
            continue
        lookup[("ps_key", key)] = key
        lookup[("ps_pair", base, partial)] = key
        lookup.setdefault(("ps_id", base), key)
        lookup.setdefault(("pp_partial_no", partial), key)
        if row["part_no"]:
            lookup.setdefault(("part_no", row["part_no"]), key)
        if row["bom_code"]:
            lookup.setdefault(("bom_code", row["bom_code"]), key)
        if row["part_no"] and row["bom_code"]:
            lookup[("triple", base, row["part_no"], row["bom_code"])] = key
            lookup[("partial_triple", partial, row["part_no"], row["bom_code"])] = key
    return lookup


def upsert_bom_flow(con, part_id, part_no, bom_code):
    flow_code = bom_code
    flow = one(con.execute("SELECT * FROM part_flow_header WHERE part_id = ? AND flow_code = ?", (part_id, flow_code)))
    if flow:
        flow_id = flow["flow_id"]
        if flow_erp_sync_locked(con, flow_id):
            return flow_id, 0
    else:
        cur = con.execute(
            "INSERT INTO part_flow_header (part_id, flow_code, flow_name, is_default) VALUES (?, ?, ?, 1)",
            (part_id, flow_code, ""),
        )
        flow_id = cur.lastrowid
    con.execute("UPDATE part_flow_header SET flow_name = '' WHERE flow_id = ?", (flow_id,))
    con.execute("UPDATE part_flow_header SET is_default = CASE WHEN flow_id = ? THEN 1 ELSE 0 END WHERE part_id = ?", (flow_id, part_id))

    stage_rows = rows(
        con.execute(
            """
            SELECT * FROM erp_bom_op_stage_staging
            WHERE bom_code = ? AND (inventory_code = ? OR COALESCE(inventory_code, '') = '')
            ORDER BY COALESCE(seq_index, 0), CAST(COALESCE(NULLIF(stage_no, ''), '999999') AS INTEGER), stage_no, stage_id
            """,
            (bom_code, part_no),
        )
    )
    existing_steps = rows(
        con.execute(
            """
            SELECT *
            FROM part_flow_steps
            WHERE flow_id = ?
            ORDER BY seq_no, step_id
            """,
            (flow_id,),
        )
    )
    existing_by_step_id = {step["step_id"]: step for step in existing_steps}
    for idx, stage in enumerate(stage_rows, 1):
        preferred_machine = compact_text(stage.get("machine_no"))
        op_no = compact_text(stage.get("stage_no")) or f"OP{idx*10}"
        op_type = compact_text(stage.get("stage_desc")) or op_no
        machine_category = machine_category_for(con, preferred_machine)
        seq_no = (int(stage.get("seq_index") or 0) or idx) * 10
        cycle_time = parse_number(stage.get("cycle_time"), 1)
        setup_time = parse_number(stage.get("setup_time"), 0)
        is_last_op = 1 if idx == len(stage_rows) else 0
        if idx <= len(existing_steps):
            step_id = existing_steps[idx - 1]["step_id"]
            con.execute(
                """
                UPDATE part_flow_steps
                SET seq_no = ?, op_no = ?, op_type = ?, machine_category = ?, preferred_machine = ?,
                    cycle_time = ?, setup_time = ?, is_last_op = ?
                WHERE step_id = ? AND flow_id = ?
                """,
                (
                    seq_no,
                    op_no,
                    op_type,
                    machine_category,
                    preferred_machine,
                    cycle_time,
                    setup_time,
                    is_last_op,
                    step_id,
                    flow_id,
                ),
            )
        else:
            con.execute(
                """
                INSERT INTO part_flow_steps (flow_id, seq_no, op_no, op_type, machine_category, preferred_machine, cycle_time, setup_time, is_last_op)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    flow_id,
                    seq_no,
                    op_no,
                    op_type,
                    machine_category,
                    preferred_machine,
                    cycle_time,
                    setup_time,
                    is_last_op,
                ),
            )

    if len(existing_steps) > len(stage_rows):
        surplus_steps = existing_steps[len(stage_rows):]
        surplus_ids = [step["step_id"] for step in surplus_steps]
        referenced_ids = {
            row["flow_step_id"]
            for row in con.execute(
                f"SELECT DISTINCT flow_step_id FROM planning_block WHERE flow_step_id IN ({','.join('?' for _ in surplus_ids)})",
                surplus_ids,
            ).fetchall()
        } if surplus_ids else set()
        removable_ids = [step_id for step_id in surplus_ids if step_id not in referenced_ids]
        if removable_ids:
            placeholders = ",".join("?" for _ in removable_ids)
            con.execute(
                f"DELETE FROM part_flow_steps WHERE step_id IN ({placeholders})",
                removable_ids,
            )
        if referenced_ids:
            # Preserve any still-referenced legacy steps so existing planning rows keep valid foreign keys.
            con.execute(
                "UPDATE part_flow_steps SET is_last_op = 0 WHERE step_id IN ({})".format(",".join("?" for _ in referenced_ids)),
                list(referenced_ids),
            )
    return flow_id, len(stage_rows)


def merge_erp_process_sheets(con, sync_batch_id):
    inserted = updated = skipped = 0
    active_rows = rows(con.execute("SELECT ps_id, part_no, bom_code FROM erp_active_orders_staging WHERE sync_batch_id = ?", (sync_batch_id,)))
    active_ps_ids = {compact_text(row["ps_id"]) for row in active_rows if compact_text(row["ps_id"])}
    staged_rows = rows(
        con.execute(
            """
            SELECT * FROM erp_process_sheets_staging
            WHERE sync_batch_id = ?
            ORDER BY due_date, ps_id
            """,
            (sync_batch_id,),
        )
    )
    for row in staged_rows:
        source_ps_id = compact_text(row.get("ps_id"))
        partial_no = compact_text(row.get("pp_partial_no")) or "1"
        ps_id = process_sheet_key(source_ps_id, partial_no)
        part_no = compact_text(row.get("part_no"))
        if not source_ps_id or not part_no:
            skipped += 1
            continue
        description = compact_text(row.get("description"))
        part_id = find_or_create_part(con, part_no, description)
        existing = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if active_ps_ids:
            status = "ACTIVE" if source_ps_id in active_ps_ids else "COMPLETED"
        else:
            status = map_ps_status(row.get("status"))
        qty = workbook_partial_qty(row)
        if existing and ps_erp_sync_locked(con, ps_id):
            skipped += 1
        elif existing:
            con.execute(
                """
                UPDATE process_sheet
                SET part_id = ?, source_ps_id = ?, pp_partial_no = ?, inv_code = ?, inv_desc = ?, order_date = ?, due_date = ?,
                    total_qty = ?, status = ?, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ?
                """,
                (
                    part_id,
                    source_ps_id,
                    partial_no,
                    part_no,
                    description,
                    parse_date_text(row.get("order_date")),
                    parse_date_text(row.get("due_date")),
                    qty,
                    status,
                    ps_id,
                ),
            )
            updated += 1
        else:
            con.execute(
                """
                INSERT INTO process_sheet (ps_id, source_ps_id, pp_partial_no, part_id, inv_code, inv_desc, order_date, due_date, total_qty, status)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    ps_id,
                    source_ps_id,
                    partial_no,
                    part_id,
                    part_no,
                    description,
                    parse_date_text(row.get("order_date")),
                    parse_date_text(row.get("due_date")),
                    qty,
                    status,
                ),
            )
            inserted += 1

        con.execute(
            """
            INSERT INTO process_sheet_erp_link (ps_id, pp_partial_no, bom_code, erp_status, last_sync_at)
            VALUES (?, ?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id) DO UPDATE SET
              pp_partial_no = excluded.pp_partial_no,
              bom_code = excluded.bom_code,
              erp_status = excluded.erp_status,
              last_sync_at = CURRENT_TIMESTAMP
            """,
            (ps_id, partial_no, compact_text(row.get("bom_code")), compact_text(row.get("status"))),
        )

    return inserted, updated, skipped


def merge_erp_materials(con, sync_batch_id):
    inserted = updated = skipped = 0
    material_descriptions = bom_material_map(con)
    ps_by_key = ps_lookup_from_staging(con)
    materials_by_ps = {}

    for row in rows(con.execute("SELECT * FROM erp_materials_per_ps_staging WHERE sync_batch_id = ?", (sync_batch_id,))):
        ps_id = ""
        part_no = compact_text(row.get("inventory_code"))
        partial_no = compact_text(row.get("pp_partial_no"))
        staged_ps = one(
            con.execute(
                "SELECT ps_id, pp_partial_no, bom_code FROM erp_process_sheets_staging WHERE pp_partial_no = ? AND part_no = ?",
                (partial_no, part_no),
            )
        )
        if staged_ps:
            ps_id = process_sheet_key(staged_ps["ps_id"], staged_ps["pp_partial_no"])
        if not ps_id:
            key_val = partial_no
            if key_val and ("pp_partial_no", key_val) in ps_by_key:
                ps_id = ps_by_key[("pp_partial_no", key_val)]
        if not ps_id:
            skipped += 1
            continue
        erp_link = one(con.execute("SELECT bom_code FROM process_sheet_erp_link WHERE ps_id = ?", (ps_id,))) or {}
        bom_code = erp_link.get("bom_code", "")
        material_code = compact_text(row.get("material_inventory_code"))
        label = material_descriptions.get((bom_code, part_no, material_code), material_code)
        materials_by_ps.setdefault(ps_id, []).append(label)

    for ps_id, labels in materials_by_ps.items():
        material_name = ", ".join(sorted(dict.fromkeys([label for label in labels if label])))
        existing = one(con.execute("SELECT mat_id FROM process_sheet_material WHERE ps_id = ?", (ps_id,)))
        if existing:
            con.execute(
                """
                UPDATE process_sheet_material
                SET material_name = ?, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ?
                """,
                (material_name, ps_id),
            )
            updated += 1
        else:
            con.execute(
                """
                INSERT INTO process_sheet_material (ps_id, material_name, material_ready, material_ready_qty, order_status)
                VALUES (?, ?, 0, 0, 'TO_ORDER')
                """,
                (ps_id, material_name),
            )
            inserted += 1
    return inserted, updated, skipped


def merge_erp_bom_flows(con):
    inserted = updated = skipped = 0
    route_rows = rows(
        con.execute(
            """
            SELECT DISTINCT inventory_code AS part_no, bom_code
            FROM erp_bom_op_stage_staging
            WHERE COALESCE(bom_code, '') <> '' AND COALESCE(inventory_code, '') <> ''
            """
        )
    )
    flow_by_key = {}
    for route in route_rows:
        bom_code = compact_text(route["bom_code"])
        part_no = compact_text(route["part_no"])
        key = (part_no, bom_code)
        part_id = find_or_create_part(con, part_no, "")
        existing_flow = one(
            con.execute(
                "SELECT flow_id FROM part_flow_header WHERE part_id = ? AND flow_code = ?",
                (part_id, bom_code),
            )
        )
        if existing_flow and flow_erp_sync_locked(con, int(existing_flow["flow_id"])):
            skipped += 1
            flow_by_key[key] = int(existing_flow["flow_id"])
            continue
        stage_count = one(
            con.execute(
                "SELECT COUNT(*) cnt FROM erp_bom_op_stage_staging WHERE bom_code = ? AND (inventory_code = ? OR COALESCE(inventory_code, '') = '')",
                (bom_code, part_no),
            )
        )["cnt"]
        if stage_count <= 0:
            flow_by_key[key] = None
            skipped += 1
            continue
        flow_id, _ = upsert_bom_flow(con, part_id, part_no, bom_code)
        flow_by_key[key] = flow_id
        if existing_flow:
            updated += 1
        else:
            inserted += 1

    ps_rows = rows(
        con.execute(
            """
            SELECT ps.ps_id, ps.part_id, ps.inv_code AS part_no, ps.selected_flow_id, l.bom_code
            FROM process_sheet ps
            JOIN process_sheet_erp_link l ON l.ps_id = ps.ps_id
            WHERE COALESCE(l.bom_code, '') <> '' AND COALESCE(ps.inv_code, '') <> ''
            """
        )
    )
    for ps in ps_rows:
        key = (compact_text(ps["part_no"]), compact_text(ps["bom_code"]))
        flow_id = flow_by_key.get(key)
        if flow_id and ps_erp_sync_locked(con, ps["ps_id"]):
            skipped += 1
            continue
        if flow_id and ps["selected_flow_id"] != flow_id:
            con.execute(
                "UPDATE process_sheet SET selected_flow_id = ?, updated_at = CURRENT_TIMESTAMP WHERE ps_id = ?",
                (flow_id, ps["ps_id"]),
            )
    return inserted, updated, skipped


def merge_erp_actuals(con, sync_batch_id):
    inserted = updated = skipped = 0
    ps_by_key = ps_lookup_from_staging(con)
    actuals = {}
    rejects = {}
    for row in rows(con.execute("SELECT * FROM erp_workorder_tracker_staging WHERE sync_batch_id = ?", (sync_batch_id,))):
        ps_id = ""
        source_pp_no = compact_text(row.get("source_pp_no"))
        partial_seq_no = compact_text(row.get("partial_seq_no"))
        inventory_code = compact_text(row.get("inventory_code"))
        bom_code = compact_text(row.get("bom_code"))
        if source_pp_no and partial_seq_no and ("ps_pair", source_pp_no, partial_seq_no) in ps_by_key:
            ps_id = ps_by_key[("ps_pair", source_pp_no, partial_seq_no)]
        if not ps_id:
            for key_name in ("source_pp_no", "partial_seq_no", "inventory_code", "bom_code"):
                key_val = compact_text(row.get(key_name))
                if key_name == "partial_seq_no" and key_val and ("pp_partial_no", key_val) in ps_by_key:
                    ps_id = ps_by_key[("pp_partial_no", key_val)]
                    break
                if key_name == "inventory_code" and key_val and ("part_no", key_val) in ps_by_key:
                    ps_id = ps_by_key[("part_no", key_val)]
                    break
                if key_name == "bom_code" and key_val and ("bom_code", key_val) in ps_by_key:
                    ps_id = ps_by_key[("bom_code", key_val)]
                    break
                if key_name == "source_pp_no" and key_val and ("ps_id", key_val) in ps_by_key:
                    ps_id = ps_by_key[("ps_id", key_val)]
                    break
        if not ps_id:
            skipped += 1
            continue
        actuals[ps_id] = max(actuals.get(ps_id, 0), parse_number(row.get("total_acc_qty_produced")))
        rejects[ps_id] = max(rejects.get(ps_id, 0), parse_number(row.get("total_rej_qty_produced")))
    for ps_id, actual_qty in actuals.items():
        existing = one(con.execute("SELECT ps_id FROM process_sheet_erp_actual WHERE ps_id = ?", (ps_id,)))
        con.execute(
            """
            INSERT INTO process_sheet_erp_actual (ps_id, actual_qty, reject_qty, updated_at)
            VALUES (?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id) DO UPDATE SET
              actual_qty = excluded.actual_qty,
              reject_qty = excluded.reject_qty,
              updated_at = CURRENT_TIMESTAMP
            """,
            (ps_id, actual_qty, rejects.get(ps_id, 0)),
        )
        if existing:
            updated += 1
        else:
            inserted += 1
    return inserted, updated, skipped


def recent_erp_syncs(con, limit=10):
    return rows(con.execute("SELECT * FROM erp_sync_log ORDER BY sync_id DESC LIMIT ?", (limit,)))


def run_erp_excel_sync(con, file_storage, active_orders_file=None):
    sync_batch_id = uuid4().hex
    workbook_name = getattr(file_storage, "filename", "") or "upload.xlsx"
    active_orders_name = getattr(active_orders_file, "filename", "") or ""
    workbook_label = workbook_name if not active_orders_name else f"{workbook_name} + {active_orders_name}"
    con.execute(
        """
        INSERT INTO erp_sync_log (sync_batch_id, source_name, workbook_name, status)
        VALUES (?, 'excel-upload', ?, 'STARTED')
        """,
        (sync_batch_id, workbook_label),
    )
    try:
        backfill_erp_sync_locks(con)
        workbook_data = load_erp_workbook(file_storage)
        if active_orders_file and getattr(active_orders_file, "filename", ""):
            workbook_data["Active Orders"] = load_active_orders_workbook(active_orders_file)
        clear_erp_staging(con)
        insert_erp_staging_rows(con, sync_batch_id, workbook_data)
        try:
            ps_inserted, ps_updated, ps_skipped = merge_erp_process_sheets(con, sync_batch_id)
        except Exception as exc:
            raise RuntimeError(f"process_sheets stage failed: {exc}") from exc
        try:
            mat_inserted, mat_updated, mat_skipped = merge_erp_materials(con, sync_batch_id)
        except Exception as exc:
            raise RuntimeError(f"materials stage failed: {exc}") from exc
        try:
            flow_inserted, flow_updated, flow_skipped = merge_erp_bom_flows(con)
        except Exception as exc:
            raise RuntimeError(f"flows stage failed: {exc}") from exc
        try:
            actual_inserted, actual_updated, actual_skipped = merge_erp_actuals(con, sync_batch_id)
        except Exception as exc:
            raise RuntimeError(f"actuals stage failed: {exc}") from exc
        backfill_erp_sync_locks(con)

        summary = {
            "process_sheets": {"inserted": ps_inserted, "updated": ps_updated, "skipped": ps_skipped},
            "materials": {"inserted": mat_inserted, "updated": mat_updated, "skipped": mat_skipped},
            "flows": {"inserted": flow_inserted, "updated": flow_updated, "skipped": flow_skipped},
            "actuals": {"inserted": actual_inserted, "updated": actual_updated, "skipped": actual_skipped},
        }
        total_inserted = ps_inserted + mat_inserted + flow_inserted + actual_inserted
        total_updated = ps_updated + mat_updated + flow_updated + actual_updated
        total_skipped = ps_skipped + mat_skipped + flow_skipped + actual_skipped
        con.execute(
            """
            UPDATE erp_sync_log
            SET status = 'SUCCESS', message = ?, inserted_count = ?, updated_count = ?, skipped_count = ?, completed_at = CURRENT_TIMESTAMP
            WHERE sync_batch_id = ?
            """,
            (json.dumps(summary), total_inserted, total_updated, total_skipped, sync_batch_id),
        )
        return {"sync_batch_id": sync_batch_id, "summary": summary}
    except Exception as exc:
        con.execute(
            """
            UPDATE erp_sync_log
            SET status = 'FAILED', message = ?, completed_at = CURRENT_TIMESTAMP
            WHERE sync_batch_id = ?
            """,
            (str(exc), sync_batch_id),
        )
        raise


def run_active_orders_only_sync(con, active_orders_file):
    sync_batch_id = uuid4().hex
    workbook_name = getattr(active_orders_file, "filename", "") or "active_orders.xlsx"
    con.execute(
        """
        INSERT INTO erp_sync_log (sync_batch_id, source_name, workbook_name, status)
        VALUES (?, 'active-orders-upload', ?, 'STARTED')
        """,
        (sync_batch_id, workbook_name),
    )
    try:
        backfill_erp_sync_locks(con)
        active_rows = load_active_orders_workbook(active_orders_file)
        insert_active_orders_staging_rows(con, sync_batch_id, active_rows)
        updated, skipped = apply_active_orders_statuses(con, sync_batch_id)
        backfill_erp_sync_locks(con)
        summary = {
            "active_order_status": {"inserted": 0, "updated": updated, "skipped": skipped},
        }
        con.execute(
            """
            UPDATE erp_sync_log
            SET status = 'SUCCESS', message = ?, inserted_count = ?, updated_count = ?, skipped_count = ?, completed_at = CURRENT_TIMESTAMP
            WHERE sync_batch_id = ?
            """,
            (json.dumps(summary), 0, updated, skipped, sync_batch_id),
        )
        return {"sync_batch_id": sync_batch_id, "summary": summary}
    except Exception as exc:
        con.execute(
            """
            UPDATE erp_sync_log
            SET status = 'FAILED', message = ?, completed_at = CURRENT_TIMESTAMP
            WHERE sync_batch_id = ?
            """,
            (str(exc), sync_batch_id),
        )
        raise


def hhmm(value):
    value = int(value or 0)
    return f"{value // 60:02d}:{value % 60:02d}"


def display_hhmm(value):
    # Row tables already show the plan date separately, so keep the time
    # display local to that day instead of exposing carry-over suffixes.
    value = int(value or 0) % 1440
    return hhmm(value)


def whole_minutes(value, minimum=1):
    return max(minimum, int(round(float(value or 0))))


def duration_for_units(cycle_time, units, setup_time=0):
    return whole_minutes((setup_time or 0) + (units * (cycle_time or 0)), minimum=0)


def max_units_fit(available_mins, cycle_time, units_remaining, setup_time=0):
    units_remaining = max(0, int(units_remaining))
    if units_remaining <= 0:
        return 0
    low, high = 0, units_remaining
    while low < high:
        mid = (low + high + 1) // 2
        if duration_for_units(cycle_time, mid, setup_time) <= available_mins:
            low = mid
        else:
            high = mid - 1
    return low


def ensure_db():
    first_run = not DB_PATH.exists()
    with db() as con:
        con.executescript((ROOT / "schema.sql").read_text(encoding="utf-8"))
        process_sheet_cols = [row[1] for row in con.execute("PRAGMA table_info(process_sheet)").fetchall()]
        stage_cols = [row[1] for row in con.execute("PRAGMA table_info(erp_bom_op_stage_staging)").fetchall()]
        planning_row_info = con.execute("PRAGMA table_info(planning_row)").fetchall()
        planning_row_cols = [row[1] for row in planning_row_info]
        history_row_cols = [row[1] for row in con.execute("PRAGMA table_info(history_row)").fetchall()]
        planning_envelope_cols = [row[1] for row in con.execute("PRAGMA table_info(planning_envelope)").fetchall()]
        history_envelope_cols = [row[1] for row in con.execute("PRAGMA table_info(history_envelope)").fetchall()]
        if process_sheet_cols and "source_ps_id" not in process_sheet_cols:
            con.execute("ALTER TABLE process_sheet ADD COLUMN source_ps_id TEXT DEFAULT ''")
        if process_sheet_cols and "pp_partial_no" not in process_sheet_cols:
            con.execute("ALTER TABLE process_sheet ADD COLUMN pp_partial_no TEXT DEFAULT '1'")
        if process_sheet_cols and "source_ps_id" in process_sheet_cols:
            con.execute("UPDATE process_sheet SET source_ps_id = COALESCE(NULLIF(source_ps_id, ''), ps_id) WHERE COALESCE(source_ps_id, '') = ''")
        if process_sheet_cols and "pp_partial_no" in process_sheet_cols:
            con.execute("UPDATE process_sheet SET pp_partial_no = COALESCE(NULLIF(pp_partial_no, ''), '1') WHERE COALESCE(pp_partial_no, '') = ''")
        override_cols = [row[1] for row in con.execute("PRAGMA table_info(process_sheet_local_override)").fetchall()]
        if override_cols and "planner_status_override" not in override_cols:
            con.execute("ALTER TABLE process_sheet_local_override ADD COLUMN planner_status_override TEXT DEFAULT ''")
        ps_lock_cols = [row[1] for row in con.execute("PRAGMA table_info(process_sheet_erp_sync_lock)").fetchall()]
        if not ps_lock_cols:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS process_sheet_erp_sync_lock (
                  ps_id TEXT PRIMARY KEY REFERENCES process_sheet(ps_id) ON DELETE CASCADE,
                  locked INTEGER NOT NULL DEFAULT 1,
                  updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
        flow_override_cols = [row[1] for row in con.execute("PRAGMA table_info(part_flow_local_override)").fetchall()]
        if not flow_override_cols:
            con.execute(
                """
                CREATE TABLE IF NOT EXISTS part_flow_local_override (
                  flow_id INTEGER PRIMARY KEY REFERENCES part_flow_header(flow_id) ON DELETE CASCADE,
                  erp_sync_locked INTEGER NOT NULL DEFAULT 0,
                  updated_at TEXT DEFAULT CURRENT_TIMESTAMP
                )
                """
            )
        elif "erp_sync_locked" not in flow_override_cols:
            con.execute("ALTER TABLE part_flow_local_override ADD COLUMN erp_sync_locked INTEGER NOT NULL DEFAULT 0")
        if stage_cols and "inventory_code" not in stage_cols:
            con.execute("ALTER TABLE erp_bom_op_stage_staging ADD COLUMN inventory_code TEXT DEFAULT ''")
        if stage_cols and "seq_index" not in stage_cols:
            con.execute("ALTER TABLE erp_bom_op_stage_staging ADD COLUMN seq_index INTEGER DEFAULT 0")
        if planning_row_cols and "actual_out_set" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN actual_out_set INTEGER NOT NULL DEFAULT 0")
            con.execute("UPDATE planning_row SET actual_out_set = 1 WHERE COALESCE(actual_out, 0) > 0")
        if planning_row_cols and "envelope_id" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN envelope_id INTEGER NOT NULL DEFAULT 0")
        if planning_row_cols and "seq_no" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN seq_no INTEGER NOT NULL DEFAULT 0")
        if planning_row_cols and "op_no" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN op_no TEXT DEFAULT ''")
        if planning_row_cols and "machine_id" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN machine_id INTEGER NOT NULL DEFAULT 0")
        if planning_row_cols and "machine_code" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN machine_code TEXT DEFAULT ''")
        if planning_row_cols and "flow_step_id" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN flow_step_id INTEGER NOT NULL DEFAULT 0")
        if planning_row_cols and "split_group_id" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN split_group_id TEXT DEFAULT ''")
        if planning_row_cols and "split_piece_id" not in planning_row_cols:
            con.execute("ALTER TABLE planning_row ADD COLUMN split_piece_id TEXT DEFAULT ''")
        if planning_envelope_cols and "split_group_id" not in planning_envelope_cols:
            con.execute("ALTER TABLE planning_envelope ADD COLUMN split_group_id TEXT DEFAULT ''")
        if planning_envelope_cols and "split_piece_id" not in planning_envelope_cols:
            con.execute("ALTER TABLE planning_envelope ADD COLUMN split_piece_id TEXT DEFAULT ''")
        legacy_block_col = "block" + "_id"
        if any(row[1] == legacy_block_col and int(row[3] or 0) == 1 for row in planning_row_info):
            con.commit()
            con.execute("PRAGMA foreign_keys = OFF")
            con.execute("DROP TABLE IF EXISTS planning_row_legacy")
            con.execute("ALTER TABLE planning_row RENAME TO planning_row_legacy")
            con.execute(
                """
                CREATE TABLE planning_row (
                  row_id INTEGER PRIMARY KEY AUTOINCREMENT,
                  envelope_id INTEGER NOT NULL DEFAULT 0,
                  ps_id TEXT NOT NULL REFERENCES process_sheet(ps_id) ON DELETE CASCADE,
                  plan_date TEXT NOT NULL,
                  start_min INTEGER NOT NULL,
                  end_min INTEGER NOT NULL,
                  qty REAL NOT NULL DEFAULT 0,
                  actual_out REAL NOT NULL DEFAULT 0,
                  actual_out_set INTEGER NOT NULL DEFAULT 0,
                  setup_mins REAL NOT NULL DEFAULT 0,
                  locked INTEGER DEFAULT 0,
                  seq_no INTEGER NOT NULL DEFAULT 0,
                  op_no TEXT DEFAULT '',
                  machine_id INTEGER NOT NULL DEFAULT 0,
                  machine_code TEXT DEFAULT '',
                  flow_step_id INTEGER NOT NULL DEFAULT 0,
                  split_group_id TEXT DEFAULT '',
                  split_piece_id TEXT DEFAULT ''
                )
                """
            )
            con.execute(
                """
                INSERT INTO planning_row (
                    row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                    setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                )
                SELECT
                    row_id, COALESCE(envelope_id, 0), ps_id, plan_date, start_min, end_min, qty,
                    COALESCE(actual_out, 0), COALESCE(actual_out_set, 0), COALESCE(setup_mins, 0),
                    COALESCE(locked, 0), COALESCE(seq_no, 0), COALESCE(op_no, ''),
                    COALESCE(machine_id, 0), COALESCE(machine_code, ''), COALESCE(flow_step_id, 0), '', ''
                FROM planning_row_legacy
                """
            )
            con.execute("DROP TABLE planning_row_legacy")
            con.execute("PRAGMA foreign_keys = ON")
            con.commit()
            planning_row_info = con.execute("PRAGMA table_info(planning_row)").fetchall()
            planning_row_cols = [row[1] for row in planning_row_info]
        if planning_envelope_cols:
            con.execute(
                """
                INSERT OR IGNORE INTO planning_envelope (
                    ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id,
                    total_qty, locked, archived, envelope_start, envelope_end
                )
                SELECT pb.ps_id, pb.seq_no, pb.op_no, pb.machine_id, pb.machine_code, pb.flow_step_id,
                       '', '', pb.total_qty, pb.locked, pb.archived, pb.envelope_start, pb.envelope_end
                FROM planning_block pb
                """
            )
            con.execute(
                """
                UPDATE planning_row
                SET envelope_id = COALESCE((
                    SELECT pe.envelope_id
                    FROM planning_envelope pe
                    WHERE pe.ps_id = planning_row.ps_id
                      AND pe.seq_no = planning_row.seq_no
                      AND pe.op_no = planning_row.op_no
                      AND pe.machine_id = planning_row.machine_id
                      AND pe.machine_code = planning_row.machine_code
                      AND pe.flow_step_id = planning_row.flow_step_id
                      AND COALESCE(pe.split_piece_id, '') = COALESCE(planning_row.split_piece_id, '')
                    LIMIT 1
                ), envelope_id)
                WHERE COALESCE(envelope_id, 0) = 0
                """
            )
        if planning_row_cols and planning_envelope_cols:
            _backfill_split_piece_ids(con)
            _sync_all_planning_envelopes(con)
        if history_row_cols and "actual_out_set" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN actual_out_set INTEGER NOT NULL DEFAULT 0")
            con.execute("UPDATE history_row SET actual_out_set = 1 WHERE COALESCE(actual_out, 0) > 0")
        if history_row_cols and "hist_envelope_id" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN hist_envelope_id INTEGER NOT NULL DEFAULT 0")
        if history_row_cols and "seq_no" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN seq_no INTEGER NOT NULL DEFAULT 0")
        if history_row_cols and "op_no" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN op_no TEXT DEFAULT ''")
        if history_row_cols and "machine_id" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN machine_id INTEGER NOT NULL DEFAULT 0")
        if history_row_cols and "machine_code" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN machine_code TEXT DEFAULT ''")
        if history_row_cols and "flow_step_id" not in history_row_cols:
            con.execute("ALTER TABLE history_row ADD COLUMN flow_step_id INTEGER NOT NULL DEFAULT 0")
        if history_row_cols and "seq_no" in history_row_cols and "op_no" in history_row_cols and "machine_code" in history_row_cols:
            con.execute(
                """
                UPDATE history_row
                SET seq_no = COALESCE(seq_no, 0),
                    op_no = COALESCE(op_no, ''),
                    machine_code = COALESCE(machine_code, '')
                WHERE COALESCE(seq_no, 0) = 0 OR COALESCE(op_no, '') = '' OR COALESCE(machine_code, '') = ''
                """
            )
        if history_envelope_cols:
            con.execute(
                """
                INSERT OR IGNORE INTO history_envelope (
                    ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id,
                    total_qty, actual_out, locked, archived_at
                )
                SELECT hr.ps_id, hr.seq_no, hr.op_no, COALESCE(hr.machine_id, 0), hr.machine_code, COALESCE(hr.flow_step_id, 0),
                       COALESCE(SUM(hr.qty), 0), COALESCE(SUM(hr.actual_out), 0), 0, CURRENT_TIMESTAMP
                FROM history_row hr
                GROUP BY hr.ps_id, hr.seq_no, hr.op_no, hr.machine_id, hr.machine_code, hr.flow_step_id
                """
            )
            con.execute(
                """
                UPDATE history_row
                SET hist_envelope_id = COALESCE((
                    SELECT he.hist_envelope_id
                    FROM history_envelope he
                    WHERE he.ps_id = history_row.ps_id
                      AND he.seq_no = history_row.seq_no
                      AND he.op_no = history_row.op_no
                      AND he.machine_code = history_row.machine_code
                      AND he.flow_step_id = history_row.flow_step_id
                    LIMIT 1
                ), hist_envelope_id)
                WHERE COALESCE(hist_envelope_id, 0) = 0
                """
            )
        if first_run or con.execute("SELECT COUNT(*) FROM parts").fetchone()[0] == 0:
            con.executescript((ROOT / "seed_demo.sql").read_text(encoding="utf-8"))
        today = date.today()
        for offset in range(-30, 120):
            d = today + timedelta(days=offset)
            con.execute(
                "INSERT OR IGNORE INTO calendar_days (work_date, is_working_day) VALUES (?, ?)",
                (d.isoformat(), 0 if d.weekday() >= 5 else 1),
            )
        backfill_erp_sync_locks(con)


def manual_actual_qty(con, ps_id):
    row = one(con.execute("SELECT actual_qty FROM process_sheet_manual_actual WHERE ps_id = ?", (ps_id,)))
    return row["actual_qty"] if row else 0


def erp_actual_qty(con, ps_id):
    row = one(con.execute("SELECT actual_qty FROM process_sheet_erp_actual WHERE ps_id = ?", (ps_id,)))
    return row["actual_qty"] if row else 0


def ps_completion_override(con, ps_id):
    row = one(con.execute("SELECT force_completed FROM process_sheet_local_override WHERE ps_id = ?", (ps_id,)))
    if not row:
        return None
    return "COMPLETED" if int(row.get("force_completed") or 0) == 1 else "INCOMPLETE"


def ps_force_completed(con, ps_id):
    return ps_completion_override(con, ps_id) == "COMPLETED"


def ps_erp_sync_locked(con, ps_id):
    row = one(con.execute("SELECT locked FROM process_sheet_erp_sync_lock WHERE ps_id = ?", (ps_id,)))
    return bool(row and int(row.get("locked") or 0) == 1)


def flow_erp_sync_locked(con, flow_id):
    row = one(con.execute("SELECT erp_sync_locked FROM part_flow_local_override WHERE flow_id = ?", (flow_id,)))
    return bool(row and int(row.get("erp_sync_locked") or 0) == 1)


def mark_ps_erp_sync_locked(con, ps_id, locked=1):
    if int(bool(locked)):
        con.execute(
            """
            INSERT INTO process_sheet_erp_sync_lock (ps_id, locked, updated_at)
            VALUES (?, 1, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id) DO UPDATE SET
              locked = 1,
              updated_at = CURRENT_TIMESTAMP
            """,
            (ps_id,),
        )
    else:
        con.execute("DELETE FROM process_sheet_erp_sync_lock WHERE ps_id = ?", (ps_id,))


def mark_flow_erp_sync_locked(con, flow_id, locked=1):
    con.execute(
        """
        INSERT INTO part_flow_local_override (flow_id, erp_sync_locked, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(flow_id) DO UPDATE SET
          erp_sync_locked = excluded.erp_sync_locked,
          updated_at = CURRENT_TIMESTAMP
        """,
        (flow_id, int(bool(locked))),
    )


def backfill_erp_sync_locks(con):
    locked_ps = 0
    unlocked_ps = 0
    locked_flows = 0

    ps_rows = rows(con.execute("SELECT ps_id, selected_flow_id, total_qty, status FROM process_sheet ORDER BY ps_id"))
    for ps in ps_rows:
        ps_id = ps["ps_id"]
        if not ps_id:
            continue
        totals = ps_totals(con, ps_id)
        truly_completed = compact_text(ps.get("status")) == "COMPLETED" and totals.get("planner_status") == "COMPLETED"
        before = ps_erp_sync_locked(con, ps_id)
        if truly_completed:
            mark_ps_erp_sync_locked(con, ps_id, 1)
            if not before:
                locked_ps += 1
        elif before:
            mark_ps_erp_sync_locked(con, ps_id, 0)
            unlocked_ps += 1

    flow_ids = [
        int(row["selected_flow_id"])
        for row in rows(
            con.execute(
                """
                SELECT DISTINCT ps.selected_flow_id
                FROM process_sheet ps
                JOIN planning_row pr ON pr.ps_id = ps.ps_id
                WHERE COALESCE(ps.selected_flow_id, 0) > 0
                """
            )
        )
    ]
    for flow_id in sorted(set(flow_ids)):
        before = flow_erp_sync_locked(con, flow_id)
        mark_flow_erp_sync_locked(con, flow_id, 1)
        if not before:
            locked_flows += 1

    return {"locked_ps": locked_ps, "unlocked_ps": unlocked_ps, "locked_flows": locked_flows}


def step_completion_override(con, ps_id, step_id):
    row = one(
        con.execute(
            "SELECT force_completed FROM process_sheet_step_local_override WHERE ps_id = ? AND step_id = ?",
            (ps_id, step_id),
        )
    )
    if not row:
        return None
    return "COMPLETED" if int(row.get("force_completed") or 0) == 1 else "INCOMPLETE"


def step_force_completed(con, ps_id, step_id):
    return step_completion_override(con, ps_id, step_id) == "COMPLETED"


def sr_series_number(ps_id):
    text = compact_text(ps_id).upper()
    # Support sheet IDs like N26-[SR]03 and N26-[SR]03::1.
    # We only need the leading series number before the SR marker.
    match = re.match(r"^[A-Z](\d+)-\[SR\]", text)
    return int(match.group(1)) if match else None


def erp_active_sheet_index(con):
    active_rows = rows(con.execute("SELECT ps_id, part_no, bom_code FROM erp_active_orders_staging"))
    ps_ids = {compact_text(row["ps_id"]) for row in active_rows if compact_text(row["ps_id"])}
    return {
        "available": bool(active_rows),
        "ps_ids": ps_ids,
        "route_keys": set(),
    }


def process_sheet_step_metrics(con, ps_id, steps):
    metrics = []
    overall_planned = 0.0
    overall_finished = 0.0
    timeline_rows = ps_timeline_rows(con, ps_id)
    active_by_step = {}
    for row in timeline_rows:
        step_id = int(row.get("flow_step_id") or 0)
        active_by_step.setdefault(step_id, []).append(row)
    hist_rows = rows(
        con.execute(
            """
            SELECT
                seq_no,
                COALESCE(SUM(qty), 0) planned_qty,
                COALESCE(SUM(actual_out), 0) actual_qty,
                COALESCE(SUM(CASE WHEN COALESCE(actual_out_set, 0) = 1 THEN COALESCE(actual_out, 0) ELSE COALESCE(qty, 0) END), 0) effective_qty
            FROM history_row
            WHERE ps_id = ?
            GROUP BY seq_no
            """,
            (ps_id,),
        )
    )
    hist_by_seq = {int(r["seq_no"] or 0): r for r in hist_rows}
    for step in steps:
        active_rows = active_by_step.get(int(step["step_id"]), [])
        active = {
            "planned_qty": sum(float(r.get("qty") or 0) for r in active_rows),
            "actual_qty": sum(
                float(r.get("actual_out") or 0)
                for r in active_rows
                if int(r.get("actual_out_set") or 0) == 1
            ),
            "effective_qty": sum(
                float(r.get("actual_out") or 0) if int(r.get("actual_out_set") or 0) == 1 else float(r.get("qty") or 0)
                for r in active_rows
            ),
        }
        hist = hist_by_seq.get(int(step["seq"] or 0), {"planned_qty": 0, "actual_qty": 0})
        planned_qty = float(active["effective_qty"] or 0) + float(hist.get("effective_qty") or 0)
        actual_qty = float(active["actual_qty"] or 0) + float(hist["actual_qty"] or 0)
        overall_planned += planned_qty
        metrics.append({
            "step_id": step["step_id"],
            "seq": step["seq"],
            "op_no": step["op_no"],
            "op_type": step["op_type"],
            "planned_qty": planned_qty,
            "raw_planned_qty": float(active["planned_qty"] or 0) + float(hist["planned_qty"] or 0),
            "actual_qty": actual_qty,
            "remaining_qty": max(0.0, float(step.get("total_qty") or 0) - actual_qty),
        })

    last = metrics[-1] if metrics else {}
    overall_finished = float(last.get("actual_qty") or 0)
    return {
        "steps": metrics,
        "overall_planned_qty": overall_planned,
        "overall_finished_qty": overall_finished,
        "last_op_seq": last.get("seq"),
        "last_op_step_id": last.get("step_id"),
        "last_op_no": last.get("op_no", ""),
        "last_op_type": last.get("op_type", ""),
        "last_op_planned_qty": last.get("planned_qty", 0),
        "last_op_actual_qty": last.get("actual_qty", 0),
    }


def compute_process_sheet_status(con, ps, erp_index=None):
    total_qty = float(ps["total_qty"] or 0)
    ps_id = ps["ps_id"]
    source_ps_id = process_sheet_source_id(ps)
    partial_no = process_sheet_partial_no(ps)
    sr_series = sr_series_number(source_ps_id)
    is_sr_special = sr_series is not None and sr_series >= 26
    manual_override = ps_completion_override(con, ps_id)
    erp_index = erp_active_sheet_index(con) if erp_index is None else erp_index
    link = one(con.execute("SELECT bom_code FROM process_sheet_erp_link WHERE ps_id = ?", (ps_id,))) or {}
    part_no = compact_text(ps.get("inv_code"))
    route_key = (part_no, compact_text(link.get("bom_code")))
    erp_active_match = bool(
        compact_text(source_ps_id) in erp_index["ps_ids"]
    )

    steps = flow_steps(con, ps.get("selected_flow_id")) if ps.get("selected_flow_id") else []
    step_metrics = process_sheet_step_metrics(con, ps_id, steps) if steps else {
        "steps": [],
        "overall_planned_qty": 0,
        "overall_finished_qty": 0,
        "last_op_seq": None,
        "last_op_step_id": None,
        "last_op_no": "",
        "last_op_type": "",
        "last_op_planned_qty": 0,
        "last_op_actual_qty": 0,
    }

    last_op_planned_qty = float(step_metrics["last_op_planned_qty"] or 0)
    last_op_actual_qty = float(step_metrics["last_op_actual_qty"] or 0)
    overall_planned_qty = float(step_metrics["overall_planned_qty"] or 0)
    overall_finished_qty = float(step_metrics["overall_finished_qty"] or 0)
    all_steps_planned = bool(step_metrics.get("steps")) and all(
        float(step.get("planned_qty") or 0) >= total_qty or bool(step.get("completed"))
        for step in step_metrics.get("steps", [])
    )

    status_flags = []
    status_reason = ""

    if manual_override == "COMPLETED":
        status = "COMPLETED"
        status_reason = "Manual complete override"
        status_flags.append("MANUAL_COMPLETE")
    elif manual_override == "INCOMPLETE":
        if erp_index["available"] and erp_active_match and last_op_actual_qty >= total_qty and total_qty > 0:
            status = "NEEDS_REVIEW"
            status_reason = "Manual incomplete override is temporary, but last operation actual has reached total qty"
            status_flags.extend(["MANUAL_ACTIVE", "LAST_OP_COMPLETE", "COMPLETION_REVIEW"])
        else:
            status = "UNPLANNED" if overall_planned_qty <= 0 else ("PARTIALLY_PLANNED" if (last_op_planned_qty <= 0 or not all_steps_planned) else "PLANNED")
            status_reason = "Manual active/incomplete override"
            status_flags.append("MANUAL_ACTIVE")
    elif is_sr_special:
        status = "UNPLANNED" if overall_planned_qty <= 0 else ("PARTIALLY_PLANNED" if (last_op_planned_qty <= 0 or not all_steps_planned) else "PLANNED")
        status_reason = "SR series 26+ is protected from auto-complete"
        status_flags.append("SR_PROTECTED")
    elif total_qty <= 0:
        status = "UNPLANNED"
        status_reason = "No required quantity"
        status_flags.append("NO_QTY")
    else:
        if erp_index["available"]:
            status_flags.append("ERP_ACTIVE_EXPORT" if erp_active_match else "ERP_NOT_IN_ACTIVE_EXPORT")
        else:
            status_flags.append("ERP_ACTIVE_EXPORT_UNAVAILABLE")

        if not erp_active_match:
            status = "COMPLETED"
            status_reason = "ERP active export no longer lists this sheet"
            status_flags.append("ERP_COMPLETE")
            if last_op_actual_qty > 0:
                status_flags.append("LAST_OP_PARTIAL")
        elif last_op_actual_qty >= total_qty and total_qty > 0:
            status = "NEEDS_REVIEW"
            status_reason = "ERP still lists this sheet as active, but last operation actual has reached total qty"
            status_flags.append("LAST_OP_COMPLETE")
            status_flags.append("COMPLETION_REVIEW")
        else:
            if overall_planned_qty <= 0:
                status = "UNPLANNED"
                status_reason = "ERP active export still lists this sheet, but nothing is planned yet"
                status_flags.append("LAST_OP_MISSING")
            elif last_op_planned_qty <= 0:
                status = "PARTIALLY_PLANNED"
                status_reason = "ERP active export still lists this sheet, but the last operation is not planned yet"
                status_flags.append("LAST_OP_MISSING")
            elif not all_steps_planned:
                status = "PARTIALLY_PLANNED"
                status_reason = "ERP active export still lists this sheet, but not every operation is fully planned"
                status_flags.append("PARTIAL_PLAN")
            else:
                status = "PLANNED"
                status_reason = "ERP active export still lists this sheet, and the full order is planned"
                status_flags.append("FULLY_PLANNED")
            if last_op_actual_qty > 0:
                status_flags.append("LAST_OP_PARTIAL")

    debug = {
        "ps_id": ps_id,
        "source_ps_id": source_ps_id,
        "pp_partial_no": partial_no,
        "process_sheet_key": ps_id,
        "is_sr": bool(sr_series is not None),
        "sr_series": sr_series,
        "is_sr_special": is_sr_special,
        "erp_active_export_available": bool(erp_index["available"]),
        "erp_active_exported": erp_active_match,
        "erp_raw_status": compact_text(ps.get("status")),
        "total_qty": total_qty,
        "planned_qty": min(total_qty, overall_planned_qty),
        "finished_qty": min(total_qty, last_op_actual_qty),
        "last_op_planned_qty": last_op_planned_qty,
        "last_op_actual_qty": last_op_actual_qty,
        "overall_planned_qty": overall_planned_qty,
        "overall_finished_qty": overall_finished_qty,
        "manual_override": manual_override,
        "manual_actual_qty": manual_actual_qty(con, ps_id),
        "erp_actual_qty": erp_actual_qty(con, ps_id),
        "status_reason": status_reason,
        "status_flags": status_flags,
        "last_op_seq": step_metrics["last_op_seq"],
        "last_op_step_id": step_metrics["last_op_step_id"],
        "last_op_no": step_metrics["last_op_no"],
        "last_op_type": step_metrics["last_op_type"],
    }
    return {
        "planner_status": status,
        "status_reason": status_reason,
        "status_flags": status_flags,
        "status_debug": debug,
        "planned_qty": debug["planned_qty"],
        "finished_qty": debug["finished_qty"],
        "remaining_qty": max(0, total_qty - debug["finished_qty"]),
    }


def latest_progress_snapshot(con, ps_id):
    return one(
        con.execute(
            """
            SELECT snapshot_id, ps_id, reason, payload, created_at
            FROM process_sheet_progress_snapshot
            WHERE ps_id = ?
            ORDER BY snapshot_id DESC
            LIMIT 1
            """,
            (ps_id,),
        )
    )


def row_actual_baseline_reason(seq_no):
    return f"ROW_ACTUAL_BASELINE_SEQ_{int(seq_no)}"


def latest_row_actual_baseline_snapshot(con, ps_id, seq_no):
    return one(
        con.execute(
            """
            SELECT snapshot_id, ps_id, reason, payload, created_at
            FROM process_sheet_progress_snapshot
            WHERE ps_id = ? AND reason = ?
            ORDER BY snapshot_id DESC
            LIMIT 1
            """,
            (ps_id, row_actual_baseline_reason(seq_no)),
        )
    )


def snapshot_row_actual_baseline(con, ps_id, seq_no):
    row_rows = [
        row
        for row in ps_timeline_rows(con, ps_id)
        if int(row.get("seq_no") or 0) == int(seq_no)
    ]
    payload = {
        "seq_no": int(seq_no),
        "row_state": [
            {
                "row_id": int(row["row_id"]),
                "plan_date": row["plan_date"],
                "start_min": int(row["start_min"]),
                "end_min": int(row["end_min"]),
                "qty": float(row.get("qty") or 0),
                "actual_out": float(row.get("actual_out") or 0),
                "actual_out_set": int(row.get("actual_out_set") or 0),
                "setup_mins": float(row.get("setup_mins") or 0),
                "locked": int(row.get("locked") or 0),
            }
            for row in row_rows
        ],
    }
    cur = con.execute(
        """
        INSERT INTO process_sheet_progress_snapshot (ps_id, reason, payload)
        VALUES (?, ?, ?)
        """,
        (ps_id, row_actual_baseline_reason(seq_no), json.dumps(payload)),
    )
    return cur.lastrowid


def snapshot_process_sheet_progress(con, ps_id, reason="MANUAL_RESET"):
    manual_row = one(con.execute("SELECT actual_qty FROM process_sheet_manual_actual WHERE ps_id = ?", (ps_id,)))
    ps_override = one(con.execute("SELECT force_completed FROM process_sheet_local_override WHERE ps_id = ?", (ps_id,)))
    step_overrides = rows(
        con.execute(
            """
            SELECT step_id, force_completed
            FROM process_sheet_step_local_override
            WHERE ps_id = ?
            ORDER BY step_id
            """,
            (ps_id,),
        )
    )
    row_actuals = rows(
        con.execute(
            """
            SELECT row_id, plan_date, start_min, end_min, actual_out, actual_out_set
            FROM planning_row
            WHERE ps_id = ? AND (COALESCE(actual_out, 0) > 0 OR COALESCE(actual_out_set, 0) = 1)
            ORDER BY plan_date, start_min, row_id
            """,
            (ps_id,),
        )
    )
    payload = {
        "manual_actual_qty": float((manual_row or {}).get("actual_qty") or 0),
        "ps_override": None if not ps_override else int(ps_override.get("force_completed") or 0),
        "step_overrides": [
            {"step_id": int(row["step_id"]), "force_completed": int(row.get("force_completed") or 0)}
            for row in step_overrides
        ],
        "row_actuals": [
            {
                "row_id": int(row["row_id"]),
                "plan_date": row["plan_date"],
                "start_min": int(row["start_min"]),
                "end_min": int(row["end_min"]),
                "actual_out": float(row.get("actual_out") or 0),
                "actual_out_set": int(row.get("actual_out_set") or 0),
            }
            for row in row_actuals
        ],
    }
    cur = con.execute(
        """
        INSERT INTO process_sheet_progress_snapshot (ps_id, reason, payload)
        VALUES (?, ?, ?)
        """,
        (ps_id, reason, json.dumps(payload)),
    )
    return cur.lastrowid


def reset_process_sheet_progress(con, ps_id):
    con.execute(
        "UPDATE planning_row SET actual_out = 0, actual_out_set = 0 WHERE ps_id = ?",
        (ps_id,),
    )
    con.execute(
        """
        INSERT INTO process_sheet_manual_actual (ps_id, actual_qty, updated_at)
        VALUES (?, 0, CURRENT_TIMESTAMP)
        ON CONFLICT(ps_id) DO UPDATE SET actual_qty = 0, updated_at = CURRENT_TIMESTAMP
        """,
        (ps_id,),
    )
    con.execute(
        """
        INSERT INTO process_sheet_local_override (ps_id, force_completed, updated_at)
        VALUES (?, 0, CURRENT_TIMESTAMP)
        ON CONFLICT(ps_id) DO UPDATE SET force_completed = 0, updated_at = CURRENT_TIMESTAMP
        """,
        (ps_id,),
    )
    con.execute("DELETE FROM process_sheet_step_local_override WHERE ps_id = ?", (ps_id,))
    step_ids = [
        row["step_id"]
        for row in con.execute(
            """
            SELECT pfs.step_id
            FROM process_sheet ps
            JOIN part_flow_steps pfs ON pfs.flow_id = ps.selected_flow_id
            WHERE ps.ps_id = ?
            ORDER BY pfs.seq_no, pfs.step_id
            """,
            (ps_id,),
        )
    ]
    for step_id in step_ids:
        con.execute(
            """
            INSERT INTO process_sheet_step_local_override (ps_id, step_id, force_completed, updated_at)
            VALUES (?, ?, 0, CURRENT_TIMESTAMP)
            """,
            (ps_id, step_id),
        )


def restore_process_sheet_progress(con, ps_id, snapshot_id=None):
    snapshot = latest_progress_snapshot(con, ps_id) if snapshot_id is None else one(
        con.execute(
            """
            SELECT snapshot_id, ps_id, reason, payload, created_at
            FROM process_sheet_progress_snapshot
            WHERE ps_id = ? AND snapshot_id = ?
            """,
            (ps_id, snapshot_id),
        )
    )
    if not snapshot:
        raise ValueError("No saved progress snapshot was found")

    payload = json.loads(snapshot.get("payload") or "{}")
    con.execute(
        "UPDATE planning_row SET actual_out = 0, actual_out_set = 0 WHERE ps_id = ?",
        (ps_id,),
    )

    restored_rows = 0
    for row in payload.get("row_actuals", []):
        matched = one(con.execute("SELECT row_id FROM planning_row WHERE row_id = ? AND ps_id = ?", (row.get("row_id"), ps_id)))
        if matched:
            con.execute(
                """
                UPDATE planning_row
                SET actual_out = ?, actual_out_set = ?
                WHERE row_id = ? AND ps_id = ?
                """,
                (float(row.get("actual_out") or 0), int(row.get("actual_out_set") or 0), int(row["row_id"]), ps_id),
            )
            restored_rows += 1
            continue

        matched = one(
            con.execute(
                """
                SELECT row_id
                FROM planning_row
                WHERE ps_id = ? AND envelope_id = ? AND plan_date = ? AND start_min = ? AND end_min = ?
                """,
                (
                    ps_id,
                    int(row.get("envelope_id") or 0),
                    row.get("plan_date"),
                    int(row.get("start_min") or 0),
                    int(row.get("end_min") or 0),
                ),
            )
        )
        if matched:
            con.execute(
                """
                UPDATE planning_row
                SET actual_out = ?, actual_out_set = ?
                WHERE row_id = ? AND ps_id = ?
                """,
                (float(row.get("actual_out") or 0), int(row.get("actual_out_set") or 0), int(matched["row_id"]), ps_id),
            )
            restored_rows += 1

    manual_actual = float(payload.get("manual_actual_qty") or 0)
    con.execute(
        """
        INSERT INTO process_sheet_manual_actual (ps_id, actual_qty, updated_at)
        VALUES (?, ?, CURRENT_TIMESTAMP)
        ON CONFLICT(ps_id) DO UPDATE SET actual_qty = excluded.actual_qty, updated_at = CURRENT_TIMESTAMP
        """,
        (ps_id, manual_actual),
    )

    if payload.get("ps_override") is None:
        con.execute("DELETE FROM process_sheet_local_override WHERE ps_id = ?", (ps_id,))
    else:
        con.execute(
            """
            INSERT INTO process_sheet_local_override (ps_id, force_completed, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id) DO UPDATE SET force_completed = excluded.force_completed, updated_at = CURRENT_TIMESTAMP
            """,
            (ps_id, int(payload.get("ps_override") or 0)),
        )

    con.execute("DELETE FROM process_sheet_step_local_override WHERE ps_id = ?", (ps_id,))
    for row in payload.get("step_overrides", []):
        con.execute(
            """
            INSERT INTO process_sheet_step_local_override (ps_id, step_id, force_completed, updated_at)
            VALUES (?, ?, ?, CURRENT_TIMESTAMP)
            """,
            (ps_id, int(row["step_id"]), int(row.get("force_completed") or 0)),
        )

    return {
        "snapshot_id": int(snapshot["snapshot_id"]),
        "restored_rows": restored_rows,
    }


def flow_steps(con, flow_id):
    ensure_part_flow_step_columns(con)
    return rows(
        con.execute(
            """
            SELECT step_id, seq_no AS seq, op_no, op_type, machine_category, preferred_machine,
                   cycle_time, setup_time, is_last_op
            FROM part_flow_steps WHERE flow_id = ? ORDER BY seq_no, step_id
            """,
            (flow_id,),
        )
    )


def route_label(con, flow_id):
    if not flow_id:
        return ""
    return ", ".join(step["op_no"] for step in flow_steps(con, flow_id))


def material_for_ps(con, ps_id):
    mat = one(con.execute("SELECT * FROM process_sheet_material WHERE ps_id = ?", (ps_id,)))
    if not mat:
        return None
    mat["order_logs"] = rows(
        con.execute("SELECT * FROM process_sheet_material_order_log WHERE mat_id = ? ORDER BY order_date, log_id", (mat["mat_id"],))
    )
    return mat


def material_for_ps_cached(con, ps_id, cache=None):
    if cache is None:
        return material_for_ps(con, ps_id)
    if ps_id not in cache:
        cache[ps_id] = material_for_ps(con, ps_id)
    return cache[ps_id]


def ensure_material_record(con, ps_id):
    con.execute(
        """
        INSERT OR IGNORE INTO process_sheet_material (ps_id, material_name, material_ready, material_ready_qty, order_status)
        VALUES (?, '', 0, 0, 'TO_ORDER')
        """,
        (ps_id,),
    )
    return material_for_ps(con, ps_id)


def support_type_lead_days(support_type):
    support_type = compact_text(support_type).upper()
    if support_type == "MATERIAL":
        return MATERIAL_LEAD_DAYS
    if support_type == "TOOLLIST":
        return TOOLLIST_LEAD_DAYS
    return PROGRAM_LEAD_DAYS


def ensure_support_record(con, ps_id, support_type):
    support_type = compact_text(support_type).upper()
    con.execute(
        """
        INSERT OR IGNORE INTO process_sheet_support (ps_id, support_type, status)
        VALUES (?, ?, 'PENDING')
        """,
        (ps_id, support_type),
    )
    row = one(con.execute("SELECT * FROM process_sheet_support WHERE ps_id = ? AND support_type = ?", (ps_id, support_type)))
    return row


def support_record_for_ps(con, ps_id, support_type):
    return one(con.execute("SELECT * FROM process_sheet_support WHERE ps_id = ? AND support_type = ?", (ps_id, compact_text(support_type).upper())))


def sync_support_need_by(con, ps_id, support_type, expected_start):
    support_type = compact_text(support_type).upper()
    record = ensure_support_record(con, ps_id, support_type)
    need_by_date = subtract_working_days(con, expected_start[:10] if expected_start else "", support_type_lead_days(support_type))
    if (record.get("need_by_date") or "") != need_by_date:
        con.execute(
            "UPDATE process_sheet_support SET need_by_date = ?, updated_at = CURRENT_TIMESTAMP WHERE support_id = ?",
            (need_by_date, record["support_id"]),
        )
        record["need_by_date"] = need_by_date
    return record


def ps_totals(con, ps_id, light=False, erp_index=None):
    ps = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
    if not ps:
        return {"planned_qty": 0, "finished_qty": 0, "last_op_planned_qty": 0, "last_op_actual_qty": 0}
    if light:
        return compute_process_sheet_status_light(con, ps, erp_index=erp_index)
    return compute_process_sheet_status(con, ps, erp_index=erp_index)


def build_process_sheet_list_cache(con, ps_rows):
    ps_ids = [ps["ps_id"] for ps in ps_rows if ps.get("ps_id")]
    if not ps_ids:
        return {
            "timeline_by_ps": {},
            "history_by_ps": {},
            "override_by_ps": {},
            "span_by_ps": {},
        }

    placeholders = ",".join("?" for _ in ps_ids)

    timeline_by_ps = {}
    for row in rows(
        con.execute(
            f"""
            SELECT
                pr.row_id,
                pr.plan_date,
                pr.start_min,
                pr.end_min,
                pr.qty,
                pr.actual_out,
                pr.actual_out_set,
                pr.setup_mins,
                pr.locked AS row_locked,
                pr.seq_no,
                pr.op_no,
                pr.machine_id,
                pr.machine_code,
                pr.flow_step_id,
                pr.ps_id,
                pr.envelope_id,
                pr.split_group_id,
                pr.split_piece_id
            FROM planning_row pr
            WHERE pr.ps_id IN ({placeholders})
            ORDER BY pr.ps_id, pr.plan_date, pr.start_min, pr.end_min, pr.seq_no, pr.row_id
            """,
            ps_ids,
        )
    ):
        timeline_by_ps.setdefault(row["ps_id"], []).append(row)

    history_by_ps = {}
    for row in rows(
        con.execute(
            f"""
            SELECT
                ps_id,
                seq_no,
                COALESCE(SUM(qty), 0) AS planned_qty,
                COALESCE(SUM(actual_out), 0) AS actual_qty
            FROM history_row
            WHERE ps_id IN ({placeholders})
            GROUP BY ps_id, seq_no
            """,
            ps_ids,
        )
    ):
        history_by_ps.setdefault(row["ps_id"], []).append(row)

    override_by_ps = {
        row["ps_id"]: ("COMPLETED" if int(row.get("force_completed") or 0) == 1 else "INCOMPLETE")
        for row in rows(
            con.execute(
                f"""
                SELECT ps_id, force_completed
                FROM process_sheet_local_override
                WHERE ps_id IN ({placeholders})
                """,
                ps_ids,
            )
        )
    }

    span_by_ps = {}
    for ps_id, timeline_rows in timeline_by_ps.items():
        if not timeline_rows:
            continue
        span_by_ps[ps_id] = {
            "expected_start": min((f"{r['plan_date']} {hhmm(r['start_min'])}" for r in timeline_rows), default=""),
            "expected_end": max((f"{r['plan_date']} {hhmm(r['end_min'])}" for r in timeline_rows), default=""),
            "machines_used": ",".join(sorted({compact_text(r.get("machine_code")) for r in timeline_rows if compact_text(r.get("machine_code"))})),
        }

    return {
        "timeline_by_ps": timeline_by_ps,
        "history_by_ps": history_by_ps,
        "override_by_ps": override_by_ps,
        "span_by_ps": span_by_ps,
    }


def compute_process_sheet_status_light(con, ps, erp_index=None, cache=None):
    total_qty = float(ps["total_qty"] or 0)
    ps_id = ps["ps_id"]
    source_ps_id = process_sheet_source_id(ps)
    partial_no = process_sheet_partial_no(ps)
    sr_series = sr_series_number(source_ps_id)
    is_sr_special = sr_series is not None and sr_series >= 26
    erp_index = erp_active_sheet_index(con) if erp_index is None else erp_index
    erp_active_match = bool(compact_text(source_ps_id) in erp_index["ps_ids"])
    cache = cache or {}
    manual_override = cache.get("override_by_ps", {}).get(ps_id)
    timeline_rows = cache.get("timeline_by_ps", {}).get(ps_id)
    if timeline_rows is None:
        timeline_rows = ps_timeline_rows(con, ps_id)

    seq_metrics = {}
    for row in timeline_rows:
        seq = int(row.get("seq_no") or 0)
        bucket = seq_metrics.setdefault(seq, {"planned_qty": 0.0, "actual_qty": 0.0, "effective_qty": 0.0})
        qty = float(row.get("qty") or 0)
        bucket["planned_qty"] += qty
        if int(row.get("actual_out_set") or 0) == 1:
            actual_qty = float(row.get("actual_out") or 0)
            bucket["actual_qty"] += actual_qty
            bucket["effective_qty"] += actual_qty
        else:
            bucket["effective_qty"] += qty

    hist_rows = cache.get("history_by_ps", {}).get(ps_id)
    if hist_rows is None:
        hist_rows = rows(
            con.execute(
                """
                SELECT seq_no,
                       COALESCE(SUM(qty), 0) planned_qty,
                       COALESCE(SUM(actual_out), 0) actual_qty,
                       COALESCE(SUM(CASE WHEN COALESCE(actual_out_set, 0) = 1 THEN COALESCE(actual_out, 0) ELSE COALESCE(qty, 0) END), 0) effective_qty
                FROM history_row
                WHERE ps_id = ?
                GROUP BY seq_no
                """,
                (ps_id,),
            )
        )
    for row in hist_rows:
        seq = int(row["seq_no"] or 0)
        bucket = seq_metrics.setdefault(seq, {"planned_qty": 0.0, "actual_qty": 0.0, "effective_qty": 0.0})
        bucket["planned_qty"] += float(row["planned_qty"] or 0)
        bucket["actual_qty"] += float(row["actual_qty"] or 0)
        bucket["effective_qty"] += float(row["effective_qty"] or 0)

    steps = flow_steps(con, ps.get("selected_flow_id")) if ps.get("selected_flow_id") else []
    step_metrics = []
    overall_planned_qty = 0.0
    overall_finished_qty = 0.0
    last_op_planned_qty = 0.0
    last_op_actual_qty = 0.0
    last_seq = None
    last_step_id = None
    last_op_no = ""
    last_op_type = ""
    if steps:
        for step in steps:
            seq = int(step.get("seq") or 0)
            bucket = seq_metrics.get(seq, {"planned_qty": 0.0, "actual_qty": 0.0, "effective_qty": 0.0})
            planned_qty = min(total_qty, float(bucket.get("effective_qty") or 0))
            actual_qty = float(bucket.get("actual_qty") or 0)
            completed = step_completion_override(con, ps_id, step["step_id"]) == "COMPLETED" or actual_qty >= total_qty
            overall_planned_qty += planned_qty
            step_metrics.append(
                {
                    "step_id": step["step_id"],
                    "seq": step["seq"],
                    "op_no": step["op_no"],
                    "op_type": step["op_type"],
                    "planned_qty": planned_qty,
                    "actual_qty": actual_qty,
                    "completed": completed,
                }
            )
        last_step = step_metrics[-1] if step_metrics else {}
        last_op_planned_qty = float(last_step.get("planned_qty") or 0)
        last_op_actual_qty = float(last_step.get("actual_qty") or 0)
        overall_finished_qty = last_op_actual_qty
        last_seq = last_step.get("seq")
        last_step_id = last_step.get("step_id")
        last_op_no = last_step.get("op_no", "")
        last_op_type = last_step.get("op_type", "")
    else:
        seq_keys = sorted(seq_metrics.keys())
        last_seq = seq_keys[-1] if seq_keys else None
        last_op_planned_qty = min(total_qty, float(seq_metrics.get(last_seq, {}).get("effective_qty", 0) or 0)) if last_seq is not None else 0.0
        last_op_actual_qty = float(seq_metrics.get(last_seq, {}).get("actual_qty", 0) or 0) if last_seq is not None else 0.0
        overall_planned_qty = sum(min(total_qty, float(v.get("effective_qty") or 0)) for v in seq_metrics.values())
        overall_finished_qty = last_op_actual_qty
        last_step_id = None
        last_op_no = ""
        last_op_type = ""
    all_steps_planned = bool(step_metrics) and all(
        float(step.get("planned_qty") or 0) >= total_qty or bool(step.get("completed"))
        for step in step_metrics
    )

    status_flags = []
    status_reason = ""

    if manual_override == "COMPLETED":
        status = "COMPLETED"
        status_reason = "Manual complete override"
        status_flags.append("MANUAL_COMPLETE")
    elif manual_override == "INCOMPLETE":
        if erp_index["available"] and erp_active_match and last_op_actual_qty >= total_qty and total_qty > 0:
            status = "NEEDS_REVIEW"
            status_reason = "Manual incomplete override is temporary, but last operation actual has reached total qty"
            status_flags.extend(["MANUAL_ACTIVE", "LAST_OP_COMPLETE", "COMPLETION_REVIEW"])
        else:
            status = "UNPLANNED" if overall_planned_qty <= 0 else ("PARTIALLY_PLANNED" if (last_op_planned_qty <= 0 or not all_steps_planned) else "PLANNED")
            status_reason = "Manual active/incomplete override"
            status_flags.append("MANUAL_ACTIVE")
    elif is_sr_special:
        status = "UNPLANNED" if overall_planned_qty <= 0 else ("PARTIALLY_PLANNED" if (last_op_planned_qty <= 0 or not all_steps_planned) else "PLANNED")
        status_reason = "SR series 26+ is protected from auto-complete"
        status_flags.append("SR_PROTECTED")
    elif total_qty <= 0:
        status = "UNPLANNED"
        status_reason = "No required quantity"
        status_flags.append("NO_QTY")
    else:
        if erp_index["available"]:
            status_flags.append("ERP_ACTIVE_EXPORT" if erp_active_match else "ERP_NOT_IN_ACTIVE_EXPORT")
        else:
            status_flags.append("ERP_ACTIVE_EXPORT_UNAVAILABLE")

        if not erp_active_match:
            status = "COMPLETED"
            status_reason = "ERP active export no longer lists this sheet"
            status_flags.append("ERP_COMPLETE")
            if last_op_actual_qty > 0:
                status_flags.append("LAST_OP_PARTIAL")
        elif last_op_actual_qty >= total_qty and total_qty > 0:
            status = "NEEDS_REVIEW"
            status_reason = "ERP still lists this sheet as active, but last operation actual has reached total qty"
            status_flags.append("LAST_OP_COMPLETE")
            status_flags.append("COMPLETION_REVIEW")
        else:
            if overall_planned_qty <= 0:
                status = "UNPLANNED"
                status_reason = "ERP active export still lists this sheet, but nothing is planned yet"
                status_flags.append("LAST_OP_MISSING")
            elif last_op_planned_qty <= 0:
                status = "PARTIALLY_PLANNED"
                status_reason = "ERP active export still lists this sheet, but the last operation is not planned yet"
                status_flags.append("LAST_OP_MISSING")
            elif not all_steps_planned:
                status = "PARTIALLY_PLANNED"
                status_reason = "ERP active export still lists this sheet, but not every operation is fully planned"
                status_flags.append("PARTIAL_PLAN")
            else:
                status = "PLANNED"
                status_reason = "ERP active export still lists this sheet, and the full order is planned"
                status_flags.append("FULLY_PLANNED")
            if last_op_actual_qty > 0:
                status_flags.append("LAST_OP_PARTIAL")

    debug = {
        "ps_id": ps_id,
        "source_ps_id": source_ps_id,
        "pp_partial_no": partial_no,
        "process_sheet_key": ps_id,
        "is_sr": bool(sr_series is not None),
        "sr_series": sr_series,
        "is_sr_special": is_sr_special,
        "erp_active_export_available": bool(erp_index["available"]),
        "erp_active_exported": erp_active_match,
        "erp_raw_status": compact_text(ps.get("status")),
        "total_qty": total_qty,
        "planned_qty": min(total_qty, overall_planned_qty),
        "finished_qty": min(total_qty, last_op_actual_qty),
        "last_op_planned_qty": last_op_planned_qty,
        "last_op_actual_qty": last_op_actual_qty,
        "overall_planned_qty": overall_planned_qty,
        "overall_finished_qty": overall_finished_qty,
        "manual_override": manual_override,
        "status_reason": status_reason,
        "status_flags": status_flags,
        "last_op_seq": last_seq,
        "last_op_step_id": last_step_id,
        "last_op_no": last_op_no,
        "last_op_type": last_op_type,
    }
    return {
        "planner_status": status,
        "status_reason": status_reason,
        "status_flags": status_flags,
        "status_debug": debug,
        "planned_qty": debug["planned_qty"],
        "finished_qty": debug["finished_qty"],
        "last_op_planned_qty": last_op_planned_qty,
        "last_op_actual_qty": last_op_actual_qty,
        "overall_planned_qty": overall_planned_qty,
        "overall_finished_qty": overall_finished_qty,
        "last_op_seq": last_seq,
        "last_op_step_id": None,
        "last_op_no": "",
        "last_op_type": "",
        "remaining_qty": max(0, total_qty - overall_planned_qty),
    }


def step_progress(con, ps_id, step, total_qty):
    total_qty = float(total_qty or 0)
    seq_no = step["seq"]
    flow_step_id = step["step_id"]
    active_rows = [r for r in ps_timeline_rows(con, ps_id) if int(r.get("flow_step_id") or 0) == int(flow_step_id)]
    active = {
        "planned_qty": sum(float(r.get("qty") or 0) for r in active_rows),
        "actual_qty": sum(float(r.get("actual_out") or 0) for r in active_rows),
        "effective_qty": sum(
            float(r.get("actual_out") or 0) if int(r.get("actual_out_set") or 0) == 1 else float(r.get("qty") or 0)
            for r in active_rows
        ),
        "expected_start": min(
            (f"{r['plan_date']} {hhmm(r['start_min'])}" for r in active_rows),
            default="",
        ),
        "expected_end": max(
            (f"{r['plan_date']} {hhmm(r['end_min'])}" for r in active_rows),
            default="",
        ),
    }
    hist = one(
        con.execute(
            """
            SELECT COALESCE(SUM(hr.qty), 0) planned_qty,
                   COALESCE(SUM(hr.actual_out), 0) actual_qty,
                   COALESCE(SUM(CASE WHEN COALESCE(hr.actual_out_set, 0) = 1 THEN COALESCE(hr.actual_out, 0) ELSE COALESCE(hr.qty, 0) END), 0) effective_qty,
                   MIN(hr.plan_date || ' ' || printf('%02d:%02d', hr.start_min / 60, hr.start_min % 60)) expected_start,
                   MAX(hr.plan_date || ' ' || printf('%02d:%02d', hr.end_min / 60, hr.end_min % 60)) expected_end
            FROM history_row hr
            WHERE hr.ps_id = ? AND hr.seq_no = ?
            """,
            (ps_id, seq_no),
        )
    )
    raw_planned_qty = float((active["planned_qty"] or 0) + (hist["planned_qty"] or 0))
    finished_qty = float((active["actual_qty"] or 0) + (hist["actual_qty"] or 0))
    planned_qty = float((active["effective_qty"] or 0) + (hist["effective_qty"] or 0))
    expected_start = active["expected_start"] or hist["expected_start"] or ""
    expected_end = active["expected_end"] or hist["expected_end"] or ""
    if total_qty <= 0:
        plan_status = "UNPLANNED"
    elif finished_qty >= total_qty:
        plan_status = "COMPLETED"
    elif planned_qty >= total_qty:
        plan_status = "PLANNED"
    elif planned_qty > 0 or finished_qty > 0:
        plan_status = "PARTIALLY_PLANNED"
    else:
        plan_status = "UNPLANNED"
    return {
        "planned_qty": min(total_qty, planned_qty),
        "raw_planned_qty": min(total_qty, raw_planned_qty),
        "finished_qty": min(total_qty, finished_qty),
        "remaining_qty": max(0, total_qty - planned_qty),
        "expected_start": expected_start,
        "expected_end": expected_end,
        "plan_status": plan_status,
        "warnings": [] if plan_status == "COMPLETED" else (["UNPLANNED"] if plan_status == "UNPLANNED" else []),
    }


def planner_status(con, ps, totals):
    return compute_process_sheet_status(con, ps)["planner_status"]


def warnings_for(con, ps, totals, material_cache=None, support_sync=True, schedule=None):
    status_info = totals if isinstance(totals, dict) and totals.get("planner_status") else compute_process_sheet_status(con, ps)
    if status_info["planner_status"] == "COMPLETED":
        return []
    if float(ps["total_qty"] or 0) <= 0:
        return []
    warnings = []
    if ps["due_date"] and ps["due_date"] < date.today().isoformat() and status_info["planner_status"] != "COMPLETED":
        warnings.append("OVERDUE")
    if status_info["planned_qty"] <= 0:
        warnings.append("UNPLANNED")
    if status_info["planner_status"] == "NEEDS_REVIEW":
        warnings.append("COMPLETION_REVIEW")
    mat = material_for_ps_cached(con, ps["ps_id"], material_cache)
    if mat:
        covered = (mat["material_ready_qty"] or 0) + sum(log["received_qty"] or 0 for log in mat["order_logs"])
        if not mat["material_ready"]:
            warnings.append("MATERIAL_PENDING")
        if covered < ps["total_qty"]:
            warnings.append("MATERIAL_SHORTAGE")
    span = schedule or schedule_span(con, ps["ps_id"]) or {}
    expected_start = span.get("expected_start") or ""
    if expected_start:
        if support_sync:
            program = sync_support_need_by(con, ps["ps_id"], "PROGRAM", expected_start)
            toollist = sync_support_need_by(con, ps["ps_id"], "TOOLLIST", expected_start)
            for prefix, record in (("PROGRAM", program), ("TOOLLIST", toollist)):
                status = compact_text(record.get("status")).upper()
                promised = compact_text(record.get("promised_date"))
                need_by = compact_text(record.get("need_by_date"))
                if status not in {"READY", "NOT_REQUIRED"}:
                    warnings.append(f"{prefix}_PENDING")
                if promised and need_by and promised > need_by:
                    warnings.append(f"{prefix}_LATE")
        else:
            program = support_record_for_ps(con, ps["ps_id"], "PROGRAM")
            toollist = support_record_for_ps(con, ps["ps_id"], "TOOLLIST")
            for prefix, record in (("PROGRAM", program), ("TOOLLIST", toollist)):
                if not record:
                    warnings.append(f"{prefix}_PENDING")
                    continue
                status = compact_text(record.get("status")).upper()
                promised = compact_text(record.get("promised_date"))
                need_by = compact_text(record.get("need_by_date"))
                if status not in {"READY", "NOT_REQUIRED"}:
                    warnings.append(f"{prefix}_PENDING")
                if promised and need_by and promised > need_by:
                    warnings.append(f"{prefix}_LATE")
    return warnings


def step_statuses(con, ps):
    steps = flow_steps(con, ps["selected_flow_id"])
    planned = {r["flow_step_id"] for r in con.execute("SELECT DISTINCT flow_step_id FROM planning_block WHERE ps_id = ? AND archived = 0", (ps["ps_id"],))}
    history = {r["op_no"] for r in con.execute("SELECT DISTINCT op_no FROM history WHERE ps_id = ?", (ps["ps_id"],))}
    found_next = False
    for step in steps:
        step.update(step_progress(con, ps["ps_id"], step, ps["total_qty"]))
        step["completion_override"] = step_completion_override(con, ps["ps_id"], step["step_id"])
        step["force_completed"] = step["completion_override"] == "COMPLETED"
        step["force_incomplete"] = step["completion_override"] == "INCOMPLETE"
        if step["force_completed"]:
            step["op_status"] = "COMPLETED"
            step["plan_status"] = "COMPLETED"
            step["remaining_qty"] = 0
            step["warnings"] = []
        elif step["force_incomplete"]:
            if step["plan_status"] == "COMPLETED":
                step["warnings"] = []
                step["op_status"] = "COMPLETED"
            else:
                if step["planned_qty"] >= ps["total_qty"]:
                    step["plan_status"] = "PLANNED"
                elif step["planned_qty"] > 0 or step["finished_qty"] > 0:
                    step["plan_status"] = "PARTIAL"
                else:
                    step["plan_status"] = "UNPLANNED"
                step["warnings"] = [] if step["plan_status"] == "COMPLETED" else (["UNPLANNED"] if step["plan_status"] == "UNPLANNED" else [])
                step["op_status"] = "NOT_READY" if step["plan_status"] == "UNPLANNED" else step["plan_status"]
        elif step["op_no"] in history:
            step["op_status"] = "COMPLETED"
        elif step["step_id"] in planned:
            step["op_status"] = "PLANNED"
        else:
            step["op_status"] = "NOT_READY"
            if not found_next:
                step["is_next"] = 1
                found_next = True
        step.setdefault("is_next", 0)
    return steps


def block_day_rows(rows_for_block):
    grouped = []
    for row in rows_for_block:
        same_date_as_previous = grouped and grouped[-1]["plan_date"] == row["plan_date"]
        previous_has_actual = same_date_as_previous and int(grouped[-1].get("actual_out_set") or 0) == 1
        current_has_actual = int(row.get("actual_out_set") or 0) == 1
        current_locked = int(row.get("locked") or 0) == 1
        if same_date_as_previous and not previous_has_actual and not current_has_actual:
            current = grouped[-1]
            current["end_hhmm"] = row["end_hhmm"]
            current["end_min"] = row["end_min"]
            current["qty"] += float(row.get("qty") or 0)
            current["actual_out"] += float(row.get("actual_out") or 0)
            current["actual_out_set"] = 1 if current.get("actual_out_set") or int(row.get("actual_out_set") or 0) == 1 else 0
            current["row_locked"] = 1 if int(current.get("row_locked") or 0) == 1 or current_locked else 0
            current["source_count"] += 1
            current["row_id"] = None
            current["last_row_id"] = row["row_id"]
        else:
            grouped.append({
                "plan_date": row["plan_date"],
                "start_hhmm": row["start_hhmm"],
                "end_hhmm": row["end_hhmm"],
                "start_min": row["start_min"],
                "end_min": row["end_min"],
                "qty": float(row.get("qty") or 0),
                "actual_out": float(row.get("actual_out") or 0),
                "actual_out_set": int(row.get("actual_out_set") or 0),
                "row_locked": current_locked,
                "row_id": row["row_id"],
                "last_row_id": row["row_id"],
                "source_count": 1,
            })
    return grouped


def ps_timeline_rows(con, ps_id):
    return rows(
        con.execute(
            """
            SELECT
                pr.row_id,
                pr.plan_date,
                pr.start_min,
                pr.end_min,
                pr.qty,
                pr.actual_out,
                pr.actual_out_set,
                pr.setup_mins,
                pr.locked AS row_locked,
                pr.seq_no,
                pr.op_no,
                pr.machine_id,
                pr.machine_code,
                pr.flow_step_id,
                pr.ps_id,
                pr.envelope_id,
                pr.split_group_id,
                pr.split_piece_id
            FROM planning_row pr
            WHERE pr.ps_id = ?
            ORDER BY pr.plan_date, pr.start_min, pr.end_min, pr.seq_no, pr.row_id
            """,
            (ps_id,),
        )
    )


def contiguous_envelope_for_row(con, timeline_rows, anchor_row_id):
    anchor_index = next(
        (idx for idx, row in enumerate(timeline_rows) if int(row.get("row_id") or 0) == int(anchor_row_id or 0)),
        None,
    )
    if anchor_index is None:
        return []
    anchor = timeline_rows[anchor_index]
    anchor_key = (
        int(anchor.get("seq_no") or 0),
        int(anchor.get("op_no") or 0),
        int(anchor.get("machine_id") or 0),
        compact_text(anchor.get("machine_code")),
        int(anchor.get("flow_step_id") or 0),
    )

    lane_rows = [
        row
        for row in timeline_rows
        if (
            int(row.get("seq_no") or 0) == anchor_key[0]
            and int(row.get("op_no") or 0) == anchor_key[1]
            and int(row.get("machine_id") or 0) == anchor_key[2]
            and compact_text(row.get("machine_code")) == anchor_key[3]
            and int(row.get("flow_step_id") or 0) == anchor_key[4]
        )
    ]
    lane_index = next(
        (idx for idx, row in enumerate(lane_rows) if int(row.get("row_id") or 0) == int(anchor_row_id or 0)),
        None,
    )
    if lane_index is None:
        return []

    def same_envelope(prev_row, next_row):
        return (
            int(prev_row.get("seq_no") or 0) == anchor_key[0]
            and int(next_row.get("seq_no") or 0) == anchor_key[0]
            and int(prev_row.get("op_no") or 0) == anchor_key[1]
            and int(next_row.get("op_no") or 0) == anchor_key[1]
            and int(prev_row.get("machine_id") or 0) == anchor_key[2]
            and int(next_row.get("machine_id") or 0) == anchor_key[2]
            and compact_text(prev_row.get("machine_code")) == anchor_key[3]
            and compact_text(next_row.get("machine_code")) == anchor_key[3]
            and int(prev_row.get("flow_step_id") or 0) == anchor_key[4]
            and int(next_row.get("flow_step_id") or 0) == anchor_key[4]
            and rows_share_visible_envelope(con, prev_row, next_row, ignore_envelope_id=int(anchor.get("envelope_id") or 0) or None)
        )

    start = lane_index
    while start > 0 and same_envelope(lane_rows[start - 1], lane_rows[start]):
        start -= 1
    end = lane_index
    while end + 1 < len(lane_rows) and same_envelope(lane_rows[end], lane_rows[end + 1]):
        end += 1
    return lane_rows[start:end + 1]


def rows_share_visible_envelope(con, prev_row, next_row, ignore_envelope_id=None):
    if not prev_row or not next_row:
        return False
    if int(prev_row.get("seq_no") or 0) != int(next_row.get("seq_no") or 0):
        return False
    if int(prev_row.get("machine_id") or 0) != int(next_row.get("machine_id") or 0):
        return False
    if compact_text(prev_row.get("op_no")) != compact_text(next_row.get("op_no")):
        return False
    machine = one(con.execute("SELECT shift_profile FROM machines WHERE machine_id = ?", (int(prev_row.get("machine_id") or 0),))) or {}
    include_overnight = compact_text(machine.get("shift_profile")) == "24HR"
    return not has_machine_blockers_between(
        con,
        int(prev_row.get("machine_id") or 0),
        prev_row,
        next_row,
        ignore_envelope_id=ignore_envelope_id,
        include_overnight=include_overnight,
    )


def envelope_rows_for_row_id_and_date(con, row_id, plan_date=None):
    row = one(con.execute("SELECT ps_id FROM planning_row WHERE row_id = ?", (row_id,)))
    if not row:
        return []
    timeline_rows = ps_timeline_rows(con, row["ps_id"])
    envelope = contiguous_envelope_for_row(con, timeline_rows, row_id)
    if not envelope:
        return []
    if plan_date:
        envelope = [r for r in envelope if r.get("plan_date") == plan_date]
    return envelope


def normalize_standard_rows(con, row_ids):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return 0
    ps_ids = sorted({
        int(r["ps_id"])
        for r in rows(con.execute(
            f"SELECT DISTINCT ps_id FROM planning_row WHERE row_id IN ({','.join('?' for _ in row_ids)})",
            row_ids,
        ))
        if r.get("ps_id") is not None
    })
    if len(ps_ids) != 1:
        return 0
    timeline_rows = ps_timeline_rows(con, ps_ids[0])
    row_lookup = {
        int(r["row_id"]): r
        for r in timeline_rows
        if int(r.get("row_id") or 0) in row_ids
    }
    rows_for_group = [row_lookup[rid] for rid in row_ids if rid in row_lookup]
    if not rows_for_group:
        return 0
    machine = one(con.execute("SELECT shift_profile FROM machines WHERE machine_id = ?", (int(rows_for_group[0]["machine_id"]),)))
    if machine and machine.get("shift_profile") == "24HR":
        return 0
    step = one(con.execute("SELECT cycle_time FROM part_flow_steps WHERE step_id = ?", (int(rows_for_group[0]["flow_step_id"]),)))
    cycle_time = float(step["cycle_time"] or 0) if step else 0.0

    grouped = {}
    for row in rows_for_group:
        grouped.setdefault(row["plan_date"], []).append(row)

    changed = 0
    for plan_date, day_rows in grouped.items():
        if len(day_rows) <= 1:
            continue
        day_rows.sort(key=lambda r: (r["start_min"], r["end_min"], r["row_id"]))
        first = day_rows[0]
        total_qty = sum(float(r.get("qty") or 0) for r in day_rows)
        total_setup = sum(float(r.get("setup_mins") or 0) for r in day_rows)
        total_actual = sum(float(r.get("actual_out") or 0) for r in day_rows)
        actual_set = 1 if any(int(r.get("actual_out_set") or 0) == 1 for r in day_rows) else 0
        start_min = int(first["start_min"])
        take = duration_for_units(cycle_time or 0, total_qty, total_setup)
        end_min = standard_day_finish_min(start_min, take)

        con.execute(
            """
            UPDATE planning_row
            SET start_min = ?, end_min = ?, qty = ?, actual_out = ?, actual_out_set = ?, setup_mins = ?
            WHERE row_id = ?
            """,
            (start_min, end_min, total_qty, total_actual, actual_set, total_setup, first["row_id"]),
        )
        for extra in day_rows[1:]:
            con.execute("DELETE FROM planning_row WHERE row_id = ?", (extra["row_id"],))
        changed += 1

    if changed:
        update_envelope_summary_from_rows(con, row_ids)
    return changed


def normalize_all_standard_rows(con):
    candidate_groups = []
    for ps in rows(con.execute("SELECT DISTINCT ps_id FROM planning_row ORDER BY ps_id")):
        timeline = ps_timeline_rows(con, ps["ps_id"])
        grouped = {}
        for row in timeline:
            machine = one(con.execute("SELECT shift_profile FROM machines WHERE machine_id = ?", (int(row.get("machine_id") or 0),)))
            if machine and machine.get("shift_profile") == "24HR":
                continue
            key = (
                int(row.get("seq_no") or 0),
                int(row.get("machine_id") or 0),
                compact_text(row.get("op_no")),
                row.get("plan_date"),
            )
            grouped.setdefault(key, []).append(int(row.get("row_id") or 0))
        for row_ids in grouped.values():
            row_ids = [rid for rid in row_ids if rid > 0]
            if len(row_ids) > 1:
                candidate_groups.append({"row_ids": row_ids})
    normalized = 0
    for group in candidate_groups:
        normalized += normalize_standard_rows(con, group["row_ids"])
    return {"groups_checked": len(candidate_groups), "days_normalized": normalized}


def update_envelope_summary_from_rows(con, row_ids):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return None
    q_marks = ",".join("?" for _ in row_ids)
    stats = one(
        con.execute(
            f"""
            SELECT COALESCE(SUM(qty), 0) total_qty,
                   MIN(plan_date || ' ' || printf('%02d:%02d', start_min / 60, start_min % 60)) envelope_start,
                   MAX(plan_date || ' ' || printf('%02d:%02d', end_min / 60, end_min % 60)) envelope_end
            FROM planning_row
            WHERE row_id IN ({q_marks})
            """,
            row_ids,
        )
    )
    return {
        "total_qty": stats["total_qty"] or 0,
        "envelope_start": stats["envelope_start"],
        "envelope_end": stats["envelope_end"],
    }


def sched_row_payload(row):
    return {
        "plan_date": row["plan_date"],
        "start_min": int(row["start_min"]),
        "end_min": int(row["end_min"]),
        "qty": float(row.get("qty") or 0),
        "setup_mins": float(row.get("setup_mins") or 0),
        "split_piece_id": compact_text(row.get("split_piece_id")),
        "split_group_id": compact_text(row.get("split_group_id")),
    }


def preceding_envelope_allows_setup_carryover(con, ps_id, machine_id, start_date, start_min, step=None, ignore_envelope_id=None):
    step_id = int(step["step_id"]) if step and step.get("step_id") is not None else None
    if not ps_id or not machine_id or step_id is None:
        return False
    ps = one(con.execute("SELECT part_id FROM process_sheet WHERE ps_id = ?", (ps_id,)))
    if not ps or ps.get("part_id") is None:
        return False
    target_key = datetime_key(start_date, start_min)
    timeline_rows = rows(
        con.execute(
            """
            SELECT
                pr.row_id,
                pr.plan_date,
                pr.start_min,
                pr.end_min,
                pr.machine_id,
                pr.flow_step_id,
                pr.envelope_id,
                pr.split_piece_id
            FROM planning_row pr
            JOIN process_sheet ps ON ps.ps_id = pr.ps_id
            WHERE ps.part_id = ?
              AND pr.machine_id = ?
              AND pr.flow_step_id = ?
            ORDER BY pr.plan_date, pr.start_min, pr.end_min, pr.row_id
            """,
            (int(ps["part_id"]), int(machine_id), step_id),
        )
    )
    earlier_rows = [
        row for row in timeline_rows
        if datetime_key(row.get("plan_date"), int(row.get("end_min") or 0)) <= target_key
    ]
    if not earlier_rows:
        return False
    latest = max(earlier_rows, key=lambda r: datetime_key(r.get("plan_date"), int(r.get("end_min") or 0)))
    probe_next = {
        "plan_date": start_date,
        "start_min": int(start_min),
        "end_min": int(start_min),
    }
    machine = one(con.execute("SELECT shift_profile FROM machines WHERE machine_id = ?", (int(machine_id),)))
    include_overnight = bool(machine and machine.get("shift_profile") == "24HR")
    return not has_machine_blockers_between(con, machine_id, latest, probe_next, ignore_envelope_id, include_overnight=include_overnight)


def effective_setup_time_for_new_envelope(con, ps_id, machine_id, start_date, start_min, setup_time, step=None, ignore_envelope_id=None):
    if not float(setup_time or 0):
        return 0.0
    if preceding_envelope_allows_setup_carryover(con, ps_id, machine_id, start_date, start_min, step, ignore_envelope_id=ignore_envelope_id):
        return 0.0
    return float(setup_time or 0)


def first_schedulable_start(con, machine, earliest_date, earliest_min, cycle_time, carry_mode="weekdays", reference_date=None, ignore_envelope_id=None):
    anchor_day = date.fromisoformat(earliest_date)
    day = anchor_day.isoformat() if machine["shift_profile"] == "24HR" else next_plan_day(con, anchor_day, carry_mode, reference_date or anchor_day).isoformat()
    earliest_min = max(0, int(earliest_min))
    min_duration = max(int(round(float(cycle_time or 0))), 1)

    for _ in range(180):
        if machine["shift_profile"] == "24HR":
            occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
            start = max(earliest_min if day == earliest_date else 0, 0)
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= start:
                    continue
                if occ_start > start:
                    if occ_start - start >= min_duration:
                        return day, start
                    start = occ_end
                    continue
                start = max(start, occ_end)
            if 1440 - start >= min_duration:
                return day, start
            day, earliest_min = next_calendar_day_start(day)
            continue

        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, earliest_min if day == earliest_date else win_start)
            if seg_start >= win_end:
                continue
            occupied = [(s, e) for s, e in occupied_segments(con, machine["machine_id"], day, ignore_envelope_id) if s < win_end and e > seg_start]
            for free_start, free_end in free_segments(seg_start, win_end, occupied):
                if free_end - free_start >= min_duration:
                    return day, free_start

        day, earliest_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)

    return earliest_date, earliest_min


def create_block_for_rows(con, ps_id, step, machine, rows_for_block, setup_charged=None):
    if not rows_for_block:
        return None
    first = rows_for_block[0]
    last = rows_for_block[-1]
    seq_value = step.get("seq", step.get("seq_no"))
    split_group_id = compact_text(first.get("split_group_id"))
    split_piece_id = compact_text(first.get("split_piece_id"))
    waive_setup = preceding_envelope_allows_setup_carryover(
        con, ps_id, machine["machine_id"], first["plan_date"], first["start_min"], step
    )
    effective_setup_charged = bool(setup_charged if setup_charged is not None else rows_for_block[0].get("setup_mins", 0))
    if waive_setup:
        effective_setup_charged = False
    envelope_meta = (
        ps_id,
        seq_value,
        step["op_no"],
        machine["machine_id"],
        machine["machine_code"],
        step["step_id"],
        split_piece_id,
    )
    envelope_row = one(
        con.execute(
            """
            SELECT envelope_id
            FROM planning_envelope
            WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
              AND COALESCE(split_piece_id, '') = ? AND COALESCE(archived, 0) = 0
              AND envelope_start = ? AND envelope_end = ?
            LIMIT 1
            """,
            (*envelope_meta, f"{first['plan_date']} {hhmm(first['start_min'])}", f"{last['plan_date']} {hhmm(last['end_min'])}"),
        )
    )
    if envelope_row:
        con.execute(
            """
            UPDATE planning_envelope
            SET locked = ?,
                archived = 0,
                envelope_start = ?,
                envelope_end = ?,
                split_piece_id = ?,
                updated_at = CURRENT_TIMESTAMP
            WHERE envelope_id = ?
            """,
            (
                int(effective_setup_charged),
                f"{first['plan_date']} {hhmm(first['start_min'])}",
                f"{last['plan_date']} {hhmm(last['end_min'])}",
                split_piece_id,
                int(envelope_row["envelope_id"]),
            ),
        )
    else:
        con.execute(
            """
            INSERT INTO planning_envelope (
                ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, split_piece_id,
                total_qty, locked, archived, envelope_start, envelope_end
            )
            VALUES (?, ?, ?, ?, ?, ?, ?, 0, ?, 0, ?, ?)
            """,
            (
                *envelope_meta,
                int(effective_setup_charged),
                f"{first['plan_date']} {hhmm(first['start_min'])}",
                f"{last['plan_date']} {hhmm(last['end_min'])}",
            ),
        )
        envelope_row = one(
            con.execute(
                """
                SELECT envelope_id
                FROM planning_envelope
                WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                  AND COALESCE(split_piece_id, '') = ? AND COALESCE(archived, 0) = 0
                  AND envelope_start = ? AND envelope_end = ?
                LIMIT 1
                """,
                (*envelope_meta, f"{first['plan_date']} {hhmm(first['start_min'])}", f"{last['plan_date']} {hhmm(last['end_min'])}"),
            )
        )
    envelope_id = int(envelope_row["envelope_id"]) if envelope_row else 0
    cur = con.execute(
        """
        INSERT INTO planning_block (ps_id, flow_step_id, op_no, seq_no, machine_id, machine_code,
                                    total_qty, setup_charged, envelope_start, envelope_end, status)
        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, 'PLANNED')
        """,
        (
            ps_id,
            step["step_id"],
            step["op_no"],
            seq_value,
            machine["machine_id"],
            machine["machine_code"],
            sum(row["qty"] for row in rows_for_block),
            int(effective_setup_charged),
            f"{first['plan_date']} {hhmm(first['start_min'])}",
            f"{last['plan_date']} {hhmm(last['end_min'])}",
        ),
    )
    for row in rows_for_block:
        setup_mins = 0 if waive_setup else row["setup_mins"]
        split_group_id = compact_text(row.get("split_group_id"))
        split_piece_id = compact_text(row.get("split_piece_id"))
        con.execute(
            """
            INSERT INTO planning_row (envelope_id, split_piece_id, ps_id, plan_date, start_min, end_min, qty, setup_mins, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (envelope_id, split_piece_id, ps_id, row["plan_date"], row["start_min"], row["end_min"], row["qty"], setup_mins, seq_value, step["op_no"], machine["machine_id"], machine["machine_code"], step["step_id"], split_group_id),
        )
    inserted_row_ids = [int(r["row_id"]) for r in rows(con.execute("SELECT row_id FROM planning_row WHERE envelope_id = ? ORDER BY row_id", (envelope_id,)))]
    if inserted_row_ids:
        _sync_group_locks_from_row_ids(con, inserted_row_ids)
    return envelope_id


def gap_fully_within_breaks(gap_start, gap_end):
    if gap_end <= gap_start:
        return True
    covered_until = gap_start
    for break_start, break_end in STANDARD_BREAKS:
        if break_end <= covered_until:
            continue
        if break_start > covered_until:
            return False
        covered_until = max(covered_until, break_end)
        if covered_until >= gap_end:
            return True
    return covered_until >= gap_end


def has_machine_blockers_between(con, machine_id, previous_row, next_row, ignore_envelope_id=None, include_overnight=False):
    prev_day = date.fromisoformat(previous_row["plan_date"])
    next_day = date.fromisoformat(next_row["plan_date"])
    current = prev_day
    while current <= next_day:
        day_str = current.isoformat()
        occupied = occupied_segments(con, machine_id, day_str, ignore_envelope_id)
        if current == prev_day == next_day:
            blockers = [(s, e) for s, e in occupied if s < next_row["start_min"] and e > previous_row["end_min"]]
        elif current == prev_day:
            day_end = 1440 if include_overnight else STANDARD_END
            blockers = [(s, e) for s, e in occupied if s < day_end and e > previous_row["end_min"]]
        elif current == next_day:
            day_start = 0 if include_overnight else STANDARD_START
            blockers = [(s, e) for s, e in occupied if s < next_row["start_min"] and e > day_start]
        else:
            day_start = 0 if include_overnight else STANDARD_START
            day_end = 1440 if include_overnight else STANDARD_END
            blockers = [(s, e) for s, e in occupied if s < day_end and e > day_start]
        if blockers:
            return True
        current += timedelta(days=1)
    return False


def rows_share_envelope(con, machine, previous_row, next_row, ignore_envelope_id=None):
    if not previous_row or not next_row:
        return False
    prev_day = date.fromisoformat(previous_row["plan_date"])
    next_day = date.fromisoformat(next_row["plan_date"])
    prev_piece_id = compact_text(previous_row.get("split_piece_id"))
    next_piece_id = compact_text(next_row.get("split_piece_id"))
    if prev_piece_id or next_piece_id:
        if prev_piece_id != next_piece_id:
            return False
    same_day = prev_day == next_day
    if machine["shift_profile"] == "24HR":
        if same_day and next_row["start_min"] < previous_row["end_min"]:
            return False
        return not has_machine_blockers_between(con, machine["machine_id"], previous_row, next_row, ignore_envelope_id, include_overnight=True)

    if next_day < prev_day:
        return False
    if same_day and next_row["start_min"] < previous_row["end_min"]:
        return False
    return not has_machine_blockers_between(con, machine["machine_id"], previous_row, next_row, ignore_envelope_id, include_overnight=False)


def merge_adjacent_schedule_rows(con, machine, scheduled_rows, ignore_envelope_id=None):
    if not scheduled_rows:
        return []
    merged = [dict(scheduled_rows[0])]
    for row in scheduled_rows[1:]:
        prev = merged[-1]
        contiguous = (
            prev["plan_date"] == row["plan_date"]
            and int(prev["end_min"]) == int(row["start_min"])
            and int(prev["qty"]) >= 0
            and int(row["qty"]) >= 0
        )
        pause_gap = (
            prev["plan_date"] == row["plan_date"]
            and int(row["start_min"]) > int(prev["end_min"])
            and any(
                break_start < int(row["start_min"]) and break_end > int(prev["end_min"])
                for break_start, break_end in STANDARD_BREAKS
            )
        )
        if (contiguous or pause_gap) and rows_share_envelope(con, machine, prev, row, ignore_envelope_id):
            prev["end_min"] = int(row["end_min"])
            prev["qty"] = float(prev.get("qty") or 0) + float(row.get("qty") or 0)
            prev["setup_mins"] = float(prev.get("setup_mins") or 0)
            continue
        merged.append(dict(row))
    return merged


def create_blocks_from_schedule(
    con,
    ps_id,
    step,
    machine,
    scheduled_rows,
    ignore_envelope_id=None,
    merge_same_day_same_step=True,
    merge_adjacent_rows=True,
    split_group_id=None,
):
    envelope_ids = []
    current_rows = []
    row_source = merge_adjacent_schedule_rows(con, machine, scheduled_rows, ignore_envelope_id) if merge_adjacent_rows else list(scheduled_rows)
    split_group_id = compact_text(split_group_id)
    for row in row_source:
        if current_rows and not rows_share_envelope(con, machine, current_rows[-1], row, ignore_envelope_id):
            envelope_id = create_block_for_rows(con, ps_id, step, machine, current_rows, setup_charged=current_rows[0].get("setup_mins", 0) > 0)
            if envelope_id is not None:
                envelope_ids.append(envelope_id)
            current_rows = []
        if split_group_id and not current_rows:
            row["split_group_id"] = split_group_id
        current_rows.append(row)
    if current_rows:
        envelope_id = create_block_for_rows(con, ps_id, step, machine, current_rows, setup_charged=current_rows[0].get("setup_mins", 0) > 0)
        if envelope_id is not None:
            envelope_ids.append(envelope_id)
    return envelope_ids


def blocks_for_ps(con, ps_id):
    timeline_rows = ps_timeline_rows(con, ps_id)
    envelopes = []
    lanes = {}
    lane_order = []
    current_rows = []
    current_env_id = None

    for row in timeline_rows:
        lane_key = (
            int(row.get("seq_no") or 0),
            int(row.get("op_no") or 0),
            int(row.get("machine_id") or 0),
            compact_text(row.get("machine_code")),
            int(row.get("flow_step_id") or 0),
        )
        if lane_key not in lanes:
            lanes[lane_key] = []
            lane_order.append(lane_key)
        lanes[lane_key].append(row)

    def flush_current():
        nonlocal current_rows, current_env_id
        if not current_rows:
            return
        first = current_rows[0]
        last = current_rows[-1]
        row_ids = []
        locked = 0
        row_locked_count = 0
        total_qty = 0.0
        for row in current_rows:
            if int(row.get("row_id") or 0) > 0:
                row_ids.append(int(row["row_id"]))
            locked = 1 if locked or int(row.get("locked") or 0) == 1 else 0
            row_locked_count += 1 if int(row.get("row_locked") or 0) == 1 else 0
            total_qty += float(row.get("qty") or 0)
        row_payload = [dict(row) for row in current_rows]
        envelopes.append(
            {
                "display_env_id": f"E-{ps_id}-{first['op_no']}-{first['row_id']}",
                "row_ids": row_ids,
                "locked": locked,
                "row_locked_count": row_locked_count,
                "row_locked_all": 1 if current_rows and row_locked_count == len(current_rows) else 0,
                "op_no": first["op_no"],
                "machine_code": first["machine_code"],
                "total_qty": total_qty,
                "envelope_start": f"{first['plan_date']} {hhmm(first['start_min'])}",
                "envelope_end": f"{last['plan_date']} {hhmm(last['end_min'])}",
                "rows": row_payload,
                "day_rows": block_day_rows(current_rows),
                "row_count": len(current_rows),
            }
        )
        current_rows = []
        current_env_id = None

    for lane_key in lane_order:
        current_rows = []
        current_env_id = None
        lane_rows = lanes.get(lane_key) or []
        for row in lane_rows:
            row["start_hhmm"] = display_hhmm(row["start_min"])
            row["end_hhmm"] = display_hhmm(row["end_min"])
            env_id = int(row.get("envelope_id") or 0) or int(row.get("row_id") or 0)
            if current_rows and current_env_id is not None and (
                env_id != current_env_id or not rows_share_visible_envelope(con, current_rows[-1], row, ignore_envelope_id=current_env_id)
            ):
                flush_current()
            if not current_rows:
                current_env_id = env_id
            current_rows.append(row)
        flush_current()

    return envelopes


def _backfill_split_piece_ids(con):
    if not one(con.execute("SELECT 1 FROM planning_row WHERE COALESCE(split_piece_id, '') = '' LIMIT 1")):
        return 0
    touched = 0
    ps_ids = [row["ps_id"] for row in rows(con.execute("SELECT DISTINCT ps_id FROM planning_row ORDER BY ps_id"))]
    for ps_id in ps_ids:
        for block in blocks_for_ps(con, ps_id):
            row_ids = [int(r) for r in block.get("row_ids") or [] if int(r or 0) > 0]
            if not row_ids:
                continue
            block_rows = block.get("rows") or []
            if not block_rows:
                continue
            if any(compact_text(r.get("split_piece_id")) for r in block_rows):
                continue
            piece_id = _split_piece_marker("piece")
            _apply_split_piece_ids(con, row_ids, piece_id)
            touched += 1
    return touched


def schedule_span(con, ps_id):
    timeline_rows = ps_timeline_rows(con, ps_id)
    if not timeline_rows:
        return None
    expected_start = min((f"{r['plan_date']} {hhmm(r['start_min'])}" for r in timeline_rows), default="")
    expected_end = max((f"{r['plan_date']} {hhmm(r['end_min'])}" for r in timeline_rows), default="")
    machines_used = ",".join(sorted({compact_text(r.get("machine_code")) for r in timeline_rows if compact_text(r.get("machine_code"))}))
    return {
        "expected_start": expected_start,
        "expected_end": expected_end,
        "machines_used": machines_used,
    }


def planned_support_process_sheets(con):
    ps_rows = rows(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id ORDER BY ps.due_date, ps.ps_id"))
    items = []
    for ps in ps_rows:
        item = serialize_ps(con, ps)
        if item["planner_status"] == "COMPLETED":
            continue
        if not item.get("expected_start"):
            continue
        items.append(item)
    return items


def serialize_ps(con, ps, detail=False, material_cache=None, erp_index=None, list_cache=None):
    totals = ps_totals(con, ps["ps_id"], light=not detail, erp_index=erp_index) if detail else compute_process_sheet_status_light(con, ps, erp_index=erp_index, cache=list_cache)
    span = (list_cache or {}).get("span_by_ps", {}).get(ps["ps_id"]) or schedule_span(con, ps["ps_id"]) or {}
    planned_qty = float(totals.get("planned_qty") or 0)
    total_qty = float(ps.get("total_qty") or 0)
    has_full_plan = total_qty > 0 and planned_qty >= total_qty
    item = dict(ps)
    item.update(totals)
    item["planner_status"] = totals["planner_status"]
    item["remaining_qty"] = 0 if item["planner_status"] == "COMPLETED" else max(0, item["total_qty"] - totals["finished_qty"])
    item["warnings"] = warnings_for(
        con,
        ps,
        totals,
        material_cache=material_cache,
        support_sync=detail,
        schedule=span,
    )
    completion_override = (list_cache or {}).get("override_by_ps", {}).get(ps["ps_id"])
    if detail:
        item["force_completed"] = ps_force_completed(con, ps["ps_id"])
        item["force_incomplete"] = ps_completion_override(con, ps["ps_id"]) == "INCOMPLETE"
        item["completion_override"] = ps_completion_override(con, ps["ps_id"])
    else:
        item["force_completed"] = completion_override == "COMPLETED"
        item["force_incomplete"] = completion_override == "INCOMPLETE"
        item["completion_override"] = completion_override
    item["status_reason"] = totals.get("status_reason", "")
    item["status_flags"] = totals.get("status_flags", [])
    item["status_debug"] = totals.get("status_debug", {}) if detail else {}
    item["source_ps_id"] = ps.get("source_ps_id") or split_process_sheet_key(ps["ps_id"])[0]
    item["pp_partial_no"] = ps.get("pp_partial_no") or split_process_sheet_key(ps["ps_id"])[1]
    item["expected_start"] = span.get("expected_start", "")
    item["expected_end"] = span.get("expected_end", "") if item["planner_status"] in ("PLANNED", "COMPLETED") else ""
    item["machines_used"] = span.get("machines_used", "")
    if detail:
        latest_snapshot = latest_progress_snapshot(con, ps["ps_id"])
        item["has_progress_snapshot"] = bool(latest_snapshot)
        item["latest_progress_snapshot_at"] = latest_snapshot["created_at"] if latest_snapshot else ""
        item["route_label"] = route_label(con, ps["selected_flow_id"])
        item["steps"] = step_statuses(con, ps)
        item["blocks"] = blocks_for_ps(con, ps["ps_id"])
        item["material"] = material_for_ps_cached(con, ps["ps_id"], material_cache)
        item["available_flows"] = rows(
            con.execute(
                """
                SELECT flow_id, flow_code, flow_name, is_default
                FROM part_flow_header
                WHERE part_id = ?
                ORDER BY is_default DESC, flow_id
                """,
                (ps["part_id"],),
            )
        )
        next_step = next((s for s in item["steps"] if s.get("is_next")), None)
        item["next_eligible_seq"] = next_step["seq"] if next_step else None
    return item


def next_working_day(con, start):
    current = start
    for _ in range(60):
        cal = one(con.execute("SELECT is_working_day FROM calendar_days WHERE work_date = ?", (current.isoformat(),)))
        if cal is None or cal["is_working_day"]:
            return current
        current += timedelta(days=1)
    return start


def same_week_saturday(day_value):
    offset = 5 - day_value.weekday()
    if offset < 0:
        offset += 7
    return day_value + timedelta(days=offset)


def next_plan_day(con, start, carry_mode="weekdays", reference_date=None):
    current = start
    saturday = same_week_saturday(reference_date or start) if carry_mode == "weekend" else None
    for _ in range(120):
        if carry_mode == "weekend" and saturday and current == saturday:
            return current
        cal = one(con.execute("SELECT is_working_day FROM calendar_days WHERE work_date = ?", (current.isoformat(),)))
        if cal is None or cal["is_working_day"]:
            return current
        current += timedelta(days=1)
    return start


def previous_working_day(con, start):
    current = start
    for _ in range(120):
        cal = one(con.execute("SELECT is_working_day FROM calendar_days WHERE work_date = ?", (current.isoformat(),)))
        if cal is None or cal["is_working_day"]:
            return current
        current -= timedelta(days=1)
    return start


def subtract_working_days(con, iso_date, days):
    if not iso_date:
        return ""
    current = date.fromisoformat(iso_date)
    for _ in range(max(0, int(days or 0))):
        current = previous_working_day(con, current - timedelta(days=1))
    return current.isoformat()


def next_start(con, machine_id, carry_mode="weekdays", reference_date=None):
    latest = None
    if machine_id is not None:
        all_rows = rows(
            con.execute(
                """
                SELECT row_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                       setup_mins, locked AS row_locked, seq_no, op_no, machine_id, machine_code,
                       ps_id, flow_step_id
                FROM planning_row
                WHERE machine_id = ?
                ORDER BY plan_date, start_min, end_min, seq_no, row_id
                """,
                (machine_id,),
            )
        )
        if all_rows:
            latest = max(all_rows, key=lambda r: (r["plan_date"], int(r["end_min"] or 0)))
    if latest and latest["end_min"] < STANDARD_END:
        return latest["plan_date"], max(STANDARD_START, latest["end_min"])
    day = next_plan_day(con, date.today() if not latest else date.fromisoformat(latest["plan_date"]) + timedelta(days=1), carry_mode, reference_date)
    return day.isoformat(), STANDARD_START


def occupied_segments(con, machine_id, plan_date, ignore_envelope_id=None, ignore_row_ids=None):
    sql = """
        SELECT pr.start_min, pr.end_min
        FROM planning_row pr
        WHERE pr.machine_id = ? AND pr.plan_date = ?
    """
    params = [machine_id, plan_date]
    if ignore_envelope_id:
        sql += " AND COALESCE(pr.envelope_id, 0) != ?"
        params.append(ignore_envelope_id)
    ignore_row_ids = [int(r) for r in ignore_row_ids or [] if str(r).strip().isdigit() and int(r) > 0]
    if ignore_row_ids:
        q_marks = ",".join("?" for _ in ignore_row_ids)
        sql += f" AND pr.row_id NOT IN ({q_marks})"
        params.extend(ignore_row_ids)
    return [(r["start_min"], r["end_min"]) for r in con.execute(sql, params)]


def occupied_rows(con, machine_id, plan_date, ignore_envelope_id=None, ignore_row_ids=None):
    sql = """
        SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
               seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id
        FROM planning_row
        WHERE machine_id = ? AND plan_date = ?
    """
    params = [machine_id, plan_date]
    if ignore_envelope_id:
        sql += " AND COALESCE(envelope_id, 0) != ?"
        params.append(ignore_envelope_id)
    ignore_row_ids = [int(r) for r in ignore_row_ids or [] if str(r).strip().isdigit() and int(r) > 0]
    if ignore_row_ids:
        q_marks = ",".join("?" for _ in ignore_row_ids)
        sql += f" AND row_id NOT IN ({q_marks})"
        params.extend(ignore_row_ids)
    sql += " ORDER BY start_min, end_min, row_id"
    return rows(con.execute(sql, params))


def free_segments(window_start, window_end, occupied):
    pos = window_start
    for occ_start, occ_end in sorted(occupied):
        if occ_end <= pos or occ_start >= window_end:
            continue
        if occ_start > pos:
            yield pos, min(occ_start, window_end)
        pos = max(pos, occ_end)
        if pos >= window_end:
            return
    if pos < window_end:
        yield pos, window_end


def next_day_start(con, current_date, carry_mode="weekdays", reference_date=None):
    day = next_plan_day(con, date.fromisoformat(current_date) + timedelta(days=1), carry_mode, reference_date)
    return day.isoformat(), STANDARD_START


def next_calendar_day_start(current_date):
    day = date.fromisoformat(current_date) + timedelta(days=1)
    while day.weekday() == 6:
        day += timedelta(days=1)
    return day.isoformat(), 0


def uninterrupted_capacity_before_blocker(con, machine_id, start_date, start_min, carry_mode="weekdays", reference_date=None, ignore_envelope_id=None, max_days=30):
    capacity = 0
    current_day = date.fromisoformat(start_date)

    for offset in range(max_days):
        if offset > 0:
            current_day = next_plan_day(con, current_day + timedelta(days=1), carry_mode, reference_date or date.fromisoformat(start_date))
        day_str = current_day.isoformat()
        occupied = occupied_segments(con, machine_id, day_str, ignore_envelope_id)

        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, start_min if day_str == start_date else win_start)
            if seg_start >= win_end:
                continue

            for occ_start, occ_end in sorted(occupied):
                if occ_end <= seg_start:
                    continue
                if occ_start >= win_end:
                    break
                if occ_start > seg_start:
                    capacity += occ_start - seg_start
                return capacity

            for free_start, free_end in free_segments(seg_start, win_end, occupied):
                if free_end > free_start:
                    capacity += free_end - free_start

    return capacity


def standard_day_finish_min(start_min, work_mins):
    remaining = max(0.0, float(work_mins or 0))
    current = int(start_min or 0)
    if remaining <= 0:
        return current

    for win_start, win_end in STANDARD_WINDOWS:
        if current > win_end:
            continue
        seg_start = max(current, win_start)
        if seg_start >= win_end:
            continue
        available = win_end - seg_start
        if remaining <= available + 1e-9:
            return int(round(seg_start + remaining))
        remaining -= available
        current = win_end

    return STANDARD_END


def schedule_rows(
    con,
    machine,
    earliest_date,
    earliest_min,
    duration_mins,
    qty,
    cycle_time,
    setup_time,
    carry_mode="weekdays",
    reference_date=None,
    ignore_envelope_id=None,
    prefer_even_split=False,
):
    anchor_day = date.fromisoformat(earliest_date)
    day = anchor_day.isoformat() if machine["shift_profile"] == "24HR" else next_plan_day(con, anchor_day, carry_mode, reference_date or anchor_day).isoformat()
    earliest_min = max(0, int(earliest_min))
    remaining_qty = max(0, int(round(float(qty or 0))))
    setup_remaining = float(setup_time or 0)
    result = []
    daily_capacity = 1440 if machine["shift_profile"] == "24HR" else sum(win_end - win_start for win_start, win_end in STANDARD_WINDOWS)
    total_work_mins = max(0, whole_minutes(duration_mins, minimum=0))
    estimated_days = max(1, (total_work_mins + daily_capacity - 1) // daily_capacity) if daily_capacity else 1
    even_day_target = max(1, (remaining_qty + estimated_days - 1) // estimated_days) if prefer_even_split and remaining_qty > 0 else None

    for _ in range(180):
        if remaining_qty <= 0:
            break
        if machine["shift_profile"] == "24HR":
            occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
            start = max(earliest_min if day == earliest_date else 0, 0)
            units_now = 0
            for free_start, free_end in free_segments(start, 1440, occupied):
                if free_end <= free_start:
                    continue
                available = free_end - free_start
                required_work = setup_remaining + max(float(cycle_time or 0), 1.0)
                if setup_remaining > 0 and available + 1e-9 < required_work:
                    continue
                fit = max_units_fit(available, cycle_time, remaining_qty, setup_remaining)
                if even_day_target:
                    fit = min(fit, even_day_target)
                if fit <= 0:
                    continue
                start = free_start
                units_now = fit
                break
            if units_now <= 0 or start is None:
                day, earliest_min = next_calendar_day_start(day)
                continue
        else:
            units_now = 0
            start = None
            for win_start, win_end in STANDARD_WINDOWS:
                seg_start = max(win_start, earliest_min if day == earliest_date else win_start)
                if seg_start >= win_end:
                    continue
                occupied = [(s, e) for s, e in occupied_segments(con, machine["machine_id"], day, ignore_envelope_id) if s < win_end and e > seg_start]
                for free_start, free_end in free_segments(seg_start, win_end, occupied):
                    if free_end <= free_start:
                        continue
                    required_work = setup_remaining + max(float(cycle_time or 0), 1.0)
                    if setup_remaining > 0:
                        uninterrupted_capacity = uninterrupted_capacity_before_blocker(
                            con,
                            machine["machine_id"],
                            day,
                            free_start,
                            carry_mode,
                            reference_date or anchor_day,
                            ignore_envelope_id,
                            max_days=1,
                        )
                        if uninterrupted_capacity + 1e-9 < required_work:
                            continue
                    else:
                        uninterrupted_capacity = uninterrupted_capacity_before_blocker(
                            con,
                            machine["machine_id"],
                            day,
                            free_start,
                            carry_mode,
                            reference_date or anchor_day,
                            ignore_envelope_id,
                            max_days=1,
                        )

                    fit = max_units_fit(uninterrupted_capacity, cycle_time, remaining_qty, setup_remaining)
                    if even_day_target:
                        fit = min(fit, even_day_target)
                    if fit <= 0:
                        continue
                    start = free_start
                    units_now = fit
                    break
                if units_now > 0:
                    break
            if units_now <= 0 or start is None:
                day, earliest_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)
                continue

        setup_here = setup_remaining if units_now > 0 else 0
        take = duration_for_units(cycle_time, units_now, setup_here)
        end_min = int(start + take) if machine["shift_profile"] == "24HR" else standard_day_finish_min(start, take)
        result.append({
            "plan_date": day,
            "start_min": int(start),
            "end_min": end_min,
            "qty": float(units_now),
            "setup_mins": setup_here,
        })
        setup_remaining = 0
        remaining_qty -= units_now
        earliest_date = day
        earliest_min = end_min
        if earliest_min >= 1440:
            if machine["shift_profile"] == "24HR":
                day, earliest_min = next_calendar_day_start(day)
            else:
                day, earliest_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)

    if remaining_qty > 0:
        raise ValueError("Could not find enough machine capacity in the scheduling window")
    return result


def first_schedulable_start_for_duration(con, machine, earliest_date, earliest_min, required_mins, carry_mode="weekdays", reference_date=None, ignore_envelope_id=None):
    anchor_day = date.fromisoformat(earliest_date)
    day = anchor_day.isoformat() if machine["shift_profile"] == "24HR" else next_plan_day(con, anchor_day, carry_mode, reference_date or anchor_day).isoformat()
    earliest_min = max(0, int(earliest_min))
    required_mins = max(1, int(round(float(required_mins or 0))))

    for _ in range(180):
        if machine["shift_profile"] == "24HR":
            occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
            start = max(earliest_min if day == earliest_date else 0, 0)
            for free_start, free_end in free_segments(start, 1440, occupied):
                if free_end - free_start >= required_mins:
                    return day, free_start
            day, earliest_min = next_calendar_day_start(day)
            continue

        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, earliest_min if day == earliest_date else win_start)
            if seg_start >= win_end:
                continue
            occupied = [(s, e) for s, e in occupied_segments(con, machine["machine_id"], day, ignore_envelope_id) if s < win_end and e > seg_start]
            for free_start, free_end in free_segments(seg_start, win_end, occupied):
                if free_end - free_start >= required_mins:
                    return day, free_start
        day, earliest_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)

    return earliest_date, earliest_min


def first_unblocked_start_for_move(con, machine, earliest_date, earliest_min, carry_mode="weekdays", reference_date=None, ignore_envelope_id=None, ignore_row_ids=None):
    anchor_day = date.fromisoformat(earliest_date)
    day = anchor_day.isoformat() if machine["shift_profile"] == "24HR" else next_plan_day(con, anchor_day, carry_mode, reference_date or anchor_day).isoformat()
    earliest_min = max(0, int(earliest_min))

    for _ in range(180):
        if machine["shift_profile"] == "24HR":
            occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id, ignore_row_ids)
            start = max(earliest_min if day == earliest_date else 0, 0)
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= start:
                    continue
                if occ_start <= start < occ_end:
                    start = occ_end
                    break
            else:
                return day, start
            if start >= 1440:
                day, earliest_min = next_calendar_day_start(day)
            else:
                earliest_min = start
            continue

        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, earliest_min if day == earliest_date else win_start)
            if seg_start >= win_end:
                continue
            occupied = [(s, e) for s, e in occupied_segments(con, machine["machine_id"], day, ignore_envelope_id, ignore_row_ids) if s < win_end and e > seg_start]
            blocked = False
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= seg_start:
                    continue
                if occ_start <= seg_start < occ_end:
                    seg_start = occ_end
                    blocked = True
                    break
            if not blocked:
                return day, seg_start
            if seg_start >= win_end:
                continue
        day, earliest_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)

    return earliest_date, earliest_min


def _datetime_from_plan_min(plan_date, minute_value):
    return datetime.combine(date.fromisoformat(plan_date), datetime.min.time()) + timedelta(minutes=int(minute_value or 0))


def _shift_rows_to_anchor(rows_in, target_date, target_start_min):
    ordered = sorted(
        [dict(row) for row in rows_in],
        key=lambda r: (r["plan_date"], int(r["start_min"] or 0), int(r["end_min"] or 0), int(r.get("row_id") or 0)),
    )
    if not ordered:
        return []
    origin_start = _datetime_from_plan_min(ordered[0]["plan_date"], ordered[0]["start_min"])
    target_start = _datetime_from_plan_min(target_date, target_start_min)
    delta = target_start - origin_start
    shifted = []
    for row in ordered:
        start_dt = _datetime_from_plan_min(row["plan_date"], row["start_min"]) + delta
        end_dt = _datetime_from_plan_min(row["plan_date"], row["end_min"]) + delta
        shifted_row = dict(row)
        shifted_row["plan_date"] = start_dt.date().isoformat()
        shifted_row["start_min"] = start_dt.hour * 60 + start_dt.minute
        shifted_row["end_min"] = int(round((end_dt - datetime.combine(start_dt.date(), datetime.min.time())).total_seconds() / 60))
        shifted.append(shifted_row)
    return shifted


def _shifted_rows_blockers(con, machine_id, shifted_rows, ignore_envelope_id=None):
    blockers = []
    seen = set()
    for row in shifted_rows:
        row_start = int(row.get("start_min") or 0)
        row_end = int(row.get("end_min") or 0)
        for blocker in occupied_rows(con, machine_id, row["plan_date"], ignore_envelope_id):
            blocker_id = int(blocker["row_id"] or 0)
            if blocker_id in seen:
                continue
            blocker_start = int(blocker.get("start_min") or 0)
            blocker_end = int(blocker.get("end_min") or 0)
            if row_end <= blocker_start or row_start >= blocker_end:
                continue
            seen.add(blocker_id)
            blockers.append(dict(blocker))
    return blockers


def _planned_move_rows(con, machine, source_rows, target_date, target_start, move_qty, cycle_time, setup_time, ignore_envelope_id=None):
    planned_rows = schedule_rows(
        con,
        machine,
        target_date,
        target_start,
        whole_minutes(setup_time + move_qty * cycle_time),
        int(round(move_qty)),
        cycle_time,
        setup_time,
        ignore_envelope_id=ignore_envelope_id,
    )
    if planned_rows:
        planned_rows[0]["setup_mins"] = setup_time
        for row in planned_rows[1:]:
            row["setup_mins"] = 0
    return planned_rows


def _partition_source_rows_for_move(source_rows, move_qty, cycle_time):
    remaining_to_move = float(move_qty)
    kept_source_rows = []
    moved_source_rows = []
    for row in reversed([dict(r) for r in source_rows]):
        if remaining_to_move <= 1e-9:
            kept_source_rows.append(row)
            continue

        row_qty = float(row.get("qty") or 0)
        if row_qty <= remaining_to_move + 1e-9:
            moved_source_rows.append(row)
            remaining_to_move -= row_qty
            continue

        remove_duration = duration_for_units(cycle_time or 1, remaining_to_move, 0)
        kept_row = dict(row)
        moved_row = dict(row)
        kept_row["qty"] = row_qty - remaining_to_move
        kept_row["end_min"] = max(int(row["start_min"]), int(row["end_min"]) - int(remove_duration))
        moved_row["qty"] = float(remaining_to_move)
        moved_row["start_min"] = kept_row["end_min"]
        moved_row["end_min"] = int(row["end_min"])
        kept_source_rows.append(kept_row)
        moved_source_rows.append(moved_row)
        remaining_to_move = 0

    if remaining_to_move > 1e-9:
        raise ValueError("Could not isolate the requested quantity from the source row")

    kept_source_rows.reverse()
    moved_source_rows.reverse()
    return kept_source_rows, moved_source_rows


def schedule_rows_from_tail(
    con,
    machine,
    earliest_date,
    earliest_min,
    duration_mins,
    qty,
    cycle_time,
    setup_time,
    carry_mode="weekdays",
    reference_date=None,
    ignore_envelope_id=None,
):
    anchor_day = date.fromisoformat(earliest_date)
    day = anchor_day.isoformat() if machine["shift_profile"] == "24HR" else next_plan_day(con, anchor_day, carry_mode, reference_date or anchor_day).isoformat()
    current_min = max(0, int(earliest_min))
    remaining_qty = max(0, int(round(float(qty or 0))))
    setup_remaining = float(setup_time or 0)
    result = []
    for _ in range(180):
        if remaining_qty <= 0:
            break
        occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
        if machine["shift_profile"] == "24HR":
            start = current_min if day == earliest_date else 0
            units_now = 0
            for free_start, free_end in free_segments(start, 1440, occupied):
                available = free_end - free_start
                if available <= 0:
                    continue
                required_work = setup_remaining + max(float(cycle_time or 0), 1.0)
                if remaining_qty > 0 and setup_remaining > 0 and available + 1e-9 < required_work:
                    continue
                fit = max_units_fit(available, cycle_time, remaining_qty, setup_remaining)
                if fit <= 0:
                    continue
                start = free_start
                units_now = fit
                break
            if units_now <= 0:
                day, current_min = next_calendar_day_start(day)
                continue
        else:
            units_now = 0
            start = None
            for win_start, win_end in STANDARD_WINDOWS:
                seg_start = max(win_start, current_min if day == earliest_date else win_start)
                if seg_start >= win_end:
                    continue
                blocker_start = win_end
                for occ_start, occ_end in sorted(occupied):
                    if occ_end <= seg_start:
                        continue
                    if occ_start >= win_end:
                        break
                    blocker_start = max(seg_start, min(blocker_start, int(occ_start)))
                    break
                uninterrupted_capacity = blocker_start - seg_start
                if uninterrupted_capacity <= 0:
                    continue
                required_work = setup_remaining + max(float(cycle_time or 0), 1.0)
                if remaining_qty > 0 and setup_remaining > 0 and uninterrupted_capacity + 1e-9 < required_work:
                    continue
                fit = max_units_fit(uninterrupted_capacity, cycle_time, remaining_qty, setup_remaining)
                if fit <= 0:
                    continue
                start = seg_start
                units_now = fit
                break
            if units_now <= 0 or start is None:
                day, current_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)
                continue

        setup_here = setup_remaining if units_now > 0 else 0
        take = duration_for_units(cycle_time, units_now, setup_here)
        end_min = int(start + take) if machine["shift_profile"] == "24HR" else standard_day_finish_min(start, take)
        result.append({
            "plan_date": day,
            "start_min": int(start),
            "end_min": end_min,
            "qty": float(units_now),
            "setup_mins": setup_here,
        })
        setup_remaining = 0
        remaining_qty -= units_now
        current_min = end_min
        if current_min >= 1440:
            if machine["shift_profile"] == "24HR":
                day, current_min = next_calendar_day_start(day)
            else:
                day, current_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)
        elif machine["shift_profile"] != "24HR":
            # Keep actual-output refill on the current contiguous tail only.
            # If the tail cannot continue within the same stretch, move to the next day
            # instead of searching later free windows on the same day.
            next_blocker = None
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= current_min:
                    continue
                next_blocker = (int(occ_start), int(occ_end))
                break
            if next_blocker and current_min >= next_blocker[0]:
                day, current_min = next_day_start(con, day, carry_mode, reference_date or anchor_day)

    if remaining_qty > 0:
        raise ValueError("Could not find enough machine capacity in the scheduling window")
    return result


def existing_step_span(con, ps_id, step_id):
    return one(
        con.execute(
            """
            SELECT pr.flow_step_id,
                   MIN(pr.plan_date) start_date,
                   MIN(pr.plan_date || printf('%04d', pr.start_min)) start_key,
                   MAX(pr.plan_date) end_date,
                   MAX(pr.plan_date || printf('%04d', pr.end_min)) end_key
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.flow_step_id = ?
            GROUP BY pr.flow_step_id
            ORDER BY end_key DESC
            LIMIT 1
            """,
            (ps_id, step_id),
        )
    )


def split_datetime_key(value):
    return value[:10], int(value[10:12]) * 60 + int(value[12:14])


def datetime_key(date_str, minute):
    return f"{date_str}{int(minute):04d}"


def max_datetime_pair(left_date, left_min, right_date, right_min):
    if datetime_key(left_date, left_min) >= datetime_key(right_date, right_min):
        return left_date, int(left_min)
    return right_date, int(right_min)


def step_actual_qty(con, ps_id, seq_no):
    active = one(
        con.execute(
            """
            SELECT COALESCE(SUM(pr.actual_out), 0) actual_qty
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
            """,
            (ps_id, seq_no),
        )
    )["actual_qty"]
    hist = one(
        con.execute(
            """
            SELECT COALESCE(SUM(hr.actual_out), 0) actual_qty
            FROM history_row hr
            WHERE hr.ps_id = ? AND hr.seq_no = ?
            """,
            (ps_id, seq_no),
        )
    )["actual_qty"]
    return (active or 0) + (hist or 0)


def existing_seq_span(con, ps_id, seq_no):
    return one(
        con.execute(
            """
            SELECT MIN(pr.plan_date || printf('%04d', pr.start_min)) start_key,
                   MAX(pr.plan_date || printf('%04d', pr.end_min)) end_key
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
            """,
            (ps_id, seq_no),
        )
    )


def future_started_rows_exist(con, ps_id, from_seq, current_row_id=None, split_piece_id=None):
    params = [ps_id, from_seq]
    sql = """
        SELECT COUNT(*) started_count FROM (
            SELECT pr.row_id
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no >= ? AND pr.actual_out > 0
    """
    if current_row_id is not None:
        sql += " AND pr.row_id != ?"
        params.append(current_row_id)
    if split_piece_id is not None:
        sql += " AND COALESCE(pr.split_piece_id, '') = ?"
        params.append(compact_text(split_piece_id))
    sql += """
            UNION ALL
            SELECT hr.hist_row_id AS row_id
            FROM history_row hr
            WHERE hr.ps_id = ? AND hr.seq_no >= ? AND hr.actual_out > 0
        )
    """
    params.extend([ps_id, from_seq])
    return con.execute(sql, params).fetchone()[0] > 0


def seq_has_earlier_rows_before(con, ps_id, seq_no, row_key):
    row = one(
        con.execute(
            """
            SELECT 1
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) < ?
            LIMIT 1
            """,
            (ps_id, seq_no, row_key),
        )
    )
    return bool(row)


def first_planned_row_after_seq(con, ps_id, seq_no):
    return one(
        con.execute(
            """
            SELECT pr.seq_no, pr.plan_date, pr.start_min, pr.row_id
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no > ?
            ORDER BY pr.seq_no ASC, pr.plan_date ASC, pr.start_min ASC, pr.row_id ASC
            LIMIT 1
            """,
            (ps_id, seq_no),
        )
    )


def reapply_actual_to_seq_head(con, ps_id, seq_no, actual_value):
    row = one(
        con.execute(
            """
            SELECT pr.row_id
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
            ORDER BY pr.plan_date, pr.start_min, pr.row_id
            LIMIT 1
            """,
            (ps_id, seq_no),
        )
    )
    if row:
        con.execute(
            "UPDATE planning_row SET actual_out = ?, actual_out_set = 1 WHERE row_id = ?",
            (float(actual_value or 0), int(row["row_id"])),
        )


def replan_downstream(con, ps_id, from_seq, earliest_date, earliest_min, carry_mode="weekdays", split_piece_id=None):
    ps = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
    steps = flow_steps(con, ps["selected_flow_id"])
    scheduled = []
    skipped = []
    piece_id = compact_text(split_piece_id) if split_piece_id is not None else ""
    piece_filter_sql = " AND COALESCE(split_piece_id, '') = ?" if piece_id else ""
    planned_seq_query = """
        SELECT DISTINCT seq_no
        FROM planning_row
        WHERE ps_id = ? AND seq_no >= ? AND COALESCE(split_group_id, '') = ''
    """
    planned_seq_params = [ps_id, from_seq]
    if piece_id:
        planned_seq_query += " AND COALESCE(split_piece_id, '') = ?"
        planned_seq_params.append(piece_id)
    planned_seq_query += "\n        ORDER BY seq_no"
    planned_seq_rows = rows(con.execute(planned_seq_query, planned_seq_params))
    planned_seq_set = {int(r["seq_no"]) for r in planned_seq_rows if r.get("seq_no") is not None}
    target_steps = [s for s in steps if s["seq"] >= from_seq and int(s["seq"]) in planned_seq_set]

    def remaining_qty_for_step(step):
        actual = step_actual_qty(con, ps_id, step["seq"])
        return max(0, ps["total_qty"] - actual)

    def collect_shared_chain(start_idx):
        return [target_steps[start_idx]], start_idx

    def available_minutes_in_day(machine, day, start_min, ignore_envelope_id=None):
        start_min = max(0, int(start_min or 0))
        occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
        if machine["shift_profile"] == "24HR":
            pos = start_min
            free = 0
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= pos:
                    continue
                if occ_start > pos:
                    free += occ_start - pos
                pos = max(pos, occ_end)
                if pos >= 1440:
                    return free
            return free + max(0, 1440 - pos)

        total = 0
        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, start_min)
            if seg_start >= win_end:
                continue
            seg_end = win_end
            pos = seg_start
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= pos or occ_start >= seg_end:
                    continue
                if occ_start > pos:
                    total += occ_start - pos
                pos = max(pos, occ_end)
                if pos >= seg_end:
                    break
            if pos < seg_end:
                total += seg_end - pos
        return total

    def plan_shared_chain(chain, chain_machine, start_date, start_min):
        chain_remaining = {
            chain_step["step_id"]: remaining_qty_for_step(chain_step)
            for chain_step in chain
            if remaining_qty_for_step(chain_step) > 0
        }
        if not chain_remaining:
            return [], [], start_date, start_min

        chain_scheduled = []
        chain_skipped = []
        chain_date = start_date
        chain_min = start_min

        while any(qty > 0 for qty in chain_remaining.values()):
            day_free = available_minutes_in_day(chain_machine, chain_date, chain_min)
            if day_free <= 0:
                if chain_machine["shift_profile"] == "24HR":
                    chain_date, chain_min = next_calendar_day_start(chain_date)
                else:
                    chain_date, chain_min = next_day_start(con, chain_date, carry_mode, date.fromisoformat(start_date))
                continue

            active_steps = [s for s in chain if chain_remaining.get(s["step_id"], 0) > 0]
            if not active_steps:
                break

            chain_cycle = sum(float(s["cycle_time"] or 1) for s in active_steps)
            if chain_cycle <= 0:
                break

            chain_setup = effective_setup_time_for_new_envelope(
                con,
                ps_id,
                chain_machine["machine_id"],
                chain_date,
                chain_min,
                chain[0]["setup_time"] or 0,
                chain[0],
            )

            rounds = max_units_fit(day_free, chain_cycle, min(chain_remaining.values()), chain_setup)
            if rounds <= 0:
                progress = False
                for chain_step in active_steps:
                    step_qty = chain_remaining[chain_step["step_id"]]
                    step_cycle = float(chain_step["cycle_time"] or 1)
                    setup_here = chain_setup
                    if available_minutes_in_day(chain_machine, chain_date, chain_min) >= setup_here + step_cycle:
                        qty_alloc = min(step_qty, max_units_fit(day_free, step_cycle, step_qty, setup_here))
                        if qty_alloc > 0:
                            actual_start_date, actual_start_min = first_schedulable_start(
                                con, chain_machine, chain_date, chain_min, step_cycle, carry_mode, date.fromisoformat(chain_date)
                            )
                            duration = whole_minutes(setup_here + qty_alloc * step_cycle)
                            scheduled = schedule_rows(
                                con,
                                chain_machine,
                                actual_start_date,
                                actual_start_min,
                                duration,
                                qty_alloc,
                                step_cycle,
                                setup_here,
                                carry_mode,
                                date.fromisoformat(chain_date),
                                prefer_even_split=False,
                            )
                            envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                            chain_date = scheduled[-1]["plan_date"]
                            chain_min = scheduled[-1]["end_min"]
                            chain_remaining[chain_step["step_id"]] -= qty_alloc
                            for envelope_id, row in zip(envelope_ids, scheduled):
                                chain_scheduled.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})
                            chain_scheduled.append({"op_no": chain_step["op_no"]})
                            progress = True
                            break
                if not progress:
                    if chain_machine["shift_profile"] == "24HR":
                        chain_date, chain_min = next_calendar_day_start(chain_date)
                    else:
                        chain_date, chain_min = next_day_start(con, chain_date, carry_mode, date.fromisoformat(start_date))
                continue

            chain_round_qty = max(1, rounds)
            for chain_step in active_steps:
                qty_alloc = min(chain_remaining[chain_step["step_id"]], chain_round_qty)
                if qty_alloc <= 0:
                    continue
                actual_start_date, actual_start_min = first_schedulable_start(
                    con, chain_machine, chain_date, chain_min, chain_step["cycle_time"] or 1, carry_mode, date.fromisoformat(chain_date)
                )
                setup_here = chain_setup if chain_step == active_steps[0] else 0
                duration = whole_minutes(setup_here + qty_alloc * (chain_step["cycle_time"] or 1))
                scheduled = schedule_rows(
                    con,
                    chain_machine,
                    actual_start_date,
                    actual_start_min,
                    duration,
                    qty_alloc,
                    chain_step["cycle_time"] or 1,
                    setup_here,
                    carry_mode,
                    date.fromisoformat(chain_date),
                    prefer_even_split=False,
                )
                envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                chain_date = scheduled[-1]["plan_date"]
                chain_min = scheduled[-1]["end_min"]
                chain_remaining[chain_step["step_id"]] -= qty_alloc
                for envelope_id, row in zip(envelope_ids, scheduled):
                    chain_scheduled.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})
                chain_scheduled.append({"op_no": chain_step["op_no"]})

            while available_minutes_in_day(chain_machine, chain_date, chain_min) > 0 and any(qty > 0 for qty in chain_remaining.values()):
                progressed = False
                for chain_step in active_steps:
                    remaining = chain_remaining[chain_step["step_id"]]
                    if remaining <= 0:
                        continue
                    step_cycle = float(chain_step["cycle_time"] or 1)
                    day_free_left = available_minutes_in_day(chain_machine, chain_date, chain_min)
                    if day_free_left < step_cycle:
                        continue
                    qty_alloc = max_units_fit(day_free_left, step_cycle, remaining, 0)
                    if qty_alloc <= 0:
                        continue
                    actual_start_date, actual_start_min = first_schedulable_start(
                        con, chain_machine, chain_date, chain_min, step_cycle, carry_mode, date.fromisoformat(chain_date)
                    )
                    duration = whole_minutes(qty_alloc * step_cycle)
                    scheduled = schedule_rows(
                        con,
                        chain_machine,
                        actual_start_date,
                        actual_start_min,
                        duration,
                        qty_alloc,
                        step_cycle,
                        0,
                        carry_mode,
                        date.fromisoformat(chain_date),
                        prefer_even_split=False,
                    )
                    envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                    chain_date = scheduled[-1]["plan_date"]
                    chain_min = scheduled[-1]["end_min"]
                    chain_remaining[chain_step["step_id"]] -= qty_alloc
                    for envelope_id, row in zip(envelope_ids, scheduled):
                        chain_scheduled.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})
                    chain_scheduled.append({"op_no": chain_step["op_no"]})
                    progressed = True
                    break
                if not progressed:
                    break

            if chain_machine["shift_profile"] == "24HR":
                chain_date, chain_min = next_calendar_day_start(chain_date)
            else:
                chain_date, chain_min = next_day_start(con, chain_date, carry_mode, date.fromisoformat(chain_date))

        return chain_scheduled, chain_skipped, chain_date, chain_min

    for step in target_steps:
        block_piece_sql = " AND COALESCE(split_piece_id, '') = ?" if piece_id else ""
        locked = con.execute(
            "SELECT COUNT(*) FROM planning_block WHERE ps_id = ? AND seq_no = ? AND archived = 0 AND locked = 1" + block_piece_sql,
            ([ps_id, step["seq"], piece_id] if piece_id else [ps_id, step["seq"]]),
        ).fetchone()[0]
        actual = step_actual_qty(con, ps_id, step["seq"])
        if locked and actual > 0:
            skipped.append({
                "op_no": step["op_no"],
                "op_type": step["op_type"],
                "machine_category": step["machine_category"],
                "preferred_machine": step["preferred_machine"],
                "reason": "Already started and locked; left in place",
            })
            continue
        if actual > 0:
            skipped.append({
                "op_no": step["op_no"],
                "op_type": step["op_type"],
                "machine_category": step["machine_category"],
                "preferred_machine": step["preferred_machine"],
                "reason": "Already started; left in place",
            })
            continue

    # Clear the unlocked future chain first, so earlier replanned operations do not
    # get artificially split around stale downstream envelopes on the same machines.
    if planned_seq_set:
        q_marks = ",".join("?" for _ in planned_seq_set)
        delete_block_query = f"""
            DELETE FROM planning_block
            WHERE ps_id = ? AND archived = 0 AND locked = 0 AND seq_no IN ({q_marks})
        """
        delete_block_params = [ps_id, *sorted(planned_seq_set)]
        if piece_id:
            delete_block_query += " AND COALESCE(split_piece_id, '') = ?"
            delete_block_params.append(piece_id)
        con.execute(delete_block_query, delete_block_params)

    idx = 0
    while idx < len(target_steps):
        step = target_steps[idx]
        chain, chain_end = collect_shared_chain(idx)
        if len(chain) > 1:
            machine = pick_machine_for_step(con, chain[0], require_preferred_machine=True)
            if not machine:
                skipped.append({
                    "op_no": chain[0]["op_no"],
                    "op_type": chain[0]["op_type"],
                    "machine_category": chain[0]["machine_category"],
                    "preferred_machine": chain[0]["preferred_machine"],
                    "reason": "No preferred machine assigned",
                })
                break
            chain_scheduled, chain_skipped, earliest_date, earliest_min = plan_shared_chain(chain, machine, earliest_date, earliest_min)
            scheduled.extend([row["op_no"] for row in chain_scheduled if row.get("op_no")])
            skipped.extend(chain_skipped)
            idx = chain_end + 1
            continue

        actual = step_actual_qty(con, ps_id, step["seq"])
        remaining_qty = max(0, ps["total_qty"] - actual)
        if remaining_qty > 0:
            machine = pick_machine_for_step(con, step, require_preferred_machine=True)
            if not machine:
                skipped.append({
                    "op_no": step["op_no"],
                    "op_type": step["op_type"],
                    "machine_category": step["machine_category"],
                    "preferred_machine": step["preferred_machine"],
                    "reason": "No preferred machine assigned",
                })
                break
            actual_start_date, actual_start_min = first_schedulable_start(
                con, machine, earliest_date, earliest_min, step["cycle_time"] or 1, carry_mode, date.fromisoformat(earliest_date)
            )
            effective_setup = effective_setup_time_for_new_envelope(
                con, ps_id, machine["machine_id"], actual_start_date, actual_start_min, step["setup_time"] or 0, step
            )
            duration = whole_minutes(effective_setup + remaining_qty * (step["cycle_time"] or 1))
            scheduled_rows = schedule_rows(
                con,
                machine,
                actual_start_date,
                actual_start_min,
                duration,
                remaining_qty,
                step["cycle_time"] or 1,
                effective_setup,
                carry_mode,
                date.fromisoformat(earliest_date),
                prefer_even_split=False,
            )
            create_blocks_from_schedule(con, ps_id, step, machine, scheduled_rows)
            earliest_date = scheduled_rows[-1]["plan_date"]
            earliest_min = scheduled_rows[-1]["end_min"]
            scheduled.append(step["op_no"])
        idx += 1
    return {"scheduled_ops": scheduled, "skipped_ops": skipped}


def replan_downstream_with_floor(con, ps_id, from_seq, earliest_date, earliest_min, carry_mode="weekdays", split_piece_id=None):
    ps = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
    steps = flow_steps(con, ps["selected_flow_id"])
    piece_id = compact_text(split_piece_id) if split_piece_id is not None else ""
    planned_seq_query = """
        SELECT DISTINCT seq_no
        FROM planning_row
        WHERE ps_id = ? AND seq_no >= ? AND COALESCE(split_group_id, '') = ''
    """
    planned_seq_params = [ps_id, from_seq]
    if piece_id:
        planned_seq_query += " AND COALESCE(split_piece_id, '') = ?"
        planned_seq_params.append(piece_id)
    planned_seq_query += "\n        ORDER BY seq_no"
    planned_seq_rows = rows(con.execute(planned_seq_query, planned_seq_params))
    planned_seq_set = {int(r["seq_no"]) for r in planned_seq_rows if r.get("seq_no") is not None}
    target_steps = [s for s in steps if s["seq"] >= from_seq and int(s["seq"]) in planned_seq_set]
    existing_starts = {}

    for step in target_steps:
        actual = step_actual_qty(con, ps_id, step["seq"])
        if actual > 0:
            skipped.append({
                "op_no": step["op_no"],
                "op_type": step["op_type"],
                "machine_category": step["machine_category"],
                "preferred_machine": step["preferred_machine"],
                "reason": "Already started; left in place",
            })
            continue
        span = existing_seq_span(con, ps_id, step["seq"])
        if span and span.get("start_key"):
            existing_starts[step["seq"]] = split_datetime_key(span["start_key"])

    if planned_seq_set:
        q_marks = ",".join("?" for _ in planned_seq_set)
        delete_block_query = f"""
            DELETE FROM planning_block
            WHERE ps_id = ? AND archived = 0 AND seq_no IN ({q_marks})
        """
        delete_block_params = [ps_id, *sorted(planned_seq_set)]
        if piece_id:
            delete_block_query += " AND COALESCE(split_piece_id, '') = ?"
            delete_block_params.append(piece_id)
        con.execute(delete_block_query, delete_block_params)

    scheduled = []
    skipped = []
    floor_date, floor_min = earliest_date, int(earliest_min)

    for step in target_steps:
        actual = step_actual_qty(con, ps_id, step["seq"])
        remaining_qty = max(0, ps["total_qty"] - actual)
        if remaining_qty <= 0:
            continue
        machine = pick_machine_for_step(con, step, require_preferred_machine=True)
        if not machine:
            skipped.append({
                "op_no": step["op_no"],
                "op_type": step["op_type"],
                "machine_category": step["machine_category"],
                "preferred_machine": step["preferred_machine"],
                "reason": "No preferred machine assigned",
            })
            break

        existing_start = existing_starts.get(step["seq"])
        if existing_start:
            request_date, request_min = max_datetime_pair(floor_date, floor_min, existing_start[0], existing_start[1])
        else:
            request_date, request_min = floor_date, floor_min

        actual_start_date, actual_start_min = first_schedulable_start(
            con, machine, request_date, request_min, step["cycle_time"] or 1, carry_mode, date.fromisoformat(request_date)
        )
        effective_setup = effective_setup_time_for_new_envelope(
            con, ps_id, machine["machine_id"], actual_start_date, actual_start_min, step["setup_time"] or 0, step
        )
        duration = whole_minutes(effective_setup + remaining_qty * (step["cycle_time"] or 1))
        scheduled_rows = schedule_rows(
            con,
            machine,
            request_date,
            request_min,
            duration,
            remaining_qty,
            step["cycle_time"] or 1,
            effective_setup,
            carry_mode,
            date.fromisoformat(request_date),
            prefer_even_split=False,
        )
        create_blocks_from_schedule(con, ps_id, step, machine, scheduled_rows)
        floor_date = scheduled_rows[-1]["plan_date"]
        floor_min = scheduled_rows[-1]["end_min"]
        scheduled.append(step["op_no"])

    return {"scheduled_ops": scheduled, "skipped_ops": skipped}


def planning_anchor_before_seq(con, ps_id, seq_no):
    active = one(
        con.execute(
            """
            SELECT pr.plan_date, pr.end_min
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no < ?
            ORDER BY pr.seq_no DESC, pr.plan_date DESC, pr.end_min DESC, pr.row_id DESC
            LIMIT 1
            """,
            (ps_id, seq_no),
        )
    )
    if active:
        return active["plan_date"], int(active["end_min"])

    hist = one(
        con.execute(
            """
            SELECT hr.plan_date, hr.end_min
            FROM history_row hr
            WHERE hr.ps_id = ? AND hr.seq_no < ?
            ORDER BY hr.seq_no DESC, hr.plan_date DESC, hr.end_min DESC, hr.hist_row_id DESC
            LIMIT 1
            """,
            (ps_id, seq_no),
        )
    )
    if hist:
        return hist["plan_date"], int(hist["end_min"])

    return date.today().isoformat(), STANDARD_START


def apply_flow_timing_updates(con, flow_id, from_seq=None):
    affected = 0
    skipped = []
    ps_rows = rows(con.execute("SELECT ps_id FROM process_sheet WHERE selected_flow_id = ? ORDER BY ps_id", (flow_id,)))

    for ps_row in ps_rows:
        ps_id = ps_row["ps_id"]
        try:
            ps = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
            if not ps:
                continue

            start_seq = int(from_seq or 1)
            started_anchor = one(
                con.execute(
                    """
                    SELECT pr.seq_no, pr.row_id
                    FROM planning_row pr
                    WHERE pr.ps_id = ? AND pr.seq_no >= ? AND pr.actual_out > 0
                    ORDER BY pr.seq_no ASC, pr.plan_date DESC, pr.start_min DESC, pr.row_id DESC
                    LIMIT 1
                    """,
                    (ps_id, start_seq),
                )
            )
            if started_anchor:
                rebalance_after_row_actual(con, started_anchor["row_id"])
                affected += 1
                continue

            earliest_planned = one(
                con.execute(
                    """
                    SELECT MIN(seq_no) AS seq_no
                    FROM planning_row
                    WHERE ps_id = ? AND seq_no >= ?
                    """,
                    (ps_id, start_seq),
                )
            )
            if not earliest_planned or earliest_planned.get("seq_no") is None:
                continue

            from_seq = int(earliest_planned["seq_no"])
            locked_future = con.execute(
                """
                SELECT COUNT(*)
                FROM planning_row
                WHERE ps_id = ? AND seq_no >= ? AND locked = 1
                """,
                (ps_id, from_seq),
            ).fetchone()[0]
            if locked_future:
                skipped.append(ps_id)
                continue

            anchor_date, anchor_min = planning_anchor_before_seq(con, ps_id, from_seq)
            replan_downstream(con, ps_id, from_seq, anchor_date, anchor_min)
            affected += 1
        except Exception as exc:
            skipped.append(f"{ps_id}: {exc}")

    return {"affected_ps": affected, "skipped_ps": sorted(set(skipped))}


def rebalance_after_row_actual(con, row_id, finishup_mode="next_free"):
    target = one(
        con.execute(
            """
            SELECT pr.ps_id, pr.seq_no, pr.machine_id, pr.flow_step_id, pr.envelope_id, pr.locked, pr.split_piece_id,
                   pr.row_id, pr.plan_date, pr.start_min, pr.end_min, pr.qty, pr.actual_out,
                   pr.op_no, fs.cycle_time, fs.setup_time,
                   ps.total_qty, ps.selected_flow_id
            FROM planning_row pr
            JOIN part_flow_steps fs ON fs.step_id = pr.flow_step_id
            JOIN process_sheet ps ON ps.ps_id = pr.ps_id
            WHERE pr.row_id = ?
            """,
            (row_id,),
        )
    )
    if not target:
        raise ValueError("Planning row not found")
    if target["locked"]:
        raise ValueError("Locked blocks cannot be auto-adjusted")
    row_key = f"{target['plan_date']}{int(target['start_min']):04d}{int(target['row_id']):010d}"
    actual_done = one(
        con.execute(
            """
            SELECT COALESCE(SUM(actual_out), 0) actual_qty
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) <= ?
            """,
            (target["ps_id"], target["seq_no"], row_key),
        )
    )["actual_qty"]
    hist_actual = one(
        con.execute(
            """
            SELECT COALESCE(SUM(hr.actual_out), 0) actual_qty
            FROM history_row hr
            WHERE hr.ps_id = ? AND hr.seq_no = ?
            """,
            (target["ps_id"], target["seq_no"]),
        )
    )["actual_qty"]
    remaining_qty = max(0, target["total_qty"] - ((actual_done or 0) + (hist_actual or 0)))

    machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (target["machine_id"],)))
    future_rows = rows(
        con.execute(
            """
            SELECT pr.row_id
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) > ?
            ORDER BY pr.plan_date, pr.start_min, pr.row_id
            """,
            (target["ps_id"], target["seq_no"], row_key),
        )
    )
    if future_rows:
        con.execute(
            """
            DELETE FROM planning_row
            WHERE row_id IN (
                SELECT pr.row_id
                FROM planning_row pr
                WHERE pr.ps_id = ? AND pr.seq_no = ?
                  AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) > ?
            )
            """,
            (target["ps_id"], target["seq_no"], row_key),
        )

    group_envelope_id = int(target["envelope_id"] or 0)
    if remaining_qty > 0:
        duration = whole_minutes(remaining_qty * (target["cycle_time"] or 1))
        current_rows = rows(con.execute("SELECT qty, plan_date, start_min, end_min FROM planning_row WHERE envelope_id = ? ORDER BY plan_date, start_min, row_id", (group_envelope_id,)))
        last_row = next((row for row in reversed(current_rows) if float(row.get("qty") or 0) > 1e-9), current_rows[-1] if current_rows else None)
        if last_row:
            refill_date, refill_min = last_row["plan_date"], int(last_row["end_min"])
        else:
            refill_date, refill_min = target["plan_date"], int(target["end_min"])
        if finishup_mode == "immediate":
            refill_date, refill_min = refill_date, refill_min
        elif machine["shift_profile"] == "24HR":
            refill_date, refill_min = refill_date, refill_min
        tail_piece_id = compact_text(last_row.get("split_piece_id")) if last_row else ""
        if not tail_piece_id:
            tail_piece_id = compact_text(target.get("split_piece_id"))
        refill_rows = schedule_rows_from_tail(
            con,
            machine,
            refill_date,
            refill_min,
            duration,
            remaining_qty,
            target["cycle_time"] or 1,
            0,
            "weekdays",
            date.fromisoformat(target["plan_date"]),
            ignore_envelope_id=group_envelope_id,
        )
        for row in refill_rows:
                con.execute(
                    """
                    INSERT INTO planning_row (envelope_id, split_piece_id, ps_id, plan_date, start_min, end_min, qty, setup_mins, seq_no, op_no, machine_id, machine_code, flow_step_id)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (group_envelope_id, tail_piece_id, target["ps_id"], row["plan_date"], row["start_min"], row["end_min"], row["qty"], row["setup_mins"], target["seq_no"], target["op_no"], target["machine_id"], target["machine_code"], target["flow_step_id"]),
                )
        current_rows = rows(con.execute("SELECT qty, plan_date, end_min FROM planning_row WHERE envelope_id = ? ORDER BY plan_date, start_min, row_id", (group_envelope_id,)))
        end_date = current_rows[-1]["plan_date"]
        end_min = current_rows[-1]["end_min"]
    else:
        current_rows = rows(con.execute("SELECT qty, plan_date, end_min FROM planning_row WHERE envelope_id = ? ORDER BY plan_date, start_min, row_id", (group_envelope_id,)))
        end_date = current_rows[-1]["plan_date"]
        end_min = current_rows[-1]["end_min"]

    con.execute(
        """
        UPDATE planning_envelope
        SET total_qty = ?,
            envelope_start = (SELECT MIN(plan_date || ' ' || printf('%02d:%02d', start_min / 60, start_min % 60)) FROM planning_row WHERE envelope_id = ?),
            envelope_end = (SELECT MAX(plan_date || ' ' || printf('%02d:%02d', end_min / 60, end_min % 60)) FROM planning_row WHERE envelope_id = ?)
        WHERE envelope_id = ?
        """,
        (sum(row["qty"] for row in current_rows), group_envelope_id, group_envelope_id, group_envelope_id),
    )

    return {
        "rebalanced_step": target["op_no"],
        "end_date": end_date,
        "end_min": int(end_min),
        "shifted_ops": {"scheduled_ops": [], "skipped_ops": []},
    }


def rebalance_after_row_overrun_tail(con, row_id, delta_qty, finishup_mode="next_free"):
    target = one(
        con.execute(
            """
            SELECT pr.ps_id, pr.seq_no, pr.machine_id, pr.machine_code, pr.flow_step_id, pr.envelope_id, pr.locked,
                   pr.row_id, pr.plan_date, pr.start_min, pr.end_min, pr.qty, pr.actual_out,
                   pr.op_no, fs.cycle_time, fs.setup_time,
                   ps.total_qty, ps.selected_flow_id
            FROM planning_row pr
            JOIN part_flow_steps fs ON fs.step_id = pr.flow_step_id
            JOIN process_sheet ps ON ps.ps_id = pr.ps_id
            WHERE pr.row_id = ?
            """,
            (row_id,),
        )
    )
    if not target:
        raise ValueError("Planning row not found")
    envelope_id = int(target["envelope_id"] or 0)

    row_key = f"{target['plan_date']}{int(target['start_min']):04d}{int(target['row_id']):010d}"
    future_rows = rows(
        con.execute(
            """
            SELECT pr.row_id, pr.envelope_id, pr.split_piece_id, pr.plan_date, pr.start_min, pr.end_min, pr.qty
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) > ?
            ORDER BY pr.plan_date, pr.start_min, pr.row_id
            """,
            (target["ps_id"], target["seq_no"], row_key),
        )
    )
    qty_delta = float(delta_qty or 0)
    if qty_delta > 1e-9:
        remaining_trim = qty_delta
        for row in reversed(future_rows):
            if remaining_trim <= 1e-9:
                break
            row_qty = float(row.get("qty") or 0)
            if row_qty <= remaining_trim + 1e-9:
                con.execute("DELETE FROM planning_row WHERE row_id = ?", (row["row_id"],))
                remaining_trim = max(0.0, remaining_trim - row_qty)
                continue

            remove_duration = duration_for_units(target["cycle_time"] or 1, remaining_trim, 0)
            new_end = max(int(row["start_min"]), int(row["end_min"]) - remove_duration)
            new_qty = max(0.0, row_qty - remaining_trim)
            con.execute(
                "UPDATE planning_row SET qty = ?, end_min = ? WHERE row_id = ?",
                (new_qty, new_end, row["row_id"]),
            )
            remaining_trim = 0.0
            break
    elif qty_delta < -1e-9:
        extra_qty = abs(qty_delta)
        tail_row = None
        if future_rows:
            tail_row = future_rows[-1]
        else:
            tail_row = {
                "envelope_id": envelope_id,
                "split_piece_id": compact_text(target.get("split_piece_id")),
                "plan_date": target["plan_date"],
                "start_min": int(target["start_min"] or 0),
                "end_min": int(target["end_min"] or 0),
            }
        machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (target["machine_id"],)))
        tail_is_edited_day = compact_text(tail_row.get("plan_date")) == compact_text(target["plan_date"])
        if tail_is_edited_day:
            if machine["shift_profile"] == "24HR":
                refill_date, refill_min = next_calendar_day_start(target["plan_date"])
            else:
                refill_date, refill_min = next_day_start(con, target["plan_date"], finishup_mode, date.fromisoformat(target["plan_date"]))
        else:
            refill_date, refill_min = tail_row["plan_date"], int(tail_row["end_min"] or 0)
        duration = whole_minutes(extra_qty * (target["cycle_time"] or 1))
        tail_piece_id = compact_text(tail_row.get("split_piece_id")) or compact_text(target.get("split_piece_id"))
        tail_envelope_id = envelope_id
        refill_rows = schedule_rows_from_tail(
            con,
            machine,
            refill_date,
            refill_min,
            duration,
            extra_qty,
            target["cycle_time"] or 1,
            0,
            "weekdays",
            date.fromisoformat(refill_date),
            ignore_envelope_id=tail_envelope_id,
        )
        for row in refill_rows:
            con.execute(
                """
                INSERT INTO planning_row (envelope_id, split_piece_id, ps_id, plan_date, start_min, end_min, qty, setup_mins, seq_no, op_no, machine_id, machine_code, flow_step_id)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (tail_envelope_id, tail_piece_id, target["ps_id"], row["plan_date"], row["start_min"], row["end_min"], row["qty"], row["setup_mins"], target["seq_no"], target["op_no"], target["machine_id"], target["machine_code"], target["flow_step_id"]),
            )
        tail_row_ids = [
            int(r["row_id"])
            for r in rows(
                con.execute(
                    "SELECT row_id FROM planning_row WHERE envelope_id = ? ORDER BY plan_date, start_min, row_id",
                    (tail_envelope_id,),
                )
            )
        ]
        if tail_row_ids:
            merge_contiguous_rows_for_rows(con, tail_row_ids)

    current_rows = rows(
        con.execute(
            "SELECT qty, plan_date, end_min, envelope_id FROM planning_row WHERE ps_id = ? AND seq_no = ? ORDER BY plan_date, start_min, row_id",
            (target["ps_id"], target["seq_no"]),
        )
    )
    end_date = current_rows[-1]["plan_date"] if current_rows else target["plan_date"]
    end_min = current_rows[-1]["end_min"] if current_rows else int(target["end_min"])

    refresh_envelope_ids = sorted({
        int(r["envelope_id"])
        for r in current_rows
        if r.get("envelope_id") is not None
    })
    for refresh_envelope_id in refresh_envelope_ids:
        envelope_rows = rows(
            con.execute(
                "SELECT qty, plan_date, start_min, end_min FROM planning_row WHERE envelope_id = ? ORDER BY plan_date, start_min, row_id",
                (refresh_envelope_id,),
            )
        )
        if not envelope_rows:
            continue
        con.execute(
            """
            UPDATE planning_envelope
            SET total_qty = ?,
                envelope_start = (SELECT MIN(plan_date || ' ' || printf('%02d:%02d', start_min / 60, start_min % 60)) FROM planning_row WHERE envelope_id = ?),
                envelope_end = (SELECT MAX(plan_date || ' ' || printf('%02d:%02d', end_min / 60, end_min % 60)) FROM planning_row WHERE envelope_id = ?)
            WHERE envelope_id = ?
            """,
            (sum(row["qty"] for row in envelope_rows), refresh_envelope_id, refresh_envelope_id, refresh_envelope_id),
        )

    return {
        "rebalanced_step": target["op_no"],
        "end_date": end_date,
        "end_min": int(end_min),
        "shifted_ops": {"scheduled_ops": [], "skipped_ops": []},
        "delta_qty": qty_delta,
    }


def normalize_sequence_after_row(con, ps_id, seq_no, anchor_row_id):
    anchor = one(
        con.execute(
            """
            SELECT pr.plan_date, pr.start_min, pr.end_min
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ? AND pr.row_id = ?
            """,
            (ps_id, seq_no, anchor_row_id),
        )
    )
    if not anchor:
        return None

    def abs_min(plan_date, minute):
        return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

    def row_span_minutes(start_min, end_min):
        start_i = int(start_min or 0)
        end_i = int(end_min or 0)
        span = end_i - start_i
        if span >= 0:
            return span
        return end_i + 1440 - start_i

    anchor_key = f"{anchor['plan_date']}{int(anchor['start_min'] or 0):04d}{int(anchor_row_id):010d}"
    future_rows = rows(
        con.execute(
            """
            SELECT pr.row_id, pr.plan_date, pr.start_min, pr.end_min
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) > ?
            ORDER BY pr.plan_date, pr.start_min, pr.row_id
            """,
            (ps_id, seq_no, anchor_key),
        )
    )
    prev_end_abs = abs_min(anchor["plan_date"], anchor["end_min"])
    for row in future_rows:
        start_abs = abs_min(row["plan_date"], row["start_min"])
        duration = row_span_minutes(row["start_min"], row["end_min"])
        new_start_abs = max(start_abs, prev_end_abs)
        new_end_abs = new_start_abs + duration
        if new_start_abs != start_abs or new_end_abs != abs_min(row["plan_date"], row["end_min"]):
            new_plan_date = date.fromordinal(new_start_abs // 1440).isoformat()
            new_start_min = int(new_start_abs % 1440)
            new_end_min = int(new_end_abs % 1440)
            con.execute(
                "UPDATE planning_row SET plan_date = ?, start_min = ?, end_min = ? WHERE row_id = ?",
                (new_plan_date, new_start_min, new_end_min, row["row_id"]),
            )
        prev_end_abs = new_end_abs
    touched_envelope_ids = sorted({
        int(r["envelope_id"])
        for r in rows(
            con.execute(
                """
                SELECT DISTINCT envelope_id
                FROM planning_row
                WHERE ps_id = ? AND seq_no = ? AND envelope_id IS NOT NULL
                """,
                (ps_id, seq_no),
            )
        )
        if r.get("envelope_id") is not None
    })
    for envelope_id in touched_envelope_ids:
        current_row_ids = [int(r["row_id"]) for r in rows(con.execute(
            "SELECT row_id FROM planning_row WHERE envelope_id = ? ORDER BY row_id",
            (envelope_id,),
        ))]
        envelope_summary = update_envelope_summary_from_rows(con, current_row_ids)
        if envelope_summary:
            con.execute(
                """
                UPDATE planning_envelope
                SET total_qty = ?, envelope_start = ?, envelope_end = ?, updated_at = CURRENT_TIMESTAMP
                WHERE envelope_id = ?
                """,
                (envelope_summary["total_qty"], envelope_summary["envelope_start"], envelope_summary["envelope_end"], envelope_id),
            )

    tail = one(
        con.execute(
            """
            SELECT pr.plan_date, pr.end_min
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
            ORDER BY pr.plan_date DESC, pr.end_min DESC, pr.row_id DESC
            LIMIT 1
            """,
            (ps_id, seq_no),
        )
    )
    if not tail:
        return None
    return {"end_date": tail["plan_date"], "end_min": int(tail["end_min"] or 0)}


def merge_contiguous_rows_for_rows(con, row_ids):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return False
    q_marks = ",".join("?" for _ in row_ids)
    rows_in_block = rows(
        con.execute(
            f"""
            SELECT row_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set, setup_mins
            FROM planning_row
            WHERE row_id IN ({q_marks})
            ORDER BY plan_date, start_min, row_id
            """,
            row_ids,
        )
    )
    if not rows_in_block:
        return False

    changed = False
    idx = 1

    def abs_min(plan_date, minute):
        return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

    def rows_are_mergeable(prev, curr):
        if prev["plan_date"] != curr["plan_date"]:
            return False
        prev_end = int(prev["end_min"] or 0)
        curr_start = int(curr["start_min"] or 0)
        if prev_end == curr_start:
            return True
        if curr_start <= prev_end:
            return False
        return any(
            break_start < curr_start and break_end > prev_end
            for break_start, break_end in STANDARD_BREAKS
        )

    while idx < len(rows_in_block):
        prev = rows_in_block[idx - 1]
        curr = rows_in_block[idx]
        directly_contiguous = prev["plan_date"] == curr["plan_date"] and int(prev["end_min"] or 0) == int(curr["start_min"] or 0)
        pause_gap = rows_are_mergeable(prev, curr)
        if not (directly_contiguous or pause_gap):
            idx += 1
            continue

        merged_qty = float(prev.get("qty") or 0) + float(curr.get("qty") or 0)
        merged_actual = float(prev.get("actual_out") or 0) + float(curr.get("actual_out") or 0)
        merged_actual_set = 1 if int(prev.get("actual_out_set") or 0) == 1 or int(curr.get("actual_out_set") or 0) == 1 else 0
        merged_setup = max(float(prev.get("setup_mins") or 0), float(curr.get("setup_mins") or 0))
        con.execute(
            """
            UPDATE planning_row
            SET end_min = ?, qty = ?, actual_out = ?, actual_out_set = ?, setup_mins = ?
            WHERE row_id = ?
            """,
            (
                int(curr["end_min"] or 0),
                merged_qty,
                merged_actual,
                merged_actual_set,
                merged_setup,
                int(prev["row_id"]),
            ),
        )
        con.execute("DELETE FROM planning_row WHERE row_id = ?", (int(curr["row_id"]),))
        changed = True
        rows_in_block[idx - 1] = {
            **prev,
            "end_min": int(curr["end_min"] or 0),
            "qty": merged_qty,
            "actual_out": merged_actual,
            "actual_out_set": merged_actual_set,
            "setup_mins": merged_setup,
        }
        del rows_in_block[idx]
    if changed:
        envelope_summary = update_envelope_summary_from_rows(con, row_ids)
        if envelope_summary:
            touched_envelopes = sorted({
                int(r["envelope_id"])
                for r in rows(con.execute(
                    f"SELECT DISTINCT envelope_id FROM planning_row WHERE row_id IN ({q_marks})",
                    row_ids,
                ))
                if r.get("envelope_id") is not None
            })
            for envelope_id in touched_envelopes:
                con.execute(
                    """
                    UPDATE planning_envelope
                    SET total_qty = ?, envelope_start = ?, envelope_end = ?
                    WHERE envelope_id = ?
                    """,
                    (envelope_summary["total_qty"], envelope_summary["envelope_start"], envelope_summary["envelope_end"], envelope_id),
                )
    return changed


def validate_sequence_after_row(con, ps_id, seq_no, anchor_row_id):
    anchor = one(
        con.execute(
            """
            SELECT pr.plan_date, pr.start_min, pr.end_min, pr.split_piece_id
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ? AND pr.row_id = ?
            """,
            (ps_id, seq_no, anchor_row_id),
        )
    )
    if not anchor:
        return []

    def abs_min(plan_date, minute):
        return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

    def row_span_minutes(start_min, end_min):
        start_i = int(start_min or 0)
        end_i = int(end_min or 0)
        span = end_i - start_i
        if span >= 0:
            return span
        return end_i + 1440 - start_i

    anchor_key = f"{anchor['plan_date']}{int(anchor['start_min'] or 0):04d}{int(anchor_row_id):010d}"
    anchor_piece_id = compact_text(anchor.get("split_piece_id"))
    future_rows = rows(
        con.execute(
            """
            SELECT pr.row_id, pr.plan_date, pr.start_min, pr.end_min, pr.qty
            FROM planning_row pr
            WHERE pr.ps_id = ? AND pr.seq_no = ?
              AND (pr.plan_date || printf('%04d', pr.start_min) || printf('%010d', pr.row_id)) > ?
              AND COALESCE(pr.split_piece_id, '') = ?
            ORDER BY pr.plan_date, pr.start_min, pr.row_id
            """,
            (ps_id, seq_no, anchor_key, anchor_piece_id),
        )
    )
    issues = []
    prev_end_abs = abs_min(anchor["plan_date"], anchor["end_min"])
    for row in future_rows:
        start_abs = abs_min(row["plan_date"], row["start_min"])
        duration = row_span_minutes(row["start_min"], row["end_min"])
        end_abs = start_abs + duration
        if start_abs < prev_end_abs - 1e-9:
            issues.append(
                {
                    "row_id": int(row["row_id"]),
                    "plan_date": row["plan_date"],
                    "start_min": int(row["start_min"] or 0),
                    "end_min": int(row["end_min"] or 0),
                    "expected_start_min": int(prev_end_abs % 1440),
                    "expected_start_date": date.fromordinal(prev_end_abs // 1440).isoformat(),
                }
            )
        prev_end_abs = max(prev_end_abs, end_abs)
    return issues


def auto_plan(con, ps_id):
    ps = one(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id WHERE ps.ps_id = ?", (ps_id,)))
    if not ps:
        raise ValueError("Process sheet not found")
    completion_override = ps_completion_override(con, ps_id)
    current_status = compute_process_sheet_status(con, ps)["planner_status"]
    if current_status == "COMPLETED" and completion_override != "INCOMPLETE":
        raise ValueError("Process sheet is already fully completed")
    con.execute(
        """
        DELETE FROM planning_envelope
        WHERE ps_id = ? AND locked = 0 AND archived = 0
          AND envelope_id NOT IN (
            SELECT DISTINCT envelope_id FROM planning_row WHERE ps_id = ? AND actual_out > 0
          )
        """,
        (ps_id, ps_id),
    )

    steps = flow_steps(con, ps["selected_flow_id"])
    if not steps:
        raise ValueError("Selected flow has no operations")

    totals = ps_totals(con, ps_id)
    qty_to_plan = max(0, ps["total_qty"] - totals["finished_qty"])
    has_step_overrides = any(step_completion_override(con, ps_id, step["step_id"]) is not None for step in steps)
    all_steps_force_completed = bool(steps) and all(step_force_completed(con, ps_id, step["step_id"]) for step in steps)
    if qty_to_plan <= 0 and has_step_overrides and not all_steps_force_completed:
        qty_to_plan = max(0, int(round(float(ps["total_qty"] or 0))))
    if qty_to_plan <= 0:
        raise ValueError("Process sheet is already fully completed")

    earliest_date = date.today().isoformat()
    earliest_min = STANDARD_START
    planned = []
    skipped = []
    step_metrics = {row["step_id"]: row for row in process_sheet_step_metrics(con, ps_id, steps)["steps"]}

    def step_remaining_qty(step):
        metric = step_metrics.get(step["step_id"], {})
        planned_qty = float(metric.get("planned_qty") or 0)
        actual_qty = float(metric.get("actual_qty") or 0)
        committed_qty = max(planned_qty, actual_qty)
        return max(0, float(ps["total_qty"] or 0) - committed_qty)

    def collect_shared_chain(start_idx):
        return [steps[start_idx]], start_idx

    def weighted_integer_split(total, items):
        total = max(0, int(total))
        if total <= 0 or not items:
            return [0 for _ in items]
        weights = [max(0.0, float(weight)) for _, weight in items]
        weight_total = sum(weights) or 1.0
        base = []
        fractions = []
        for idx, (_, weight) in enumerate(items):
            exact = total * (weight / weight_total)
            whole = int(exact)
            base.append(min(int(items[idx][0]), whole))
            fractions.append((exact - whole, idx))
        assigned = sum(base)
        remainder = total - assigned
        for _, idx in sorted(fractions, key=lambda item: (-item[0], item[1])):
            if remainder <= 0:
                break
            ceiling = int(items[idx][0])
            if base[idx] >= ceiling:
                continue
            base[idx] += 1
            remainder -= 1
        if remainder > 0:
            for idx in range(len(base)):
                while remainder > 0 and base[idx] < int(items[idx][0]):
                    base[idx] += 1
                    remainder -= 1
        return base

    def available_minutes_in_day(machine, day, start_min, ignore_envelope_id=None):
        start_min = max(0, int(start_min or 0))
        occupied = occupied_segments(con, machine["machine_id"], day, ignore_envelope_id)
        if machine["shift_profile"] == "24HR":
            pos = start_min
            free = 0
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= pos:
                    continue
                if occ_start > pos:
                    free += occ_start - pos
                pos = max(pos, occ_end)
                if pos >= 1440:
                    return free
            return free + max(0, 1440 - pos)

        total = 0
        for win_start, win_end in STANDARD_WINDOWS:
            seg_start = max(win_start, start_min)
            if seg_start >= win_end:
                continue
            seg_end = win_end
            pos = seg_start
            for occ_start, occ_end in sorted(occupied):
                if occ_end <= pos or occ_start >= seg_end:
                    continue
                if occ_start > pos:
                    total += occ_start - pos
                pos = max(pos, occ_end)
                if pos >= seg_end:
                    break
            if pos < seg_end:
                total += seg_end - pos
        return total

    idx = 0
    while idx < len(steps):
        step = steps[idx]
        chain, chain_end = collect_shared_chain(idx)
        if len(chain) > 1:
            chain_machine = pick_machine_for_step(con, chain[0], require_preferred_machine=True)
            if not chain_machine:
                skipped.append({
                    "op_no": chain[0]["op_no"],
                    "op_type": chain[0]["op_type"],
                    "machine_category": chain[0]["machine_category"],
                    "preferred_machine": chain[0]["preferred_machine"],
                    "reason": "No preferred machine assigned",
                })
                break

            chain_remaining = {
                chain_step["step_id"]: step_remaining_qty(chain_step)
                for chain_step in chain
                if step_remaining_qty(chain_step) > 0
            }
            if not chain_remaining:
                idx = chain_end + 1
                continue

            chain_date = earliest_date
            chain_min = earliest_min
            while any(qty > 0 for qty in chain_remaining.values()):
                day_free = available_minutes_in_day(chain_machine, chain_date, chain_min)
                if day_free <= 0:
                    if chain_machine["shift_profile"] == "24HR":
                        chain_date, chain_min = next_calendar_day_start(chain_date)
                    else:
                        chain_date, chain_min = next_day_start(con, chain_date, "weekdays", date.fromisoformat(earliest_date))
                    continue

                active_steps = [s for s in chain if chain_remaining.get(s["step_id"], 0) > 0]
                if not active_steps:
                    break

                chain_cycle = sum(float(s["cycle_time"] or 1) for s in active_steps)
                if chain_cycle <= 0:
                    break

                chain_setup = effective_setup_time_for_new_envelope(
                    con,
                    ps_id,
                    chain_machine["machine_id"],
                    chain_date,
                    chain_min,
                    chain[0]["setup_time"] or 0,
                    chain[0],
                )

                rounds = max_units_fit(day_free, chain_cycle, min(chain_remaining.values()), chain_setup)
                if rounds <= 0:
                    # Try to fit a single unit of the earliest remaining step before advancing the day.
                    progress = False
                    for chain_step in active_steps:
                        step_qty = chain_remaining[chain_step["step_id"]]
                        step_cycle = float(chain_step["cycle_time"] or 1)
                        setup_here = chain_setup
                        if available_minutes_in_day(chain_machine, chain_date, chain_min) >= setup_here + step_cycle:
                            qty_alloc = min(step_qty, max_units_fit(day_free, step_cycle, step_qty, setup_here))
                            if qty_alloc > 0:
                                actual_start_date, actual_start_min = first_schedulable_start(
                                    con, chain_machine, chain_date, chain_min, step_cycle
                                )
                                duration = whole_minutes(setup_here + qty_alloc * step_cycle)
                                scheduled = schedule_rows(
                                    con,
                                    chain_machine,
                                    actual_start_date,
                                    actual_start_min,
                                    duration,
                                    qty_alloc,
                                    step_cycle,
                                    setup_here,
                                    prefer_even_split=False,
                                )
                                envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                                chain_date = scheduled[-1]["plan_date"]
                                chain_min = scheduled[-1]["end_min"]
                                chain_remaining[chain_step["step_id"]] -= qty_alloc
                                for envelope_id, row in zip(envelope_ids, scheduled):
                                    planned.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})
                                progress = True
                                break
                    if not progress:
                        if chain_machine["shift_profile"] == "24HR":
                            chain_date, chain_min = next_calendar_day_start(chain_date)
                        else:
                            chain_date, chain_min = next_day_start(con, chain_date, "weekdays", date.fromisoformat(earliest_date))
                    continue

                chain_round_qty = max(1, rounds)
                for chain_step in active_steps:
                    qty_alloc = min(chain_remaining[chain_step["step_id"]], chain_round_qty)
                    if qty_alloc <= 0:
                        continue
                    actual_start_date, actual_start_min = first_schedulable_start(
                        con, chain_machine, chain_date, chain_min, chain_step["cycle_time"] or 1
                    )
                    setup_here = chain_setup if chain_step == active_steps[0] else 0
                    duration = whole_minutes(setup_here + qty_alloc * (chain_step["cycle_time"] or 1))
                    scheduled = schedule_rows(
                        con,
                        chain_machine,
                        actual_start_date,
                        actual_start_min,
                        duration,
                        qty_alloc,
                        chain_step["cycle_time"] or 1,
                        setup_here,
                        prefer_even_split=False,
                    )
                    envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                    chain_date = scheduled[-1]["plan_date"]
                    chain_min = scheduled[-1]["end_min"]
                    chain_remaining[chain_step["step_id"]] -= qty_alloc
                    for envelope_id, row in zip(envelope_ids, scheduled):
                        planned.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})

                # If there is still room after the main round, try to consume leftovers
                # with the earliest remaining steps in order.
                while available_minutes_in_day(chain_machine, chain_date, chain_min) > 0 and any(qty > 0 for qty in chain_remaining.values()):
                    progressed = False
                    for chain_step in active_steps:
                        remaining = chain_remaining[chain_step["step_id"]]
                        if remaining <= 0:
                            continue
                        step_cycle = float(chain_step["cycle_time"] or 1)
                        day_free_left = available_minutes_in_day(chain_machine, chain_date, chain_min)
                        if day_free_left < step_cycle:
                            continue
                        qty_alloc = max_units_fit(day_free_left, step_cycle, remaining, 0)
                        if qty_alloc <= 0:
                            continue
                        actual_start_date, actual_start_min = first_schedulable_start(
                            con, chain_machine, chain_date, chain_min, step_cycle
                        )
                        duration = whole_minutes(qty_alloc * step_cycle)
                        scheduled = schedule_rows(
                            con,
                            chain_machine,
                            actual_start_date,
                            actual_start_min,
                            duration,
                            qty_alloc,
                            step_cycle,
                            0,
                            prefer_even_split=False,
                        )
                        envelope_ids = create_blocks_from_schedule(con, ps_id, chain_step, chain_machine, scheduled)
                        chain_date = scheduled[-1]["plan_date"]
                        chain_min = scheduled[-1]["end_min"]
                        chain_remaining[chain_step["step_id"]] -= qty_alloc
                        for envelope_id, row in zip(envelope_ids, scheduled):
                            planned.append({"envelope_id": envelope_id, "op_no": chain_step["op_no"], "qty_planned": row["qty"], "existing": False})
                        progressed = True
                        break
                    if not progressed:
                        break

                if chain_machine["shift_profile"] == "24HR":
                    chain_date, chain_min = next_calendar_day_start(chain_date)
                else:
                    chain_date, chain_min = next_day_start(con, chain_date, "weekdays", date.fromisoformat(chain_date))

            earliest_date, earliest_min = chain_date, chain_min
            idx = chain_end + 1
            continue

        existing = existing_step_span(con, ps_id, step["step_id"])
        metric = step_metrics.get(step["step_id"], {})
        step_planned_qty = float(metric.get("planned_qty") or 0)
        step_actual = float(metric.get("actual_qty") or 0)
        step_committed_qty = max(step_planned_qty, step_actual)
        step_remaining_qty = max(0, float(ps["total_qty"] or 0) - step_committed_qty)
        step_completed = step_force_completed(con, ps_id, step["step_id"]) or step_actual >= ps["total_qty"]
        if step_completed:
            if existing and existing["end_key"]:
                earliest_date, earliest_min = split_datetime_key(existing["end_key"])
            idx += 1
            continue
        if existing and existing["end_key"]:
            earliest_date, earliest_min = split_datetime_key(existing["end_key"])
            if step_remaining_qty <= 0:
                idx += 1
                continue
            existing_machine = one(con.execute("SELECT machine_id FROM planning_row WHERE ps_id = ? AND seq_no = ? LIMIT 1", (ps_id, step["seq"])))
            if not existing_machine:
                idx += 1
                continue
            machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (existing_machine["machine_id"],)))
            if not machine:
                skipped.append({
                    "op_no": step["op_no"],
                    "op_type": step["op_type"],
                    "machine_category": step["machine_category"],
                    "preferred_machine": step["preferred_machine"],
                    "reason": "Existing envelope machine could not be resolved",
                })
                break
            actual_start_date, actual_start_min = first_schedulable_start(
                con, machine, earliest_date, earliest_min, step["cycle_time"] or 1
            )
            duration = whole_minutes(step_remaining_qty * (step["cycle_time"] or 1))
            scheduled = schedule_rows(
                con,
                machine,
                actual_start_date,
                actual_start_min,
                duration,
                step_remaining_qty,
                step["cycle_time"] or 1,
                0,
                prefer_even_split=False,
            )
            envelope_ids = create_blocks_from_schedule(con, ps_id, step, machine, scheduled)
            earliest_date = scheduled[-1]["plan_date"]
            earliest_min = scheduled[-1]["end_min"]
            for envelope_id, row in zip(envelope_ids, scheduled):
                planned.append({"envelope_id": envelope_id, "op_no": step["op_no"], "qty_planned": row["qty"], "existing": False})
            idx += 1
            continue
        machine = pick_machine_for_step(con, step, require_preferred_machine=True)
        if not machine:
            skipped.append({
                "op_no": step["op_no"],
                "op_type": step["op_type"],
                "machine_category": step["machine_category"],
                "preferred_machine": step["preferred_machine"],
                "reason": "No preferred machine assigned",
            })
            break

        actual_start_date, actual_start_min = first_schedulable_start(
            con, machine, earliest_date, earliest_min, step["cycle_time"] or 1
        )
        effective_setup = effective_setup_time_for_new_envelope(
            con, ps_id, machine["machine_id"], actual_start_date, actual_start_min, step["setup_time"] or 0, step
        )
        duration = whole_minutes(effective_setup + step_remaining_qty * (step["cycle_time"] or 1))
        scheduled = schedule_rows(
            con,
            machine,
            earliest_date,
            earliest_min,
            duration,
            step_remaining_qty,
            step["cycle_time"] or 1,
            effective_setup,
            prefer_even_split=False,
        )
        envelope_ids = create_blocks_from_schedule(con, ps_id, step, machine, scheduled)
        earliest_date = scheduled[-1]["plan_date"]
        earliest_min = scheduled[-1]["end_min"]
        for envelope_id, row in zip(envelope_ids, scheduled):
            planned.append({"envelope_id": envelope_id, "op_no": step["op_no"], "qty_planned": row["qty"]})
        idx += 1

    if not planned and skipped:
        raise ValueError("No schedulable operations found. Remaining operations are still missing a usable machine.")
    return {"blocks": planned, "ops_planned": len(planned), "qty_planned": qty_to_plan, "skipped_ops": skipped}


@app.route("/")
def root():
    return redirect("/process-sheets")


@app.route("/<page>")
def page(page):
    mapping = {
        "process-sheets": "process_sheets.html",
        "erp-sync": "erp_sync.html",
        "parts-flows": "parts_flows.html",
        "machines": "machines.html",
        "calendar": "calendar.html",
        "materials": "materials.html",
        "program-toollist": "program_toollist.html",
        "staffing": "staffing.html",
        "summary": "summary.html",
        "gantt": "gantt.html",
        "history": "history.html",
    }
    if page not in mapping:
        return api_error("Not found", 404)
    return render_template(mapping[page])


@app.get("/api/parts")
def api_parts():
    with db() as con:
        parts = rows(con.execute("SELECT * FROM parts ORDER BY part_name"))
        for part in parts:
            part["flows"] = rows(con.execute("SELECT flow_id, flow_code, is_default FROM part_flow_header WHERE part_id = ? ORDER BY is_default DESC, flow_id", (part["part_id"],)))
        return jsonify(parts)


@app.post("/api/parts")
def create_part():
    data = request.get_json() or {}
    with db() as con:
        cur = con.execute("INSERT INTO parts (part_name, part_desc) VALUES (?, ?)", (data.get("part_name"), data.get("part_desc", "")))
        return jsonify({"part_id": cur.lastrowid})


@app.put("/api/parts/<int:part_id>")
def update_part(part_id):
    data = request.get_json() or {}
    with db() as con:
        con.execute("UPDATE parts SET part_name = ?, part_desc = ? WHERE part_id = ?", (data.get("part_name"), data.get("part_desc", ""), part_id))
        return jsonify({"ok": True})


@app.get("/api/parts/<int:part_id>/flows")
def api_part_flows(part_id):
    with db() as con:
        ensure_part_flow_step_columns(con)
        flows = rows(con.execute("SELECT * FROM part_flow_header WHERE part_id = ? ORDER BY is_default DESC, flow_id", (part_id,)))
        for flow in flows:
            flow["steps"] = flow_steps(con, flow["flow_id"])
        return jsonify(flows)


@app.post("/api/flows")
def create_flow():
    data = request.get_json() or {}
    with db() as con:
        cur = con.execute(
            "INSERT INTO part_flow_header (part_id, flow_code, flow_name, is_default) VALUES (?, ?, ?, ?)",
            (data.get("part_id"), data.get("flow_code"), data.get("flow_name", ""), int(bool(data.get("is_default")))),
        )
        return jsonify({"flow_id": cur.lastrowid})


@app.put("/api/flows/<int:flow_id>")
def update_flow(flow_id):
    data = request.get_json() or {}
    with db() as con:
        ensure_part_flow_step_columns(con)
        flow = one(con.execute("SELECT part_id FROM part_flow_header WHERE flow_id = ?", (flow_id,)))
        if not flow:
            return api_error("Flow not found", 404)
        steps = data.get("steps", [])
        if not steps:
            return api_error("Flow must have at least one step")
        if sum(1 for step in steps if step.get("is_last_op")) != 1:
            return api_error("Flow must have exactly one last operation step")
        if data.get("is_default"):
            con.execute("UPDATE part_flow_header SET is_default = 0 WHERE part_id = ?", (flow["part_id"],))
        con.execute(
            "UPDATE part_flow_header SET flow_code = ?, flow_name = ?, is_default = ? WHERE flow_id = ?",
            (data.get("flow_code"), data.get("flow_name", ""), int(bool(data.get("is_default"))), flow_id),
        )
        existing_steps = rows(
            con.execute(
                """
                SELECT step_id, seq_no, op_no, op_type, machine_category, preferred_machine, cycle_time, setup_time, is_last_op
                FROM part_flow_steps
                WHERE flow_id = ?
                ORDER BY seq_no, step_id
                """,
                (flow_id,),
            )
        )
        existing_by_id = {row["step_id"]: row for row in existing_steps}
        existing_ids = set(existing_by_id.keys())
        incoming_ids = {int(step["step_id"]) for step in steps if step.get("step_id")}
        changed_from_seq = None

        removed_ids = [step_id for step_id in existing_ids if step_id not in incoming_ids]
        if removed_ids:
            for step_id in removed_ids:
                old_seq = int(existing_by_id[step_id]["seq_no"])
                changed_from_seq = old_seq if changed_from_seq is None else min(changed_from_seq, old_seq)
            placeholders = ",".join("?" for _ in removed_ids)
            con.execute(
                f"UPDATE planning_block SET flow_step_id = NULL WHERE flow_step_id IN ({placeholders})",
                removed_ids,
            )
            con.execute(
                f"DELETE FROM part_flow_steps WHERE step_id IN ({placeholders})",
                removed_ids,
            )

        for idx, step in enumerate(steps, 1):
            step_id = int(step["step_id"]) if step.get("step_id") else None
            params = (
                idx,
                step.get("op_no"),
                step.get("op_type"),
                step.get("machine_category", "").upper(),
                step.get("preferred_machine", ""),
                step.get("cycle_time") or 1,
                step.get("setup_time") or 0,
                int(bool(step.get("is_last_op"))),
            )
            if step_id and step_id in existing_ids:
                previous = existing_by_id[step_id]
                if (
                    int(previous["seq_no"]) != idx
                    or compact_text(previous["op_no"]) != compact_text(step.get("op_no"))
                    or compact_text(previous["op_type"]) != compact_text(step.get("op_type"))
                    or compact_text(previous["machine_category"]).upper() != compact_text(step.get("machine_category")).upper()
                    or compact_text(previous["preferred_machine"]) != compact_text(step.get("preferred_machine"))
                    or float(previous["cycle_time"] or 0) != float(step.get("cycle_time") or 1)
                    or float(previous["setup_time"] or 0) != float(step.get("setup_time") or 0)
                    or int(previous["is_last_op"] or 0) != int(bool(step.get("is_last_op")))
                ):
                    changed_from_seq = idx if changed_from_seq is None else min(changed_from_seq, idx)
                con.execute(
                    """
                    UPDATE part_flow_steps
                    SET seq_no = ?, op_no = ?, op_type = ?, machine_category = ?, preferred_machine = ?,
                        cycle_time = ?, setup_time = ?, is_last_op = ?
                    WHERE step_id = ? AND flow_id = ?
                    """,
                    params + (step_id, flow_id),
                )
            else:
                con.execute(
                    """
                    INSERT INTO part_flow_steps (flow_id, seq_no, op_no, op_type, machine_category, preferred_machine, cycle_time, setup_time, is_last_op)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (flow_id,) + params,
                )
                changed_from_seq = idx if changed_from_seq is None else min(changed_from_seq, idx)
        refresh = apply_flow_timing_updates(con, flow_id, changed_from_seq)
        mark_flow_erp_sync_locked(con, flow_id, 1)
        return jsonify({"ok": True, **refresh})


@app.post("/api/flows/<int:flow_id>/duplicate")
def duplicate_flow(flow_id):
    with db() as con:
        ensure_part_flow_step_columns(con)
        flow = one(con.execute("SELECT * FROM part_flow_header WHERE flow_id = ?", (flow_id,)))
        if not flow:
            return api_error("Flow not found", 404)
        code = f"{flow['flow_code']}-COPY"
        cur = con.execute(
            "INSERT INTO part_flow_header (part_id, flow_code, flow_name, is_default) VALUES (?, ?, ?, 0)",
            (flow["part_id"], code, flow["flow_name"]),
        )
        for step in flow_steps(con, flow_id):
            con.execute(
                """
                INSERT INTO part_flow_steps (flow_id, seq_no, op_no, op_type, machine_category, preferred_machine, cycle_time, setup_time, is_last_op)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (cur.lastrowid, step["seq"], step["op_no"], step["op_type"], step["machine_category"], step["preferred_machine"], step["cycle_time"], step["setup_time"], step["is_last_op"]),
            )
        return jsonify({"flow_id": cur.lastrowid, "flow_code": code})


@app.get("/api/process-sheets")
def api_process_sheets():
    with db() as con:
        data = rows(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id ORDER BY ps.due_date, ps.ps_id"))
        list_cache = build_process_sheet_list_cache(con, data)
        material_cache = {}
        erp_index = erp_active_sheet_index(con)
        result = [serialize_ps(con, ps, material_cache=material_cache, erp_index=erp_index, list_cache=list_cache) for ps in data]
        q = (request.args.get("search") or "").lower()
        if q:
            result = [p for p in result if q in " ".join(str(p.get(k, "")).lower() for k in ("ps_id", "source_ps_id", "pp_partial_no", "part_name", "inv_code", "inv_desc"))]
        for key in ("status", "planner_status"):
            val = request.args.get(key)
            if not val:
                continue
            target_key = "planner_status" if key == "status" else key
            result = [p for p in result if p[target_key] == val]
        if request.args.get("show_completed") != "1":
            result = [p for p in result if p["planner_status"] != "COMPLETED"]
        if request.args.get("overdue_only") == "1":
            result = [p for p in result if "OVERDUE" in p["warnings"]]
        if request.args.get("shortage_only") == "1":
            result = [p for p in result if "MATERIAL_SHORTAGE" in p["warnings"]]
        mat_ready = request.args.get("mat_ready")
        if mat_ready in {"0", "1"}:
            want_ready = (mat_ready == "1")
            filtered = []
            for p in result:
                mat = material_for_ps_cached(con, p["ps_id"], material_cache)
                if bool(mat and mat["material_ready"]) == want_ready:
                    filtered.append(p)
            result = filtered
        return jsonify(result)


@app.get("/api/process-sheets/<ps_id>")
def get_process_sheet(ps_id):
    with db() as con:
        ps = one(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id WHERE ps.ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        return jsonify(serialize_ps(con, ps, detail=True))


@app.post("/api/process-sheets")
def create_process_sheet():
    data = request.get_json() or {}
    source_ps_id = compact_text(data.get("source_ps_id") or data.get("ps_id"))
    partial_no = compact_text(data.get("pp_partial_no")) or "1"
    ps_id = process_sheet_key(source_ps_id, partial_no)
    with db() as con:
        con.execute(
            """
            INSERT INTO process_sheet (ps_id, source_ps_id, pp_partial_no, part_id, selected_flow_id, inv_code, inv_desc, order_date, due_date, total_qty, status)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (ps_id, source_ps_id, partial_no, data.get("part_id"), data.get("selected_flow_id"), data.get("inv_code", ""), data.get("inv_desc", ""), data.get("order_date", ""), data.get("due_date", ""), data.get("total_qty"), data.get("status", "ACTIVE")),
        )
        con.execute(
            "INSERT INTO process_sheet_material (ps_id, material_name, material_ready, material_ready_qty, order_status, need_by_date) VALUES (?, ?, ?, ?, ?, ?)",
            (ps_id, data.get("material_name", ""), int(bool(data.get("material_ready"))), float(data.get("total_qty") or 0) if data.get("material_ready") else 0, "NA" if data.get("material_ready") else data.get("order_status", "TO_ORDER"), data.get("need_by_date", "")),
        )
        return jsonify({"ok": True, "ps_id": ps_id, "source_ps_id": source_ps_id, "pp_partial_no": partial_no})


@app.put("/api/process-sheets/<ps_id>")
def update_process_sheet(ps_id):
    data = request.get_json() or {}
    with db() as con:
        row = one(con.execute("SELECT source_ps_id, pp_partial_no FROM process_sheet WHERE ps_id = ?", (ps_id,))) or {}
        source_ps_id = compact_text(data.get("source_ps_id") or row.get("source_ps_id") or split_process_sheet_key(ps_id)[0])
        partial_no = compact_text(data.get("pp_partial_no") or row.get("pp_partial_no") or split_process_sheet_key(ps_id)[1]) or "1"
        if process_sheet_key(source_ps_id, partial_no) != ps_id:
            return api_error("Process sheet identity is immutable after creation. Create a new partial instead.", 400)
        con.execute(
            """
            UPDATE process_sheet SET source_ps_id = ?, pp_partial_no = ?, inv_code = ?, inv_desc = ?, order_date = ?, due_date = ?,
                   total_qty = ?, status = ?, selected_flow_id = ? WHERE ps_id = ?
            """,
            (source_ps_id, partial_no, data.get("inv_code", ""), data.get("inv_desc", ""), data.get("order_date", ""), data.get("due_date", ""), data.get("total_qty"), data.get("status", "ACTIVE"), data.get("selected_flow_id"), ps_id),
        )
        mark_ps_erp_sync_locked(con, ps_id, 1)
        return jsonify({"ok": True, "ps_id": ps_id, "source_ps_id": source_ps_id, "pp_partial_no": partial_no})


@app.get("/api/process-sheets/<ps_id>/actual-output")
def get_process_sheet_actual(ps_id):
    with db() as con:
        ps = one(con.execute("SELECT total_qty FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        ps_row = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        totals = ps_totals(con, ps_id)
        return jsonify({
            "ps_id": ps_id,
            "total_qty": ps["total_qty"],
            "manual_actual": manual_actual_qty(con, ps_id),
            "erp_actual": erp_actual_qty(con, ps_id),
            "last_op_actual": totals.get("last_op_actual_qty", 0),
            "last_op_planned": totals.get("last_op_planned_qty", 0),
            "effective_actual": totals["finished_qty"],
            "status_debug": compute_process_sheet_status(con, ps_row)["status_debug"] if ps_row else {},
        })


@app.post("/api/process-sheets/<ps_id>/actual-output")
def set_process_sheet_actual(ps_id):
    data = request.get_json() or {}
    actual_qty = max(0, float(data.get("actual_qty") or 0))
    with db() as con:
        ps = one(con.execute("SELECT total_qty FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        actual_qty = min(actual_qty, ps["total_qty"])
        con.execute(
            """
            INSERT INTO process_sheet_manual_actual (ps_id, actual_qty, updated_at)
            VALUES (?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id) DO UPDATE SET actual_qty = excluded.actual_qty, updated_at = CURRENT_TIMESTAMP
            """,
            (ps_id, actual_qty),
        )
        return jsonify({"ok": True, "actual_qty": actual_qty})


@app.post("/api/process-sheets/<ps_id>/complete")
def set_process_sheet_complete(ps_id):
    data = request.get_json() or {}
    force_completed = 1 if data.get("force_completed") else 0
    with db() as con:
        ps = one(con.execute("SELECT ps_id FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        step_ids = [row["step_id"] for row in con.execute(
            """
            SELECT pfs.step_id
            FROM process_sheet ps
            JOIN part_flow_steps pfs ON pfs.flow_id = ps.selected_flow_id
            WHERE ps.ps_id = ?
            ORDER BY pfs.seq_no, pfs.step_id
            """,
            (ps_id,),
        )]
        if force_completed:
            con.execute(
                """
                INSERT INTO process_sheet_local_override (ps_id, force_completed, updated_at)
                VALUES (?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(ps_id) DO UPDATE SET force_completed = excluded.force_completed, updated_at = CURRENT_TIMESTAMP
                """,
                (ps_id, 1),
            )
            mark_ps_erp_sync_locked(con, ps_id, 1)
            for step_id in step_ids:
                con.execute(
                    """
                    INSERT INTO process_sheet_step_local_override (ps_id, step_id, force_completed, updated_at)
                    VALUES (?, ?, ?, CURRENT_TIMESTAMP)
                    ON CONFLICT(ps_id, step_id) DO UPDATE SET force_completed = excluded.force_completed, updated_at = CURRENT_TIMESTAMP
                    """,
                    (ps_id, step_id, 1),
                )
            return jsonify({"ok": True, "force_completed": True, "completion_override": "COMPLETED"})

        snapshot_id = snapshot_process_sheet_progress(con, ps_id, "MANUAL_RESET_TO_INCOMPLETE")
        reset_process_sheet_progress(con, ps_id)
        return jsonify({
            "ok": True,
            "force_completed": False,
            "completion_override": "INCOMPLETE",
            "snapshot_id": snapshot_id,
            "message": "Progress was reset to a fresh batch and saved in a recoverable snapshot.",
        })


@app.post("/api/process-sheets/<ps_id>/steps/<int:step_id>/complete")
def set_process_sheet_step_complete(ps_id, step_id):
    data = request.get_json() or {}
    force_completed = 1 if data.get("force_completed") else 0
    with db() as con:
        ps = one(con.execute("SELECT ps_id FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        step = one(con.execute("SELECT step_id FROM part_flow_steps WHERE step_id = ?", (step_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        if not step:
            return api_error("Step not found", 404)
        con.execute(
            """
            INSERT INTO process_sheet_step_local_override (ps_id, step_id, force_completed, updated_at)
            VALUES (?, ?, ?, CURRENT_TIMESTAMP)
            ON CONFLICT(ps_id, step_id) DO UPDATE SET force_completed = excluded.force_completed, updated_at = CURRENT_TIMESTAMP
            """,
            (ps_id, step_id, force_completed),
        )
        mark_ps_erp_sync_locked(con, ps_id, 1)
        if not force_completed:
            con.execute(
                """
                INSERT INTO process_sheet_local_override (ps_id, force_completed, updated_at)
                VALUES (?, 0, CURRENT_TIMESTAMP)
                ON CONFLICT(ps_id) DO UPDATE SET force_completed = 0, updated_at = CURRENT_TIMESTAMP
                """,
                (ps_id,),
            )
        return jsonify({"ok": True, "force_completed": bool(force_completed), "completion_override": "COMPLETED" if force_completed else "INCOMPLETE"})


@app.post("/api/process-sheets/<ps_id>/restore-progress")
def restore_process_sheet_progress_api(ps_id):
    data = request.get_json() or {}
    with db() as con:
        ps = one(con.execute("SELECT ps_id FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)
        try:
            restored = restore_process_sheet_progress(con, ps_id, data.get("snapshot_id"))
        except ValueError as exc:
            return api_error(str(exc))
        return jsonify({"ok": True, **restored})


def auto_plan_op(con, ps_id, step_id):
    ps = one(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id WHERE ps.ps_id = ?", (ps_id,)))
    if not ps:
        raise ValueError("Process sheet not found")

    steps = flow_steps(con, ps["selected_flow_id"])
    target = next((step for step in steps if int(step.get("step_id") or 0) == int(step_id or 0)), None)
    if not target:
        raise ValueError("Operation not found in selected flow")

    step_metrics = {row["step_id"]: row for row in process_sheet_step_metrics(con, ps_id, steps)["steps"]}
    metric = step_metrics.get(target["step_id"], {})
    planned_qty = float(metric.get("planned_qty") or 0)
    actual_qty = float(metric.get("actual_qty") or 0)
    step_remaining_qty = max(0, float(ps["total_qty"] or 0) - max(planned_qty, actual_qty))
    if step_remaining_qty <= 0:
        return {
            "op_no": target["op_no"],
            "step_id": target["step_id"],
            "planned_qty": 0,
            "scheduled_rows": 0,
            "skipped": True,
        }

    existing = existing_step_span(con, ps_id, target["step_id"])
    machine = None
    if existing and existing.get("end_key"):
        existing_machine = one(
            con.execute(
                """
                SELECT machine_id
                FROM planning_row
                WHERE ps_id = ? AND seq_no = ?
                ORDER BY plan_date DESC, start_min DESC, row_id DESC
                LIMIT 1
                """,
                (ps_id, target["seq"]),
            )
        )
        if existing_machine:
            machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (int(existing_machine["machine_id"]),)))
        anchor_date, anchor_min = split_datetime_key(existing["end_key"])
    else:
        anchor_date, anchor_min = planning_anchor_before_seq(con, ps_id, target["seq"])

    if not machine:
        machine = pick_machine_for_step(con, target, require_preferred_machine=True)
    if not machine:
        raise ValueError("No preferred machine assigned")

    effective_setup = effective_setup_time_for_new_envelope(
        con,
        ps_id,
        machine["machine_id"],
        anchor_date,
        anchor_min,
        target["setup_time"] or 0,
        target,
    )
    duration = whole_minutes(effective_setup + step_remaining_qty * (target["cycle_time"] or 1))
    scheduled_rows = schedule_rows(
        con,
        machine,
        anchor_date,
        anchor_min,
        duration,
        step_remaining_qty,
        target["cycle_time"] or 1,
        effective_setup,
        prefer_even_split=False,
    )
    created_envelope_ids = create_blocks_from_schedule(con, ps_id, target, machine, scheduled_rows)
    return {
        "op_no": target["op_no"],
        "step_id": target["step_id"],
        "planned_qty": step_remaining_qty,
        "scheduled_rows": len(scheduled_rows),
        "created_envelope_ids": created_envelope_ids,
    }


@app.post("/api/plan/auto/<ps_id>")
def api_auto_plan(ps_id):
    with db() as con:
        try:
            result = auto_plan(con, ps_id)
            return jsonify({
                "result": result,
                "op_no": "ALL",
                "ops_planned": len(result.get("scheduled_ops", [])),
                "ops_skipped": len(result.get("skipped_ops", [])),
            })
        except ValueError as exc:
            return api_error(str(exc))
        except Exception as exc:
            return api_error(f"Auto-plan failed: {exc}")


@app.post("/api/plan/auto-op/<ps_id>/<int:step_id>")
def api_auto_plan_op(ps_id, step_id):
    with db() as con:
        try:
            result = auto_plan_op(con, ps_id, step_id)
            return jsonify({"result": result, "op_no": result.get("op_no"), "planned_rows": result.get("scheduled_rows", 0)})
        except ValueError as exc:
            return api_error(str(exc))
        except Exception as exc:
            return api_error(f"Auto-plan failed: {exc}")


@app.post("/api/process-sheets/<ps_id>/unplan-unlocked")
def api_unplan_unlocked(ps_id):
    with db() as con:
        ps = one(con.execute("SELECT ps_id FROM process_sheet WHERE ps_id = ?", (ps_id,)))
        if not ps:
            return api_error("Process sheet not found", 404)

        started_blocks = rows(
            con.execute(
                """
                SELECT DISTINCT pr.envelope_id
                FROM planning_row pr
                WHERE pr.ps_id = ? AND COALESCE(pr.actual_out, 0) > 0 AND COALESCE(pr.locked, 0) = 0 AND pr.envelope_id IS NOT NULL
                """,
                (ps_id,),
            )
        )
        if started_blocks:
            return api_error("Cannot unplan unlocked envelopes that already have actual output")

        envelope_ids = [
            int(r["envelope_id"])
            for r in rows(
                con.execute(
                    """
                    SELECT envelope_id
                    FROM planning_row
                    WHERE ps_id = ? AND envelope_id IS NOT NULL
                    GROUP BY envelope_id
                    HAVING SUM(CASE WHEN COALESCE(actual_out, 0) > 0 THEN 1 ELSE 0 END) = 0
                       AND SUM(CASE WHEN COALESCE(locked, 0) = 1 THEN 1 ELSE 0 END) = 0
                    """,
                    (ps_id,),
                )
            )
            if r.get("envelope_id") is not None
        ]
        if not envelope_ids:
            return jsonify({"ok": True, "deleted_envelopes": 0, "deleted_rows": 0})

        q_marks = ",".join("?" for _ in envelope_ids)
        group_keys = rows(
            con.execute(
                f"""
                SELECT DISTINCT ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id
                FROM planning_row
                WHERE envelope_id IN ({q_marks})
                """,
                envelope_ids,
            )
        )
        deleted_rows = con.execute(
            f"DELETE FROM planning_row WHERE envelope_id IN ({q_marks})",
            envelope_ids,
        ).rowcount
        deleted_envelopes = con.execute(
            f"DELETE FROM planning_envelope WHERE envelope_id IN ({q_marks})",
            envelope_ids,
        ).rowcount
        if group_keys:
            for key in group_keys:
                con.execute(
                    """
                    DELETE FROM planning_block
                    WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                    """,
                    (key["ps_id"], key["seq_no"], key["op_no"], key["machine_id"], key["machine_code"], key["flow_step_id"]),
                )
        return jsonify({"ok": True, "deleted_envelopes": deleted_envelopes, "deleted_rows": deleted_rows})


@app.post("/api/process-sheets/unplan-unlocked-all")
def api_unplan_unlocked_all():
    with db() as con:
        deletable_envelopes = rows(
            con.execute(
                """
                SELECT envelope_id, ps_id
                FROM planning_row
                WHERE envelope_id IS NOT NULL
                GROUP BY ps_id, envelope_id
                HAVING SUM(CASE WHEN COALESCE(actual_out, 0) > 0 THEN 1 ELSE 0 END) = 0
                   AND SUM(CASE WHEN COALESCE(locked, 0) = 1 THEN 1 ELSE 0 END) = 0
                """
            )
        )
        if not deletable_envelopes:
            return jsonify({"ok": True, "deleted_envelopes": 0, "deleted_rows": 0, "affected_ps": 0, "skipped_ps": 0})

        envelope_ids = [int(row["envelope_id"]) for row in deletable_envelopes if row.get("envelope_id") is not None]
        q_marks = ",".join("?" for _ in envelope_ids)
        group_keys = rows(
            con.execute(
                f"""
                SELECT DISTINCT ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id
                FROM planning_row
                WHERE envelope_id IN ({q_marks})
                """,
                envelope_ids,
            )
        )
        deleted_rows = con.execute(
            f"DELETE FROM planning_row WHERE envelope_id IN ({q_marks})",
            envelope_ids,
        ).rowcount
        deleted_envelopes = con.execute(
            f"DELETE FROM planning_envelope WHERE envelope_id IN ({q_marks})",
            envelope_ids,
        ).rowcount
        if group_keys:
            for key in group_keys:
                con.execute(
                    """
                    DELETE FROM planning_block
                    WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                    """,
                    (key["ps_id"], key["seq_no"], key["op_no"], key["machine_id"], key["machine_code"], key["flow_step_id"]),
                )
        target_ps = sorted({row["ps_id"] for row in deletable_envelopes})
        return jsonify({
            "ok": True,
            "deleted_envelopes": deleted_envelopes,
            "deleted_rows": deleted_rows,
            "affected_ps": len(target_ps),
            "skipped_ps": 0,
        })


@app.post("/api/admin/cleanup-legacy-blocks")
def api_cleanup_legacy_blocks():
    with db() as con:
        deletable_blocks = rows(
            con.execute(
                """
                SELECT DISTINCT envelope_id, ps_id
                FROM planning_envelope
                WHERE flow_step_id IS NULL AND archived = 0
                """
            )
        )
        target_ids = [row["envelope_id"] for row in deletable_blocks]
        affected_ps = len({row["ps_id"] for row in deletable_blocks})
        deleted_envelopes = 0
        if target_ids:
            placeholders = ",".join("?" for _ in target_ids)
            deleted_envelopes = con.execute(
                f"DELETE FROM planning_envelope WHERE envelope_id IN ({placeholders})",
                target_ids,
            ).rowcount
        return jsonify({
            "ok": True,
            "deleted_envelopes": deleted_envelopes,
            "affected_ps": affected_ps,
            "skipped_started_blocks": 0,
        })


def cleanup_legacy_process_sheet_rows(con):
    """
    Remove pre-migration process_sheet rows that duplicate the normalized composite-key rows.
    Keep the normalized row and only delete the legacy row after child records are re-homed.
    """
    legacy_rows = rows(
        con.execute(
            """
            SELECT *
            FROM process_sheet
            WHERE ps_id = source_ps_id
              AND COALESCE(pp_partial_no, '1') = '1'
              AND instr(ps_id, '::') = 0
              AND EXISTS (
                  SELECT 1
                  FROM process_sheet ps2
                  WHERE ps2.source_ps_id = process_sheet.source_ps_id
                    AND COALESCE(ps2.pp_partial_no, '1') = COALESCE(process_sheet.pp_partial_no, '1')
                    AND ps2.ps_id = process_sheet.source_ps_id || '::' || COALESCE(process_sheet.pp_partial_no, '1')
              )
            ORDER BY source_ps_id
            """
        )
    )

    deleted = 0
    skipped = 0
    migrated_children = 0
    for legacy in legacy_rows:
        canonical_ps_id = process_sheet_key(legacy["source_ps_id"], legacy["pp_partial_no"])
        if not canonical_ps_id or canonical_ps_id == legacy["ps_id"]:
            skipped += 1
            continue

        canonical = one(con.execute("SELECT * FROM process_sheet WHERE ps_id = ?", (canonical_ps_id,)))
        if not canonical:
            skipped += 1
            continue

        # Re-home schedule/archive-style children first.
        migrated_children += con.execute("UPDATE planning_block SET ps_id = ? WHERE ps_id = ?", (canonical_ps_id, legacy["ps_id"])).rowcount
        migrated_children += con.execute("UPDATE planning_row SET ps_id = ? WHERE ps_id = ?", (canonical_ps_id, legacy["ps_id"])).rowcount
        migrated_children += con.execute("UPDATE history_block SET ps_id = ? WHERE ps_id = ?", (canonical_ps_id, legacy["ps_id"])).rowcount
        migrated_children += con.execute("UPDATE history_row SET ps_id = ? WHERE ps_id = ?", (canonical_ps_id, legacy["ps_id"])).rowcount
        migrated_children += con.execute(
            "UPDATE process_sheet_progress_snapshot SET ps_id = ? WHERE ps_id = ?",
            (canonical_ps_id, legacy["ps_id"]),
        ).rowcount

        # Material rows are one-per-sheet; merge the row and move order logs if the canonical row already exists.
        legacy_mat = one(con.execute("SELECT * FROM process_sheet_material WHERE ps_id = ?", (legacy["ps_id"],)))
        canonical_mat = one(con.execute("SELECT * FROM process_sheet_material WHERE ps_id = ?", (canonical_ps_id,)))
        if legacy_mat:
            if not canonical_mat:
                con.execute(
                    "UPDATE process_sheet_material SET ps_id = ? WHERE mat_id = ?",
                    (canonical_ps_id, legacy_mat["mat_id"]),
                )
                migrated_children += 1
            else:
                merged_name = canonical_mat["material_name"] or legacy_mat["material_name"] or ""
                merged_ready = max(int(canonical_mat["material_ready"] or 0), int(legacy_mat["material_ready"] or 0))
                merged_ready_qty = max(float(canonical_mat["material_ready_qty"] or 0), float(legacy_mat["material_ready_qty"] or 0))
                merged_need_by = canonical_mat["need_by_date"] or legacy_mat["need_by_date"] or ""
                merged_note = canonical_mat["planner_note"] or legacy_mat["planner_note"] or ""
                merged_order_status = canonical_mat["order_status"] or legacy_mat["order_status"] or "TO_ORDER"
                con.execute(
                    """
                    UPDATE process_sheet_material
                    SET material_name = ?, material_ready = ?, material_ready_qty = ?, need_by_date = ?,
                        planner_note = ?, order_status = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE mat_id = ?
                    """,
                    (
                        merged_name,
                        merged_ready,
                        merged_ready_qty,
                        merged_need_by,
                        merged_note,
                        merged_order_status,
                        canonical_mat["mat_id"],
                    ),
                )
                if legacy_mat["mat_id"] != canonical_mat["mat_id"]:
                    con.execute(
                        "UPDATE process_sheet_material_order_log SET mat_id = ?, ps_id = ? WHERE mat_id = ?",
                        (canonical_mat["mat_id"], canonical_ps_id, legacy_mat["mat_id"]),
                    )
                    con.execute("DELETE FROM process_sheet_material WHERE mat_id = ?", (legacy_mat["mat_id"],))
                    migrated_children += 1

        # One-row-per-sheet tables can safely merge by keeping the latest / strongest value.
        singleton_tables = (
            "process_sheet_erp_link",
            "process_sheet_erp_actual",
            "process_sheet_local_override",
        )
        for table in singleton_tables:
            legacy_row = one(con.execute(f"SELECT * FROM {table} WHERE ps_id = ?", (legacy["ps_id"],)))
            canonical_row = one(con.execute(f"SELECT * FROM {table} WHERE ps_id = ?", (canonical_ps_id,)))
            if not legacy_row:
                continue
            if not canonical_row:
                con.execute(
                    f"UPDATE {table} SET ps_id = ? WHERE ps_id = ?",
                    (canonical_ps_id, legacy["ps_id"]),
                )
                migrated_children += 1
                continue
            if table == "process_sheet_erp_link":
                merged = {
                    "pp_partial_no": canonical_row["pp_partial_no"] or legacy_row["pp_partial_no"] or "1",
                    "bom_code": canonical_row["bom_code"] or legacy_row["bom_code"] or "",
                    "erp_status": canonical_row["erp_status"] or legacy_row["erp_status"] or "",
                }
            elif table == "process_sheet_erp_actual":
                merged = {
                    "actual_qty": max(float(canonical_row["actual_qty"] or 0), float(legacy_row["actual_qty"] or 0)),
                    "reject_qty": max(float(canonical_row["reject_qty"] or 0), float(legacy_row["reject_qty"] or 0)),
                    "source_sheet": canonical_row["source_sheet"] or legacy_row["source_sheet"] or "Workorder Tracker",
                }
            else:
                merged = {
                    "force_completed": max(int(canonical_row["force_completed"] or 0), int(legacy_row["force_completed"] or 0)),
                }
            if table == "process_sheet_erp_link":
                con.execute(
                    """
                    UPDATE process_sheet_erp_link
                    SET pp_partial_no = ?, bom_code = ?, erp_status = ?, last_sync_at = CURRENT_TIMESTAMP
                    WHERE ps_id = ?
                    """,
                    (merged["pp_partial_no"], merged["bom_code"], merged["erp_status"], canonical_ps_id),
                )
                con.execute("DELETE FROM process_sheet_erp_link WHERE ps_id = ?", (legacy["ps_id"],))
            elif table == "process_sheet_erp_actual":
                con.execute(
                    """
                    UPDATE process_sheet_erp_actual
                    SET actual_qty = ?, reject_qty = ?, source_sheet = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE ps_id = ?
                    """,
                    (merged["actual_qty"], merged["reject_qty"], merged["source_sheet"], canonical_ps_id),
                )
                con.execute("DELETE FROM process_sheet_erp_actual WHERE ps_id = ?", (legacy["ps_id"],))
            else:
                con.execute(
                    """
                    UPDATE process_sheet_local_override
                    SET force_completed = ?, updated_at = CURRENT_TIMESTAMP
                    WHERE ps_id = ?
                    """,
                    (merged["force_completed"], canonical_ps_id),
                )
                con.execute("DELETE FROM process_sheet_local_override WHERE ps_id = ?", (legacy["ps_id"],))

        # Per-step overrides: keep the strongest completion signal.
        step_rows = rows(con.execute("SELECT * FROM process_sheet_step_local_override WHERE ps_id = ?", (legacy["ps_id"],)))
        for row in step_rows:
            con.execute(
                """
                INSERT INTO process_sheet_step_local_override (ps_id, step_id, force_completed, updated_at)
                VALUES (?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(ps_id, step_id) DO UPDATE SET
                  force_completed = MAX(process_sheet_step_local_override.force_completed, excluded.force_completed),
                  updated_at = CURRENT_TIMESTAMP
                """,
                (canonical_ps_id, row["step_id"], int(row["force_completed"] or 0)),
            )
        migrated_children += len(step_rows)
        if step_rows:
            con.execute("DELETE FROM process_sheet_step_local_override WHERE ps_id = ?", (legacy["ps_id"],))

        support_rows = rows(con.execute("SELECT * FROM process_sheet_support WHERE ps_id = ?", (legacy["ps_id"],)))
        for row in support_rows:
            con.execute(
                """
                INSERT INTO process_sheet_support (ps_id, support_type, status, need_by_date, promised_date, ready_date, note, updated_by, updated_at)
                VALUES (?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)
                ON CONFLICT(ps_id, support_type) DO UPDATE SET
                  status = CASE
                      WHEN excluded.status = 'COMPLETED' OR process_sheet_support.status = 'COMPLETED' THEN 'COMPLETED'
                      WHEN excluded.status = 'READY' OR process_sheet_support.status = 'READY' THEN 'READY'
                      ELSE COALESCE(excluded.status, process_sheet_support.status)
                  END,
                  need_by_date = COALESCE(NULLIF(excluded.need_by_date, ''), process_sheet_support.need_by_date),
                  promised_date = COALESCE(NULLIF(excluded.promised_date, ''), process_sheet_support.promised_date),
                  ready_date = COALESCE(NULLIF(excluded.ready_date, ''), process_sheet_support.ready_date),
                  note = COALESCE(NULLIF(excluded.note, ''), process_sheet_support.note),
                  updated_by = COALESCE(NULLIF(excluded.updated_by, ''), process_sheet_support.updated_by),
                  updated_at = CURRENT_TIMESTAMP
                """,
                (
                    canonical_ps_id,
                    row["support_type"],
                    row["status"],
                    row["need_by_date"],
                    row["promised_date"],
                    row["ready_date"],
                    row["note"],
                    row["updated_by"],
                ),
            )
        if support_rows:
            migrated_children += len(support_rows)
            con.execute("DELETE FROM process_sheet_support WHERE ps_id = ?", (legacy["ps_id"],))

        remaining_refs = one(
            con.execute(
                """
                SELECT
                    (SELECT COUNT(*) FROM planning_block WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM planning_row WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM history_block WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM history_row WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_progress_snapshot WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_material WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_erp_link WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_erp_actual WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_local_override WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_step_local_override WHERE ps_id = ?) +
                    (SELECT COUNT(*) FROM process_sheet_support WHERE ps_id = ?) AS refs
                """,
                (
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                    legacy["ps_id"],
                ),
            )
        )["refs"]
        if remaining_refs == 0:
            con.execute("DELETE FROM process_sheet WHERE ps_id = ?", (legacy["ps_id"],))
            deleted += 1
        else:
            skipped += 1

    return {"deleted": deleted, "skipped": skipped, "migrated_children": migrated_children, "candidates": len(legacy_rows)}


@app.post("/api/admin/cleanup-legacy-process-sheets")
def api_cleanup_legacy_process_sheets():
    with db() as con:
        result = cleanup_legacy_process_sheet_rows(con)
        return jsonify({"ok": True, **result})


@app.post("/api/admin/normalize-standard-rows")
def api_normalize_standard_rows():
    with db() as con:
        result = normalize_all_standard_rows(con)
        return jsonify({"ok": True, **result})


@app.post("/api/plan/auto-all")
def api_auto_plan_all():
    planned = 0
    with db() as con:
        for row in con.execute("SELECT ps_id FROM process_sheet WHERE status = 'ACTIVE' ORDER BY due_date, ps_id"):
            try:
                auto_plan(con, row["ps_id"])
                planned += 1
            except ValueError:
                pass
        return jsonify({"planned": planned})


@app.post("/api/integrations/erp/test")
def api_test_erp_connection():
    data = request.get_json(silent=True) or {}
    try:
        result = erp_test_connection(
            host=data.get("host"),
            port=data.get("port"),
            dbname=data.get("dbname"),
            user=data.get("user"),
            password=data.get("password"),
        )
        return jsonify({"success": True, **result})
    except Exception as exc:
        return api_error(f"ERP connection test failed: {exc}")


@app.get("/api/erp-sync/history")
def api_erp_sync_history():
    with db() as con:
        return jsonify(recent_erp_syncs(con))


@app.get("/api/erp-sync/debug")
def api_erp_sync_debug():
    with db() as con:
        staged_bom = rows(
            con.execute(
                """
                SELECT bom_code, inventory_code, stage_no, stage_desc, machine_no, machine_category, cycle_time, setup_time, raw_payload
                FROM erp_bom_op_stage_staging
                ORDER BY stage_id DESC
                LIMIT 20
                """
            )
        )
        flows = rows(
            con.execute(
                """
                SELECT pfh.flow_code, p.part_name, pfs.op_no, pfs.op_type, pfs.machine_category, pfs.preferred_machine
                FROM part_flow_header pfh
                JOIN parts p ON p.part_id = pfh.part_id
                JOIN part_flow_steps pfs ON pfs.flow_id = pfh.flow_id
                WHERE pfh.flow_code LIKE 'BOM-%'
                ORDER BY pfh.flow_id DESC, pfs.seq_no
                LIMIT 40
                """
            )
        )
        return jsonify({
            "staged_bom_op_stage": staged_bom,
            "erp_bom_flows": flows,
        })


@app.post("/api/erp-sync/upload")
def api_erp_sync_upload():
    file = request.files.get("file")
    active_orders_file = request.files.get("active_orders_file")
    if not file and not active_orders_file:
        return api_error("Choose an ERP workbook or an active orders workbook first.")
    if file and getattr(file, "filename", "") and not file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        return api_error("Upload an .xlsx or .xlsm workbook.")
    if active_orders_file and getattr(active_orders_file, "filename", "") and not active_orders_file.filename.lower().endswith((".xlsx", ".xlsm", ".xltx", ".xltm")):
        return api_error("Active orders file must be an .xlsx or .xlsm workbook.")
    with db() as con:
        try:
            if file and getattr(file, "filename", ""):
                result = run_erp_excel_sync(con, file, active_orders_file=active_orders_file)
            else:
                result = run_active_orders_only_sync(con, active_orders_file)
            return jsonify({"success": True, **result})
        except Exception as exc:
            con.rollback()
            return api_error(f"ERP sync failed: {exc}")


@app.get("/api/machines")
def api_machines():
    with db() as con:
        return jsonify(rows(con.execute("SELECT * FROM machines ORDER BY machine_code")))


@app.post("/api/machines")
def create_machine():
    data = request.get_json() or {}
    with db() as con:
        cur = con.execute(
            "INSERT INTO machines (machine_code, machine_category, shift_profile, active, notes) VALUES (?, ?, ?, ?, ?)",
            (data.get("machine_code"), data.get("machine_category"), data.get("shift_profile", "STANDARD"), int(data.get("active", 1)), data.get("notes", "")),
        )
        refresh_flow_step_machine_categories(con, data.get("machine_code"))
        return jsonify({"machine_id": cur.lastrowid})


@app.put("/api/machines/<int:machine_id>")
def update_machine(machine_id):
    data = request.get_json() or {}
    with db() as con:
        existing = one(con.execute("SELECT machine_code FROM machines WHERE machine_id = ?", (machine_id,)))
        con.execute(
            "UPDATE machines SET machine_code = ?, machine_category = ?, shift_profile = ?, active = ?, notes = ? WHERE machine_id = ?",
            (data.get("machine_code"), data.get("machine_category"), data.get("shift_profile", "STANDARD"), int(data.get("active", 1)), data.get("notes", ""), machine_id),
        )
        if existing and existing.get("machine_code") and existing["machine_code"] != data.get("machine_code"):
            refresh_flow_step_machine_categories(con, existing["machine_code"])
        refresh_flow_step_machine_categories(con, data.get("machine_code"))
        return jsonify({"ok": True})


@app.get("/api/calendar")
def api_calendar():
    start = request.args.get("from", date.today().isoformat())
    end = request.args.get("to", (date.today() + timedelta(days=30)).isoformat())
    with db() as con:
        return jsonify(rows(con.execute("SELECT * FROM calendar_days WHERE work_date BETWEEN ? AND ? ORDER BY work_date", (start, end))))


@app.post("/api/calendar")
def update_calendar():
    data = request.get_json() or {}
    with db() as con:
        for item in data.get("updates", []):
            con.execute(
                """
                INSERT INTO calendar_days (work_date, is_working_day, note) VALUES (?, ?, ?)
                ON CONFLICT(work_date) DO UPDATE SET is_working_day = excluded.is_working_day, note = excluded.note
                """,
                (item.get("work_date"), int(item.get("is_working_day", 1)), item.get("note", "")),
            )
        return jsonify({"ok": True})


@app.get("/api/materials")
def api_materials():
    with db() as con:
        mats = []
        for ps in planned_support_process_sheets(con):
            mat = ensure_material_record(con, ps["ps_id"]) or {}
            expected_start = ps.get("expected_start") or ""
            mat["ps_id"] = ps["ps_id"]
            mat["part_name"] = ps.get("part_name", "")
            mat["inv_desc"] = ps.get("inv_desc", "")
            mat["total_qty"] = ps.get("total_qty", 0)
            mat["planned_qty"] = ps.get("planned_qty", 0)
            mat["finished_qty"] = ps.get("finished_qty", 0)
            mat["remaining_qty"] = ps.get("remaining_qty", 0)
            mat["expected_start"] = expected_start
            mat["need_by_date"] = subtract_working_days(con, expected_start[:10], MATERIAL_LEAD_DAYS) if expected_start else ""
            mat["order_logs"] = rows(con.execute("SELECT * FROM process_sheet_material_order_log WHERE mat_id = ? ORDER BY order_date, log_id", (mat["mat_id"],)))
            mat["covered_qty"] = (mat["material_ready_qty"] or 0) + sum(log["received_qty"] or 0 for log in mat["order_logs"])
            mats.append(mat)
        return jsonify(mats)


@app.put("/api/materials/<int:mat_id>")
def update_material(mat_id):
    data = request.get_json() or {}
    with db() as con:
        con.execute(
            """
            UPDATE process_sheet_material SET material_name = ?, material_ready = ?, material_ready_qty = ?,
                   order_status = ?, need_by_date = ?, planner_note = ? WHERE mat_id = ?
            """,
            (data.get("material_name", ""), int(bool(data.get("material_ready"))), data.get("material_ready_qty", 0), data.get("order_status", "TO_ORDER"), data.get("need_by_date", ""), data.get("planner_note", ""), mat_id),
        )
        return jsonify({"ok": True})


@app.post("/api/materials/<int:mat_id>/order-log")
def add_order_log(mat_id):
    data = request.get_json() or {}
    with db() as con:
        con.execute(
            """
            INSERT INTO process_sheet_material_order_log (mat_id, ps_id, ordered_qty, received_qty, order_date, expected_date, log_status, note)
            SELECT ?, ps_id, ?, ?, ?, ?, ?, ? FROM process_sheet_material WHERE mat_id = ?
            """,
            (mat_id, data.get("ordered_qty", 0), data.get("received_qty", 0), data.get("order_date", ""), data.get("expected_date", ""), data.get("log_status", "PENDING"), data.get("note", ""), mat_id),
        )
        return jsonify({"ok": True})


@app.get("/api/program-toollist")
def api_program_toollist():
    with db() as con:
        result = []
        for ps in planned_support_process_sheets(con):
            expected_start = ps.get("expected_start") or ""
            program = sync_support_need_by(con, ps["ps_id"], "PROGRAM", expected_start)
            toollist = sync_support_need_by(con, ps["ps_id"], "TOOLLIST", expected_start)
            result.append({
                "ps_id": ps["ps_id"],
                "part_name": ps.get("part_name", ""),
                "inv_desc": ps.get("inv_desc", ""),
                "total_qty": ps.get("total_qty", 0),
                "planned_qty": ps.get("planned_qty", 0),
                "finished_qty": ps.get("finished_qty", 0),
                "remaining_qty": ps.get("remaining_qty", 0),
                "expected_start": expected_start,
                "route_label": ps.get("route_label", ""),
                "program": program,
                "toollist": toollist,
            })
        return jsonify(result)


@app.put("/api/support/<ps_id>/<support_type>")
def update_support(ps_id, support_type):
    data = request.get_json() or {}
    support_type = compact_text(support_type).upper()
    if support_type not in {"PROGRAM", "TOOLLIST"}:
        return api_error("Unsupported support type")
    with db() as con:
        expected_start = (schedule_span(con, ps_id) or {}).get("expected_start") or ""
        record = sync_support_need_by(con, ps_id, support_type, expected_start)
        con.execute(
            """
            UPDATE process_sheet_support
            SET status = ?, promised_date = ?, ready_date = ?, note = ?, updated_by = ?, updated_at = CURRENT_TIMESTAMP
            WHERE support_id = ?
            """,
            (
                data.get("status", record.get("status", "PENDING")),
                data.get("promised_date", record.get("promised_date", "")),
                data.get("ready_date", record.get("ready_date", "")),
                data.get("note", record.get("note", "")),
                data.get("updated_by", ""),
                record["support_id"],
            ),
        )
        return jsonify({"ok": True})


@app.get("/api/staff")
def api_staff():
    with db() as con:
        return jsonify(rows(con.execute("SELECT * FROM staff ORDER BY staff_name")))


@app.post("/api/staff")
def create_staff():
    data = request.get_json() or {}
    with db() as con:
        cur = con.execute("INSERT INTO staff (staff_name, role) VALUES (?, ?)", (data.get("staff_name"), data.get("role", "MACHINIST")))
        return jsonify({"staff_id": cur.lastrowid})


@app.get("/api/staffing")
def api_staffing():
    start = request.args.get("from", date.today().isoformat())
    end = request.args.get("to", (date.today() + timedelta(days=14)).isoformat())
    with db() as con:
        return jsonify(rows(con.execute(
            """
            SELECT a.*, m.machine_code, m.machine_category, s.staff_name, s.role
            FROM machine_staff_assignment a
            JOIN machines m ON m.machine_id = a.machine_id
            JOIN staff s ON s.staff_id = a.staff_id
            WHERE a.assign_date BETWEEN ? AND ? ORDER BY a.assign_date, m.machine_code
            """,
            (start, end),
        )))


@app.post("/api/staffing")
def create_staffing():
    data = request.get_json() or {}
    with db() as con:
        cur = con.execute(
            "INSERT INTO machine_staff_assignment (machine_id, staff_id, assign_date, shift) VALUES (?, ?, ?, ?)",
            (data.get("machine_id"), data.get("staff_id"), data.get("assign_date"), data.get("shift", "DAY")),
        )
        return jsonify({"assign_id": cur.lastrowid})


@app.delete("/api/staffing/<int:assign_id>")
def delete_staffing(assign_id):
    with db() as con:
        con.execute("DELETE FROM machine_staff_assignment WHERE assign_id = ?", (assign_id,))
        return jsonify({"ok": True})


@app.get("/api/summary")
def api_summary():
    view = compact_text(request.args.get("view") or "").lower()
    start = request.args.get("from") or date.today().isoformat()
    end = request.args.get("to") or date.today().isoformat()
    category = request.args.get("category")
    with db() as con:
        if view == "machines":
            return jsonify(machine_summary(con, start, end, category))
        data = rows(con.execute("SELECT ps.*, p.part_name FROM process_sheet ps JOIN parts p ON p.part_id = ps.part_id ORDER BY ps.due_date, ps.ps_id"))
        result = []
        for ps in data:
            item = serialize_ps(con, ps)
            span = one(con.execute(
                """
                SELECT MIN(pr.plan_date || ' ' || printf('%02d:%02d', pr.start_min / 60, pr.start_min % 60)) first_start,
                       MAX(pr.plan_date || ' ' || printf('%02d:%02d', pr.end_min / 60, pr.end_min % 60)) last_end,
                       GROUP_CONCAT(DISTINCT m.machine_code) machines_used
                FROM planning_envelope pe
                JOIN planning_row pr ON pr.envelope_id = pe.envelope_id
                JOIN machines m ON m.machine_id = pe.machine_id
                WHERE pe.ps_id = ? AND pe.archived = 0
                """,
                (ps["ps_id"],),
            ))
            item.update(span)
            item["expected_start"] = item.get("first_start")
            item["expected_end"] = item.get("last_end")
            result.append(item)
        return jsonify(result)


def _range_dates(start, end):
    try:
        start_d = date.fromisoformat(start)
    except Exception:
        start_d = date.today()
    try:
        end_d = date.fromisoformat(end)
    except Exception:
        end_d = start_d
    if end_d < start_d:
        start_d, end_d = end_d, start_d
    return start_d, end_d


def _working_days_in_range(con, start_d, end_d):
    days = 0
    current = start_d
    while current <= end_d:
        cal = one(con.execute("SELECT is_working_day FROM calendar_days WHERE work_date = ?", (current.isoformat(),)))
        if cal is None or int(cal.get("is_working_day") or 0) == 1:
            days += 1
        current += timedelta(days=1)
    return days


def _machine_available_minutes(con, shift_profile, start_d, end_d):
    days = _working_days_in_range(con, start_d, end_d)
    if compact_text(shift_profile).upper() == "24HR":
        return days * 1440
    return days * sum(win_end - win_start for win_start, win_end in STANDARD_WINDOWS)


def machine_summary(con, start, end, machine_category=None):
    start_d, end_d = _range_dates(start, end)
    machine_category = compact_text(machine_category)
    machine_types = [
        compact_text(row["machine_category"])
        for row in rows(
            con.execute(
                """
                SELECT DISTINCT machine_category
                FROM machines
                WHERE active = 1
                  AND COALESCE(machine_category, '') <> ''
                ORDER BY machine_category
                """
            )
        )
        if compact_text(row["machine_category"])
    ]
    ps_total_lookup = {
        compact_text(row["source_ps_id"]) or compact_text(row["ps_id"]): float(row["total_qty"] or 0)
        for row in rows(
            con.execute(
                """
                SELECT ps_id, COALESCE(NULLIF(source_ps_id, ''), ps_id) AS source_ps_id, total_qty
                FROM process_sheet
                """
            )
        )
        if compact_text(row["source_ps_id"]) or compact_text(row["ps_id"])
    }
    active_machines = rows(con.execute(
        """
        SELECT machine_id, machine_code, machine_category, shift_profile
        FROM machines
        WHERE active = 1
        {category_clause}
        ORDER BY machine_code
        """
        .format(category_clause="AND machine_category = ?" if machine_category else ""),
        (machine_category,) if machine_category else (),
    ))
    by_machine = {}
    for machine in active_machines:
        machine_id = int(machine.get("machine_id") or 0)
        machine_code = compact_text(machine.get("machine_code")) or "UNKNOWN"
        by_machine[(machine_id, machine_code)] = {
            "machine_id": machine_id,
            "machine_code": machine_code,
            "machine_category": compact_text(machine.get("machine_category")) or "UNKNOWN",
            "shift_profile": compact_text(machine.get("shift_profile")) or "STANDARD",
            "planned_minutes": 0.0,
            "active_minutes": 0.0,
            "history_minutes": 0.0,
            "planned_qty": 0.0,
            "actual_qty": 0.0,
            "rows": 0,
            "active_rows": 0,
            "history_rows": 0,
            "locked_rows": 0,
            "ps_ids": set(),
            "order_qty_total": 0.0,
            "detail_total_qty": 0.0,
            "detail_actual_qty": 0.0,
            "details_map": {},
            "first_start": None,
            "last_end": None,
        }
    active_rows = rows(con.execute(
        """
        SELECT
            'ACTIVE' AS source,
            pe.machine_id,
            m.machine_code,
            m.machine_category,
            m.shift_profile,
            ps.ps_id,
            ps.source_ps_id,
            pe.seq_no,
            pe.flow_step_id,
            fs.op_no,
            fs.op_type,
            pr.plan_date,
            pr.start_min,
            pr.end_min,
            pr.qty,
            pr.actual_out,
            pr.actual_out_set,
            pr.locked,
            pr.envelope_id
        FROM planning_envelope pe
        JOIN planning_row pr ON pr.envelope_id = pe.envelope_id
        JOIN process_sheet ps ON ps.ps_id = pe.ps_id
        JOIN part_flow_steps fs ON fs.step_id = pe.flow_step_id
        JOIN machines m ON m.machine_id = pe.machine_id
        WHERE pe.archived = 0
          AND pr.plan_date BETWEEN ? AND ?
        """,
        (start_d.isoformat(), end_d.isoformat()),
    ))
    history_rows = rows(con.execute(
        """
        SELECT
            'HISTORY' AS source,
            he.machine_id,
            he.machine_code,
            m.machine_category,
            m.shift_profile,
            ps.ps_id,
            ps.source_ps_id,
            he.seq_no,
            he.flow_step_id,
            hr.plan_date,
            hr.start_min,
            hr.end_min,
            hr.qty,
            hr.actual_out,
            hr.actual_out_set,
            1 AS locked,
            hr.hist_envelope_id AS envelope_id,
            he.op_no,
            COALESCE(pfs.op_type, he.op_no) AS op_type
        FROM history_envelope he
        JOIN history_row hr ON hr.hist_envelope_id = he.hist_envelope_id
        JOIN process_sheet ps ON ps.ps_id = he.ps_id
        LEFT JOIN machines m ON m.machine_code = he.machine_code
        LEFT JOIN part_flow_steps pfs
               ON pfs.flow_id = ps.selected_flow_id
              AND pfs.seq_no = he.seq_no
        WHERE hr.plan_date BETWEEN ? AND ?
        """,
        (start_d.isoformat(), end_d.isoformat()),
    ))
    all_rows = active_rows + history_rows
    for row in all_rows:
        machine_code = compact_text(row.get("machine_code")) or "UNKNOWN"
        machine_id = int(row.get("machine_id") or 0)
        row_category = compact_text(row.get("machine_category"))
        if machine_category and row_category != machine_category:
            continue
        key = (machine_id, machine_code)
        bucket = by_machine.setdefault(key, {
            "machine_id": machine_id,
            "machine_code": machine_code,
            "machine_category": compact_text(row.get("machine_category")) or "UNKNOWN",
            "shift_profile": compact_text(row.get("shift_profile")) or "STANDARD",
            "planned_minutes": 0.0,
            "active_minutes": 0.0,
            "history_minutes": 0.0,
            "planned_qty": 0.0,
            "actual_qty": 0.0,
            "rows": 0,
            "active_rows": 0,
            "history_rows": 0,
            "locked_rows": 0,
            "ps_ids": set(),
            "order_qty_total": 0.0,
            "detail_total_qty": 0.0,
            "detail_actual_qty": 0.0,
            "details_map": {},
            "first_start": None,
            "last_end": None,
        })
        start_min = int(row.get("start_min") or 0)
        end_min = int(row.get("end_min") or 0)
        duration = max(0, end_min - start_min)
        ps_key = compact_text(row.get("source_ps_id")) or compact_text(row.get("ps_id"))
        op_no = compact_text(row.get("op_no")) or "0"
        op_type = compact_text(row.get("op_type")) or op_no
        op_label = op_type if op_no and op_no in op_type else f"{op_type} {op_no}".strip()
        seq_no = int(row.get("seq_no") or 0)
        flow_step_id = int(row.get("flow_step_id") or 0)
        detail_key = (ps_key, seq_no, op_no, flow_step_id)
        detail = bucket["details_map"].setdefault(detail_key, {
            "ps_id": ps_key,
            "sheet_id": compact_text(row.get("ps_id")),
            "op_no": op_no,
            "op_type": op_type,
            "op_label": op_label,
            "seq_no": seq_no,
            "flow_step_id": flow_step_id,
            "start_time": None,
            "end_time": None,
            "total_qty": ps_total_lookup.get(ps_key, 0.0),
            "planned_qty": 0.0,
            "actual_qty": 0.0,
            "row_count": 0,
        })
        row_start = f"{row.get('plan_date')} {start_min // 60:02d}:{start_min % 60:02d}"
        row_end = f"{row.get('plan_date')} {end_min // 60:02d}:{end_min % 60:02d}"
        if detail["start_time"] is None or row_start < detail["start_time"]:
            detail["start_time"] = row_start
        if detail["end_time"] is None or row_end > detail["end_time"]:
            detail["end_time"] = row_end
        detail["planned_qty"] += float(row.get("qty") or 0)
        if int(row.get("actual_out_set") or 0) == 1:
            detail["actual_qty"] += float(row.get("actual_out") or 0)
        detail["row_count"] += 1
        bucket["rows"] += 1
        bucket["planned_minutes"] += duration
        bucket["planned_qty"] += float(row.get("qty") or 0)
        bucket["actual_qty"] += float(row.get("actual_out") or 0) if int(row.get("actual_out_set") or 0) == 1 else 0.0
        bucket["locked_rows"] += 1 if int(row.get("locked") or 0) == 1 else 0
        bucket["ps_ids"].add(compact_text(row.get("source_ps_id")) or compact_text(row.get("ps_id")))
        if compact_text(row.get("source")) == "HISTORY":
            bucket["history_rows"] += 1
            bucket["history_minutes"] += duration
        else:
            bucket["active_rows"] += 1
            bucket["active_minutes"] += duration
        if bucket["first_start"] is None or row_start < bucket["first_start"]:
            bucket["first_start"] = row_start
        if bucket["last_end"] is None or row_end > bucket["last_end"]:
            bucket["last_end"] = row_end

    result = []
    for bucket in by_machine.values():
        ps_ids = sorted({ps_id for ps_id in bucket["ps_ids"] if ps_id})
        if ps_ids:
            q_marks = ",".join("?" for _ in ps_ids)
            row = one(
                con.execute(
                    f"SELECT COALESCE(SUM(total_qty), 0) AS order_qty_total FROM process_sheet WHERE COALESCE(NULLIF(source_ps_id, ''), ps_id) IN ({q_marks})",
                    ps_ids,
                )
            ) or {"order_qty_total": 0}
            bucket["order_qty_total"] = float(row.get("order_qty_total") or 0)
        available_minutes = _machine_available_minutes(con, bucket["shift_profile"], start_d, end_d)
        utilization = (bucket["planned_minutes"] / available_minutes * 100.0) if available_minutes else 0.0
        idle_minutes = max(0.0, available_minutes - bucket["planned_minutes"])
        schedule_completion = (bucket["actual_qty"] / bucket["planned_qty"] * 100.0) if bucket["planned_qty"] else 0.0
        next_avail_date, next_avail_min = next_start(con, bucket["machine_id"])
        next_available = f"{next_avail_date} {int(next_avail_min or 0) // 60:02d}:{int(next_avail_min or 0) % 60:02d}" if next_avail_date else None
        result.append({
            **bucket,
            "ps_count": len(ps_ids),
            "available_minutes": available_minutes,
            "idle_minutes": idle_minutes,
            "utilization_pct": utilization,
            "schedule_completion_pct": schedule_completion,
            "order_completion_pct": 0.0,
            "completion_pct": schedule_completion,
            "total_qty": bucket["order_qty_total"],
            "planned_qty": bucket["planned_qty"],
            "actual_qty": bucket["actual_qty"],
            "start_time": bucket["first_start"],
            "end_time": bucket["last_end"],
            "planned_hours": round(bucket["planned_minutes"] / 60.0, 2),
            "active_hours": round(bucket["active_minutes"] / 60.0, 2),
            "history_hours": round(bucket["history_minutes"] / 60.0, 2),
            "available_hours": round(available_minutes / 60.0, 2),
            "idle_hours": round(idle_minutes / 60.0, 2),
            "next_available": next_available,
            "details": [],
        })
    result.sort(key=lambda item: (-item["planned_minutes"], item["machine_code"]))
    for item in result:
        item["ps_ids"] = sorted(item["ps_ids"])
        details = list(item.pop("details_map", {}).values())
        for detail in details:
            total_qty = float(detail.get("total_qty") or 0)
            planned_qty = float(detail.get("planned_qty") or 0)
            actual_qty = float(detail.get("actual_qty") or 0)
            detail["schedule_completion_pct"] = (actual_qty / planned_qty * 100.0) if planned_qty else 0.0
            detail["order_completion_pct"] = min(100.0, (actual_qty / total_qty * 100.0) if total_qty else 0.0)
        details.sort(key=lambda d: (d["start_time"] or "", d["end_time"] or "", d["ps_id"] or "", d["op_no"] or "", d["seq_no"] or 0))
        detail_total_qty = sum(float(detail.get("total_qty") or 0) for detail in details)
        detail_actual_qty = sum(float(detail.get("actual_qty") or 0) for detail in details)
        order_completion = min(100.0, (detail_actual_qty / detail_total_qty * 100.0) if detail_total_qty else 0.0)
        item["detail_total_qty"] = detail_total_qty
        item["detail_actual_qty"] = detail_actual_qty
        item["order_actual_qty"] = detail_actual_qty
        item["order_completion_pct"] = order_completion
        item["total_qty"] = detail_total_qty
        item["details"] = details
    return {"rows": result, "machine_types": machine_types}


@app.get("/api/gantt")
def api_gantt():
    start = request.args.get("from", date.today().isoformat())
    end = request.args.get("to", (date.today() + timedelta(days=7)).isoformat())
    start_d = date.fromisoformat(start)
    end_d = date.fromisoformat(end)
    dates = [(start_d + timedelta(days=i)).isoformat() for i in range((end_d - start_d).days + 1)]
    with db() as con:
        def abs_min(plan_date, minute):
            return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

        def row_progress_qty(row):
            planned = max(0.0, float(row.get("qty") or row.get("row_qty") or 0))
            actual = max(0.0, float(row.get("actual_out") or 0))
            actual_set = int(row.get("actual_out_set") or 0) == 1
            return actual if actual_set else planned

        def cumulative_qty(rows_for_seq, point, qty_field):
            total = 0.0
            for item in rows_for_seq:
                start_abs = item["abs_start"]
                end_abs = item["abs_end"]
                qty = max(0.0, float(item.get(qty_field) or 0))
                if qty <= 0:
                    continue
                if point <= start_abs:
                    continue
                if end_abs <= start_abs or point >= end_abs:
                    total += qty
                    continue
                total += qty * ((point - start_abs) / (end_abs - start_abs))
            return total

        blocks = rows(con.execute(
            """
            SELECT pe.envelope_id, pe.envelope_id AS display_env_id, NULL AS hist_envelope_id,
                   CASE WHEN pe.locked = 1 OR COALESCE(pr.locked, 0) = 1 THEN 1 ELSE 0 END AS locked,
                   pe.machine_id, pr.row_id, pr.plan_date, pr.start_min, pr.end_min,
                   pr.qty row_qty, pr.actual_out, pr.actual_out_set, pr.setup_mins, pr.locked AS row_locked,
                   COALESCE((SELECT SUM(pr2.qty) FROM planning_row pr2 WHERE pr2.envelope_id = pe.envelope_id), 0) AS envelope_qty_all,
                   ps.ps_id, ps.source_ps_id, ps.pp_partial_no, ps.due_date,
                   fs.op_no, fs.op_type, pe.seq_no, m.machine_code,
                   0 AS is_history, 0 AS archived_view
            FROM planning_envelope pe
            JOIN planning_row pr ON pr.envelope_id = pe.envelope_id
            JOIN process_sheet ps ON ps.ps_id = pe.ps_id
            JOIN part_flow_steps fs ON fs.step_id = pe.flow_step_id
            JOIN machines m ON m.machine_id = pe.machine_id
            WHERE pe.archived = 0 AND pr.plan_date BETWEEN ? AND ?

            UNION ALL

            SELECT he.hist_envelope_id,
                   ('H-' || he.hist_envelope_id) AS display_env_id,
                   he.hist_envelope_id,
                   1 AS locked,
                   he.machine_id,
                   hr.hist_row_id AS row_id,
                   hr.plan_date,
                   hr.start_min,
                   hr.end_min,
                   hr.qty AS row_qty,
                   hr.actual_out,
                   hr.actual_out_set,
                   0 AS setup_mins,
                   1 AS row_locked,
                   COALESCE((SELECT SUM(hr2.qty) FROM history_row hr2 WHERE hr2.hist_envelope_id = he.hist_envelope_id), 0) AS envelope_qty_all,
                   ps.ps_id,
                   ps.source_ps_id,
                   ps.pp_partial_no,
                   ps.due_date,
                   he.op_no,
                   COALESCE(pfs.op_type, he.op_no) AS op_type,
                   he.seq_no,
                   he.machine_code,
                   1 AS is_history,
                   1 AS archived_view
            FROM history_envelope he
            JOIN history_row hr ON hr.hist_envelope_id = he.hist_envelope_id
            JOIN process_sheet ps ON ps.ps_id = he.ps_id
            LEFT JOIN machines m ON m.machine_code = he.machine_code
            LEFT JOIN part_flow_steps pfs
                   ON pfs.flow_id = ps.selected_flow_id
                  AND pfs.seq_no = he.seq_no
            WHERE hr.plan_date BETWEEN ? AND ?
            ORDER BY plan_date, start_min
            """,
            (start, end, start, end),
        ))

        active_rows = rows(con.execute(
            """
            SELECT pe.ps_id, pe.seq_no, pr.row_id, pr.plan_date, pr.start_min, pr.end_min,
                   pr.qty, pr.actual_out, pr.actual_out_set, pr.setup_mins, pr.locked AS row_locked,
                   COALESCE((SELECT SUM(pr2.qty) FROM planning_row pr2 WHERE pr2.envelope_id = pe.envelope_id), 0) AS envelope_qty_all
            FROM planning_envelope pe
            JOIN planning_row pr ON pr.envelope_id = pe.envelope_id
            JOIN part_flow_steps fs ON fs.step_id = pe.flow_step_id
            WHERE pe.archived = 0
            """
        ))
        history_rows = rows(con.execute(
            """
            SELECT he.ps_id, he.seq_no, hr.hist_row_id AS row_id, hr.plan_date, hr.start_min, hr.end_min,
                   hr.qty, hr.actual_out, hr.actual_out_set, 0 AS setup_mins, 1 AS row_locked,
                   COALESCE((SELECT SUM(hr2.qty) FROM history_row hr2 WHERE hr2.hist_envelope_id = he.hist_envelope_id), 0) AS envelope_qty_all
            FROM history_envelope he
            JOIN history_row hr ON hr.hist_envelope_id = he.hist_envelope_id
            """
        ))

        rows_by_ps_seq = {}
        for item in active_rows + history_rows:
            seq_no = int(item.get("seq_no") or 0)
            prepared = dict(item)
            prepared["abs_start"] = abs_min(item["plan_date"], item["start_min"])
            prepared["abs_end"] = abs_min(item["plan_date"], item["end_min"])
            prepared["planned_qty"] = max(0.0, float(item.get("qty") or 0))
            prepared["progress_qty"] = row_progress_qty(item)
            prepared["setup_mins"] = max(0.0, float(item.get("setup_mins") or 0))
            rows_by_ps_seq.setdefault((item["ps_id"], seq_no), []).append(prepared)

        for group in rows_by_ps_seq.values():
            group.sort(key=lambda r: (r["abs_start"], r["abs_end"], r.get("row_id") or 0))

        for block in blocks:
            seq_no = int(block.get("seq_no") or 0)
            current_rows = rows_by_ps_seq.get((block["ps_id"], seq_no), [])
            previous_rows = rows_by_ps_seq.get((block["ps_id"], seq_no - 1), [])
            row_start_abs = abs_min(block["plan_date"], block["start_min"])
            row_end_abs = abs_min(block["plan_date"], block["end_min"])
            block["sequence_overlap"] = False
            block["sequence_overlap_start_min"] = None
            block["sequence_overlap_ratio"] = 0
            block["setup_mins"] = max(0.0, float(block.get("setup_mins") or 0))
            if seq_no <= 1 or not current_rows or not previous_rows or row_end_abs <= row_start_abs:
                continue

            overlap_start_abs = None
            for point in range(row_start_abs, row_end_abs + 1):
                downstream_done = cumulative_qty(current_rows, point, "progress_qty")
                upstream_ready = cumulative_qty(previous_rows, point, "progress_qty")
                if downstream_done > upstream_ready + 1e-9:
                    overlap_start_abs = point
                    break

            if overlap_start_abs is None:
                continue

            local_min = block["start_min"] + (overlap_start_abs - row_start_abs)
            ratio = 1.0 if row_end_abs == row_start_abs else max(0.0, min(1.0, (local_min - block["start_min"]) / (block["end_min"] - block["start_min"])))
            block["sequence_overlap"] = True
            block["sequence_overlap_start_min"] = int(local_min)
            block["sequence_overlap_ratio"] = ratio

        envelope_anchor_by_row_id = {}
        envelope_rows_by_anchor = {}
        active_ps_ids = sorted({item["ps_id"] for item in blocks if int(item.get("is_history") or 0) == 0})
        for ps_id in active_ps_ids:
            timeline_rows = ps_timeline_rows(con, ps_id)
            for row in timeline_rows:
                row_id = int(row.get("row_id") or 0)
                if row_id <= 0 or row_id in envelope_anchor_by_row_id:
                    continue
                envelope = contiguous_envelope_for_row(con, timeline_rows, row_id)
                if not envelope:
                    continue
                anchor_row_id = min(int(r.get("row_id") or row_id) for r in envelope if r.get("row_id") is not None)
                envelope_rows_by_anchor.setdefault(anchor_row_id, [])
                for env_row in envelope:
                    env_row_id = int(env_row.get("row_id") or 0)
                    if env_row_id > 0:
                        envelope_anchor_by_row_id[env_row_id] = anchor_row_id
                        envelope_rows_by_anchor[anchor_row_id].append(env_row_id)

        for block in blocks:
            if int(block.get("is_history") or 0) == 0:
                anchor_row_id = envelope_anchor_by_row_id.get(int(block.get("row_id") or 0))
                if anchor_row_id:
                    block["display_env_id"] = f"E-{block['ps_id']}-{block['seq_no']}-{anchor_row_id}"
                    block["envelope_anchor_row_id"] = anchor_row_id
                    block["envelope_row_ids"] = envelope_rows_by_anchor.get(anchor_row_id, [int(block.get("row_id") or 0)])
            else:
                block["envelope_anchor_row_id"] = int(block.get("row_id") or 0)
                block["envelope_row_ids"] = [int(block.get("row_id") or 0)]
        return jsonify({
            "machines": rows(con.execute("SELECT * FROM machines WHERE active = 1 ORDER BY machine_code")),
            "blocks": blocks,
            "calendar": rows(con.execute("SELECT * FROM calendar_days WHERE work_date BETWEEN ? AND ? ORDER BY work_date", (start, end))),
            "dates": dates,
        })


def _group_key_from_row(row):
    return (
        compact_text(row.get("ps_id")),
        int(row.get("seq_no") or 0),
        compact_text(row.get("op_no")),
        int(row.get("machine_id") or 0),
        compact_text(row.get("machine_code")),
        int(row.get("flow_step_id") or 0),
    )


def _row_split_group_id(row):
    return compact_text(row.get("split_group_id"))


def _row_split_piece_id(row):
    return compact_text(row.get("split_piece_id"))


def _split_group_marker(prefix="split"):
    return compact_text(prefix) or "split"


def _split_piece_marker(prefix="piece"):
    base = compact_text(prefix) or "piece"
    return f"{base}-{uuid4().hex[:12]}"


def _apply_split_group_ids(con, row_ids, split_group_id):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return
    q_marks = ",".join("?" for _ in row_ids)
    con.execute(
        f"UPDATE planning_row SET split_group_id = ? WHERE row_id IN ({q_marks})",
        [split_group_id, *row_ids],
    )
    con.execute(
        f"""
        UPDATE planning_envelope
        SET split_group_id = ?, updated_at = CURRENT_TIMESTAMP
        WHERE envelope_id IN (
            SELECT DISTINCT envelope_id
            FROM planning_row
            WHERE row_id IN ({q_marks})
        )
        """,
        [split_group_id, *row_ids],
    )


def _apply_split_piece_ids(con, row_ids, split_piece_id):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return
    q_marks = ",".join("?" for _ in row_ids)
    con.execute(
        f"UPDATE planning_row SET split_piece_id = ? WHERE row_id IN ({q_marks})",
        [split_piece_id, *row_ids],
    )
    con.execute(
        f"""
        UPDATE planning_envelope
        SET split_piece_id = ?, updated_at = CURRENT_TIMESTAMP
        WHERE envelope_id IN (
            SELECT DISTINCT envelope_id
            FROM planning_row
            WHERE row_id IN ({q_marks})
        )
        """,
        [split_piece_id, *row_ids],
    )


def _sync_envelope_summary_from_envelope_id(con, envelope_id):
    envelope_id = int(envelope_id or 0)
    if envelope_id <= 0:
        return False
    envelope_rows = rows(
        con.execute(
            """
            SELECT row_id, ps_id, plan_date, start_min, end_min, qty, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
            FROM planning_row
            WHERE envelope_id = ?
            ORDER BY plan_date, start_min, end_min, row_id
            """,
            (envelope_id,),
        )
    )
    if not envelope_rows:
        return False
    first = envelope_rows[0]
    last = envelope_rows[-1]
    total_qty = sum(float(r.get("qty") or 0) for r in envelope_rows)
    locked = 1 if all(int(r.get("locked") or 0) == 1 for r in envelope_rows) else 0
    split_group_id = compact_text(first.get("split_group_id") or "")
    split_piece_id = compact_text(first.get("split_piece_id") or "")
    envelope_start = f"{first['plan_date']} {hhmm(first['start_min'])}"
    envelope_end = f"{last['plan_date']} {hhmm(last['end_min'])}"
    con.execute(
        """
        UPDATE planning_envelope
        SET total_qty = ?, locked = ?, archived = 0, envelope_start = ?, envelope_end = ?, split_group_id = ?, split_piece_id = ?, updated_at = CURRENT_TIMESTAMP
        WHERE envelope_id = ?
        """,
        (total_qty, locked, envelope_start, envelope_end, split_group_id, split_piece_id, envelope_id),
    )
    con.execute(
        """
        UPDATE planning_block
        SET total_qty = ?, locked = ?, archived = 0, envelope_start = ?, envelope_end = ?, updated_at = CURRENT_TIMESTAMP
        WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
        """,
        (
            total_qty,
            locked,
            envelope_start,
            envelope_end,
            first["ps_id"],
            first["seq_no"],
            first["op_no"],
            first["machine_id"],
            first["machine_code"],
            first["flow_step_id"],
        ),
    )
    return True


def _sync_all_planning_envelopes(con):
    ps_ids = [row["ps_id"] for row in rows(con.execute("SELECT DISTINCT ps_id FROM planning_row ORDER BY ps_id"))]
    for ps_id in ps_ids:
        row_ids = [
            int(row["row_id"])
            for row in rows(
                con.execute(
                    """
                    SELECT row_id
                    FROM planning_row
                    WHERE ps_id = ?
                    ORDER BY plan_date, start_min, row_id
                    """,
                    (ps_id,),
                )
            )
        ]
        if row_ids:
            _sync_group_locks_from_row_ids(con, row_ids)
        active_envelope_ids = [
            int(row["envelope_id"])
            for row in rows(
                con.execute(
                    """
                    SELECT DISTINCT envelope_id
                    FROM planning_row
                    WHERE ps_id = ? AND COALESCE(envelope_id, 0) > 0
                    """,
                    (ps_id,),
                )
            )
        ]
        if active_envelope_ids:
            q_marks = ",".join("?" for _ in active_envelope_ids)
            con.execute(
                f"""
                UPDATE planning_envelope
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND COALESCE(archived, 0) = 0 AND envelope_id NOT IN ({q_marks})
                """,
                [ps_id, *active_envelope_ids],
            )
        else:
            con.execute(
                """
                UPDATE planning_envelope
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND COALESCE(archived, 0) = 0
                """,
                (ps_id,),
            )
    envelope_ids = [
        int(row["envelope_id"])
        for row in rows(
            con.execute(
                """
                SELECT DISTINCT envelope_id
                FROM planning_row
                WHERE COALESCE(envelope_id, 0) > 0
                ORDER BY envelope_id
                """
            )
        )
    ]
    for envelope_id in envelope_ids:
        _sync_envelope_summary_from_envelope_id(con, envelope_id)


def _rows_for_group_key(con, group_key):
    ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id = group_key
    return rows(
        con.execute(
            """
            SELECT *
            FROM planning_row
            WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
            ORDER BY plan_date, start_min, row_id
            """,
            (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id),
        )
    )


def _sync_group_lock_from_rows(con, row_ids):
    return _sync_group_locks_from_row_ids(con, row_ids)


def _sync_group_locks_from_row_ids(con, row_ids):
    row_ids = [int(r) for r in row_ids or [] if str(r).strip().isdigit()]
    if not row_ids:
        return []
    q_marks = ",".join("?" for _ in row_ids)
    rows_in_scope = rows(
        con.execute(
            f"""
            SELECT row_id, ps_id, plan_date, start_min, end_min, qty, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, envelope_id, split_group_id, split_piece_id
            FROM planning_row
            WHERE row_id IN ({q_marks})
            ORDER BY plan_date, start_min, row_id
            """,
            row_ids,
        )
    )
    grouped = {}
    for row in rows_in_scope:
        grouped.setdefault(_group_key_from_row(row), []).append(row)

    touched_envelope_ids = []
    for group_key, group_rows in grouped.items():
        ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id = group_key
        split_group_id = compact_text(group_rows[0].get("split_group_id") or "") if group_rows else ""
        run_rows = []
        for row in group_rows:
            if run_rows and not rows_share_visible_envelope(con, run_rows[-1], row, ignore_envelope_id=int(run_rows[0].get("envelope_id") or 0)):
                run_rows = []
            run_rows.append(row)
            if not run_rows:
                continue
            next_row = group_rows[group_rows.index(row) + 1] if group_rows.index(row) + 1 < len(group_rows) else None
            if next_row is not None and rows_share_visible_envelope(con, row, next_row, ignore_envelope_id=int(run_rows[0].get("envelope_id") or 0)):
                continue
            locked = 1 if run_rows and all(int(r.get("locked") or 0) == 1 for r in run_rows) else 0
            run_start = f"{run_rows[0]['plan_date']} {hhmm(run_rows[0]['start_min'])}"
            run_end = f"{run_rows[-1]['plan_date']} {hhmm(run_rows[-1]['end_min'])}"
            envelope = one(
                con.execute(
                    """
                    SELECT envelope_id
                    FROM planning_envelope
                    WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                      AND COALESCE(archived, 0) = 0
                      AND envelope_start = ? AND envelope_end = ?
                    LIMIT 1
                    """,
                    (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, run_start, run_end),
                )
            )
            if envelope:
                envelope_id = int(envelope["envelope_id"])
            else:
                con.execute(
                    """
                    INSERT INTO planning_envelope (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id, total_qty, locked, archived, envelope_start, envelope_end)
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?, 0, ?, 0, ?, ?)
                    """,
                    (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, group_rows[0].get("split_piece_id") or "", locked, run_start, run_end),
                )
                envelope = one(
                    con.execute(
                        """
                        SELECT envelope_id
                        FROM planning_envelope
                        WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                          AND COALESCE(archived, 0) = 0
                          AND envelope_start = ? AND envelope_end = ?
                        LIMIT 1
                        """,
                        (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, run_start, run_end),
                    )
                )
                if not envelope:
                    continue
                envelope_id = int(envelope["envelope_id"])
            touched_envelope_ids.append(envelope_id)
            run_row_ids = [int(r["row_id"]) for r in run_rows]
            run_q_marks = ",".join("?" for _ in run_row_ids)
            con.execute(
                f"UPDATE planning_row SET envelope_id = ? WHERE row_id IN ({run_q_marks})",
                [envelope_id, *run_row_ids],
            )
            _sync_envelope_summary_from_envelope_id(con, envelope_id)
            run_rows = []
    return touched_envelope_ids


def _load_row_envelope(con, row_id):
    row = one(
        con.execute(
            """
            SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                   setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
            FROM planning_row
            WHERE row_id = ?
            """,
            (int(row_id),),
        )
    )
    if not row:
        return None, []
    timeline_rows = ps_timeline_rows(con, row["ps_id"])
    envelope_rows = contiguous_envelope_for_row(con, timeline_rows, row_id)
    if envelope_rows:
        return row, envelope_rows
    envelope_id = int(row.get("envelope_id") or 0)
    if envelope_id > 0:
        envelope_rows = rows(
            con.execute(
                """
                SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                       setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                FROM planning_row
                WHERE envelope_id = ?
                ORDER BY plan_date, start_min, row_id
                """,
                (envelope_id,),
            )
        )
        if envelope_rows:
            return row, envelope_rows
    group_key = _group_key_from_row(row)
    envelope_rows = list(_rows_for_group_key(con, group_key))
    return row, envelope_rows


@app.post("/api/rows/<int:row_id>/lock")
def lock_row_by_id(row_id):
    data = request.get_json() or {}
    locked = int(bool(data.get("locked")))
    with db() as con:
        row = one(con.execute("SELECT * FROM planning_row WHERE row_id = ?", (row_id,)))
        if not row:
            return api_error("Planning row not found", 404)
        con.execute("UPDATE planning_row SET locked = ? WHERE row_id = ?", (locked, row_id))
        envelope_id = int(row.get("envelope_id") or 0)
        if not envelope_id:
            _, envelope_rows = _load_row_envelope(con, row_id)
            if envelope_rows:
                envelope_id = int(envelope_rows[0].get("envelope_id") or 0)
        if envelope_id > 0:
            _sync_envelope_summary_from_envelope_id(con, envelope_id)
        return jsonify({"ok": True, "row_id": row_id, "locked": locked})


@app.post("/api/rows/<int:row_id>/envelope-lock")
def lock_envelope_by_row(row_id):
    data = request.get_json() or {}
    locked = int(bool(data.get("locked")))
    with db() as con:
        row = one(con.execute("SELECT * FROM planning_row WHERE row_id = ?", (row_id,)))
        if not row:
            return api_error("Planning row not found", 404)
        timeline_rows = ps_timeline_rows(con, row["ps_id"])
        envelope = contiguous_envelope_for_row(con, timeline_rows, row_id)
        if not envelope:
            return api_error("Envelope row not found", 404)
        row_ids = [int(r["row_id"]) for r in envelope if r.get("row_id") is not None]
        if not row_ids:
            return api_error("Envelope row not found", 404)
        q_marks = ",".join("?" for _ in row_ids)
        con.execute(f"UPDATE planning_row SET locked = ? WHERE row_id IN ({q_marks})", [locked, *row_ids])
        _sync_group_locks_from_row_ids(con, row_ids)
        return jsonify({"ok": True, "row_id": row_id, "row_ids": row_ids, "locked": locked})


@app.post("/api/process-sheets/<ps_id>/ops/<op_no>/lock")
def lock_operation(ps_id, op_no):
    data = request.get_json() or {}
    locked = int(bool(data.get("locked")))
    with db() as con:
        op_rows = rows(con.execute(
            """
            SELECT row_id
            FROM planning_row
            WHERE ps_id = ? AND op_no = ?
            ORDER BY row_id
            """,
            (ps_id, op_no),
        ))
        if not op_rows:
            return api_error("Operation not found", 404)
        row_ids = [int(r["row_id"]) for r in op_rows]
        q_marks = ",".join("?" for _ in row_ids)
        con.execute(f"UPDATE planning_row SET locked = ? WHERE row_id IN ({q_marks})", [locked, *row_ids])
        _sync_group_locks_from_row_ids(con, row_ids)
        return jsonify({"ok": True, "ps_id": ps_id, "op_no": op_no, "locked": locked})

@app.post("/api/rows/<int:row_id>/actual-output")
def actual_output_for_row(row_id):
    data = request.get_json(silent=True) or {}
    data["row_id"] = row_id
    return actual_output_by_row(data)


@app.post("/api/rows/<int:row_id>/day-actual-output")
def actual_output_for_row_day(row_id):
    data = request.get_json() or {}
    data["row_id"] = row_id
    return actual_output_by_row(data)


@app.post("/api/rows/<int:row_id>/move")
def move_row(row_id):
    data = request.get_json(silent=True) or {}
    data["row_id"] = row_id
    return move_row_by_id(data)


@app.post("/api/rows/<int:row_id>/move-qty")
def move_row_qty(row_id):
    data = request.get_json(silent=True) or {}
    data["row_id"] = row_id
    return move_row_qty_by_id(data)


@app.post("/api/rows/<int:row_id>/split")
def split_row(row_id):
    data = request.get_json(silent=True) or {}
    data["row_id"] = row_id
    return split_row_by_id(data)


@app.delete("/api/rows/<int:row_id>")
def delete_row(row_id):
    data = request.get_json(silent=True) or {}
    data["row_id"] = row_id
    return delete_row_by_id(data)


@app.post("/api/rows/<int:row_id>/history")
def archive_row(row_id):
    return archive_row_by_id({"row_id": row_id})

def actual_output_by_row(data=None):
    data = data or {}
    with db() as con:
        row_id = data.get("row_id")
        if row_id in (None, ""):
            return api_error("Planning row not found", 404)
        anchor_row, timeline_rows = _load_row_envelope(con, int(row_id))
        if not anchor_row:
            return api_error("Planning row not found", 404)
        raw_actual = data.get("actual_out") if "actual_out" in data else None
        clear_actual = raw_actual in (None, "")
        force_reflow = bool(data.get("force_reflow"))
        actual_value = 0.0 if clear_actual else max(0, float(raw_actual or 0))
        plan_date = compact_text(data.get("plan_date"))
        target_row_id = row_id
        planned_qty = 0.0
        target_row = None

        def saved_row_snapshot(snapshot_row_id):
            snapshot = one(
                con.execute(
                    """
                    SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                           seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                    FROM planning_row
                    WHERE row_id = ?
                    """,
                    (int(snapshot_row_id),),
                )
            )
            return dict(snapshot) if snapshot else None

        if plan_date:
            day_rows = [r for r in contiguous_envelope_for_row(con, timeline_rows, int(row_id)) if r.get("plan_date") == plan_date]
            if not day_rows:
                return api_error("Planning row not found", 404)
            planned_qty = sum(float(r.get("qty") or 0) for r in day_rows)
            existing_actual_qty = sum(float(r.get("actual_out") or 0) for r in day_rows if int(r.get("actual_out_set") or 0) == 1)
            existing_actual_set = any(int(r.get("actual_out_set") or 0) == 1 for r in day_rows)
            comparison_qty = existing_actual_qty if existing_actual_set else planned_qty
            target_row = next((r for r in day_rows if int(r.get("row_id") or 0) == int(row_id or day_rows[-1]["row_id"])), day_rows[-1])
            if not target_row:
                return api_error("Planning row not found", 404)
            if clear_actual:
                for day_row in day_rows:
                    con.execute(
                        "UPDATE planning_row SET actual_out = 0, actual_out_set = 0 WHERE row_id = ?",
                        (day_row["row_id"],),
                    )
                return jsonify({"ok": True, "adjusted": None, "cleared": True, "requested_actual": actual_value, "saved_row": saved_row_snapshot(day_rows[-1]["row_id"])})
            remaining = actual_value
            for idx, day_row in enumerate(day_rows):
                row_actual = max(0.0, remaining) if idx == len(day_rows) - 1 else min(remaining, float(day_row.get("qty") or 0))
                con.execute(
                    "UPDATE planning_row SET actual_out = ?, actual_out_set = 1 WHERE row_id = ?",
                    (row_actual, day_row["row_id"]),
                )
                remaining = max(0.0, remaining - row_actual)
            target_row_id = day_rows[-1]["row_id"]
        else:
            row = next((r for r in timeline_rows if int(r.get("row_id") or 0) == int(row_id)), None)
            if not row:
                return api_error("Planning row not found", 404)
            planned_qty = float(row["qty"] or 0)
            existing_actual = row or {"actual_out_set": 0, "actual_out": 0}
            existing_actual_qty = float(existing_actual.get("actual_out") or 0) if int(existing_actual.get("actual_out_set") or 0) == 1 else 0.0
            comparison_qty = existing_actual_qty if int(existing_actual.get("actual_out_set") or 0) == 1 else planned_qty
            target_row = row
            if clear_actual:
                con.execute(
                    "UPDATE planning_row SET actual_out = 0, actual_out_set = 0 WHERE row_id = ?",
                    (row_id,),
                )
                return jsonify({"ok": True, "adjusted": None, "cleared": True, "requested_actual": actual_value, "saved_row": saved_row_snapshot(row_id)})
            con.execute(
                "UPDATE planning_row SET actual_out = ?, actual_out_set = 1 WHERE row_id = ?",
                (actual_value, row_id),
            )
        adjusted = None
        locked_no_replan = False
        finishup_mode = compact_text(data.get("finishup_mode") or data.get("carry_mode") or "next_free") or "next_free"
        try:
            delta_qty = actual_value - comparison_qty
            if abs(delta_qty) <= 1e-9 and not force_reflow:
                return jsonify({"ok": True, "adjusted": None, "locked_no_replan": False, "requested_actual": actual_value, "saved_row": saved_row_snapshot(target_row_id)})

            if force_reflow and abs(delta_qty) <= 1e-9:
                delta_qty = actual_value - planned_qty

            con.execute("SAVEPOINT actual_reflow")
            adjusted = rebalance_after_row_overrun_tail(con, target_row_id, delta_qty, finishup_mode)
            seq_row_ids = [
                int(r["row_id"])
                for r in rows(
                    con.execute(
                        """
                        SELECT row_id
                        FROM planning_row
                        WHERE ps_id = ? AND seq_no = ?
                        ORDER BY plan_date, start_min, row_id
                        """,
                        (target_row["ps_id"], int(target_row["seq_no"] or 0)),
                    )
                )
            ]
            merge_contiguous_rows_for_rows(con, seq_row_ids if seq_row_ids else [target_row_id])
            _, post_rows = _load_row_envelope(con, target_row_id)
            current_row_ids = [int(r["row_id"]) for r in post_rows if r.get("row_id") is not None]
            if target_row:
                affected_row_ids = seq_row_ids if seq_row_ids else current_row_ids
                envelope_summary = update_envelope_summary_from_rows(con, affected_row_ids)
                if envelope_summary and affected_row_ids:
                    _sync_group_locks_from_row_ids(con, affected_row_ids)
                validation_issues = validate_sequence_after_row(con, target_row["ps_id"], int(target_row["seq_no"] or 0), target_row_id)
            if validation_issues and not force_reflow:
                    con.execute("ROLLBACK TO actual_reflow")
                    con.execute("RELEASE actual_reflow")
                    adjusted = adjusted or {}
                    adjusted["validation_issues"] = validation_issues
                    adjusted["clash_detected"] = True
                    return jsonify({
                        "ok": True,
                        "adjusted": adjusted,
                        "locked_no_replan": False,
                        "needs_clash_resolution": True,
                    })
            con.execute("RELEASE actual_reflow")
            return jsonify({"ok": True, "adjusted": adjusted, "locked_no_replan": False, "requested_actual": actual_value, "saved_row": saved_row_snapshot(target_row_id)})
        except ValueError as exc:
            try:
                con.execute("ROLLBACK TO actual_reflow")
                con.execute("RELEASE actual_reflow")
            except Exception:
                pass
            if "locked" in str(exc).lower():
                locked_no_replan = True
            else:
                con.rollback()
                return api_error(str(exc))
        return jsonify({"ok": True, "adjusted": adjusted, "locked_no_replan": locked_no_replan, "requested_actual": actual_value, "saved_row": saved_row_snapshot(target_row_id)})


def delete_row_by_id(data=None):
    data = data or {}
    with db() as con:
        row_id = data.get("row_id")
        if row_id in (None, ""):
            return api_error("Row id is required")
        row = one(
            con.execute(
                """
                SELECT row_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                       setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, envelope_id
                FROM planning_row
                WHERE row_id = ?
                """,
                (int(row_id),),
            )
        )
        if not row:
            return api_error("Planning row not found", 404)
        group_key = _group_key_from_row(row)
        source_rows = sorted(
            [dict(r) for r in _rows_for_group_key(con, group_key)],
            key=lambda r: (r["plan_date"], int(r["start_min"] or 0), int(r["end_min"] or 0), int(r["row_id"] or 0)),
        )
        if not source_rows:
            source_rows = [dict(row)]
        selected_index = next((idx for idx, r in enumerate(source_rows) if int(r["row_id"]) == int(row_id)), 0)
        delete_rows = source_rows[selected_index:]
        delete_row_ids = sorted({int(r["row_id"]) for r in delete_rows if int(r.get("row_id") or 0) > 0})
        if not delete_row_ids:
            delete_row_ids = [int(row_id)]
        if any(int(r.get("locked") or 0) == 1 for r in delete_rows):
            return api_error("Locked rows cannot be deleted")
        q_marks = ",".join("?" for _ in delete_row_ids)
        con.execute(f"DELETE FROM planning_row WHERE row_id IN ({q_marks})", delete_row_ids)
        remaining_rows = list(_rows_for_group_key(con, group_key))
        remaining_row_ids = [int(r["row_id"]) for r in remaining_rows]
        if remaining_row_ids:
            _sync_group_locks_from_row_ids(con, remaining_row_ids)
        else:
            con.execute(
                """
                UPDATE planning_envelope
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                """,
                group_key,
            )
            con.execute(
                """
                UPDATE planning_block
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                """,
                group_key,
            )
        return jsonify({"ok": True})


def move_row_by_id(data=None):
    try:
        data = data or (request.get_json() or {})
        with db() as con:
            row_id = data.get("row_id")
            clash_mode = compact_text(data.get("clash_mode") or "")
            requested_envelope_row_ids = [
                int(v)
                for v in (data.get("envelope_row_ids") or [])
                if str(v).strip().isdigit() and int(v) > 0
            ]
            if row_id in (None, ""):
                return api_error("Row id is required")
            selected_row = one(
                con.execute(
                    """
                    SELECT row_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                           locked, seq_no, op_no, machine_id, machine_code, flow_step_id, envelope_id, split_piece_id
                    FROM planning_row
                    WHERE row_id = ?
                    """,
                    (int(row_id),),
                )
            )
            if not selected_row:
                return api_error("Planning row not found", 404)
            _, source_rows = _load_row_envelope(con, int(row_id))
            if not source_rows:
                return api_error("Envelope has no rows", 404)
            source_rows = sorted(
                [dict(r) for r in source_rows],
                key=lambda r: (r["plan_date"], int(r["start_min"] or 0), int(r["end_min"] or 0), int(r["row_id"] or 0)),
            )
            selected_index = next((idx for idx, r in enumerate(source_rows) if int(r["row_id"]) == int(row_id)), 0)
            keep_rows = source_rows[:selected_index]
            move_rows = source_rows[selected_index:]
            source_row_id_set = {int(r["row_id"]) for r in move_rows}
            remaining_row_ids = [int(r["row_id"]) for r in keep_rows]
            move_row_id_set = {int(r["row_id"]) for r in move_rows}
            moving_entire_envelope = selected_index == 0 and bool(move_row_id_set)
            if any(int(r.get("locked") or 0) == 1 for r in move_rows):
                return api_error("Locked rows cannot be moved")
            if not moving_entire_envelope and any(float(r.get("actual_out") or 0) > 0 for r in move_rows):
                return api_error("Rows with actual output cannot be moved")
            primary_row = move_rows[0]
            step = one(con.execute("SELECT * FROM part_flow_steps WHERE step_id = ?", (int(primary_row["flow_step_id"] or 0),)))
            if not step:
                return api_error("Flow step not found", 404)
            machine_id = int(data.get("machine_id") or primary_row["machine_id"])
            machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (machine_id,)))
            if not machine:
                return api_error("Machine not found", 404)
            new_date = data.get("new_date") or primary_row["plan_date"]
            new_start = int(data.get("new_start_min", STANDARD_START))
            requested_date = new_date
            requested_start = new_start
            source_envelope_id = int(primary_row.get("envelope_id") or 0) or None
            setup_time = effective_setup_time_for_new_envelope(
                con,
                primary_row["ps_id"],
                machine_id,
                new_date,
                new_start,
                step["setup_time"] or 0,
                step,
                ignore_envelope_id=source_envelope_id,
            )
            move_qty = sum(float(r.get("qty") or 0) for r in move_rows)
            required_mins = duration_for_units(step["cycle_time"] or 1, move_qty, setup_time)
            if not clash_mode:
                clash_date, clash_start = first_unblocked_start_for_move(
                    con,
                    machine,
                    new_date,
                    new_start,
                    "weekdays",
                    date.fromisoformat(new_date),
                    ignore_envelope_id=source_envelope_id,
                    ignore_row_ids=source_row_id_set,
                )
                if clash_date != new_date or int(clash_start) != int(new_start):
                    shifted_probe = _shift_rows_to_anchor(move_rows, clash_date, clash_start)
                    return jsonify({
                        "ok": True,
                        "needs_clash_resolution": True,
                        "requested_date": requested_date,
                        "requested_start_min": requested_start,
                        "clash_date": clash_date,
                        "clash_start_min": int(clash_start),
                        "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], shifted_probe)[:8]],
                        "message": "Requested slot clashes with existing work. Choose how to resolve it.",
                        "clash_options": [
                            {"label": "Cancel", "value": "cancel"},
                            {"label": "Plan but skip clash", "value": "skip_clash"},
                            {"label": "Next free slot", "value": "next_free"},
                        ],
                    })
            if clash_mode == "next_free":
                new_date, new_start = first_unblocked_start_for_move(
                    con,
                    machine,
                    new_date,
                    new_start,
                    "weekdays",
                    date.fromisoformat(new_date),
                    ignore_envelope_id=source_envelope_id,
                    ignore_row_ids=source_row_id_set,
                )
            move_row_ids = list(move_row_id_set)
            search_date = new_date
            search_start = int(new_start)
            last_error = None
            move_split_marker = next((compact_text(r.get("split_group_id")) for r in move_rows if compact_text(r.get("split_group_id"))), "")
            move_piece_id = next((compact_text(r.get("split_piece_id")) for r in move_rows if compact_text(r.get("split_piece_id"))), "") or _split_piece_marker("piece")
            for _ in range(180):
                con.execute("SAVEPOINT move_envelope")
                try:
                    if move_row_ids:
                        q_marks = ",".join("?" for _ in move_row_ids)
                        con.execute(f"DELETE FROM planning_row WHERE row_id IN ({q_marks})", move_row_ids)
                    shifted_rows = _planned_move_rows(
                        con,
                        machine,
                        move_rows,
                        search_date,
                        search_start,
                        move_qty,
                        step["cycle_time"] or 1,
                        setup_time,
                        ignore_envelope_id=source_envelope_id,
                    )
                    for row in shifted_rows:
                        row["split_piece_id"] = move_piece_id
                    blockers = _shifted_rows_blockers(con, machine["machine_id"], shifted_rows) if machine.get("shift_profile") == "24HR" else []
                    if blockers:
                        raise ValueError("Requested slot clashes with existing work")
                    split_envelopes = any(
                        not rows_share_envelope(con, machine, shifted_rows[i], shifted_rows[i + 1])
                        for i in range(len(shifted_rows) - 1)
                    )
                    created_envelope_ids = create_blocks_from_schedule(
                        con,
                        primary_row["ps_id"],
                        step,
                        machine,
                        shifted_rows,
                        split_group_id=move_split_marker,
                        merge_adjacent_rows=False,
                    )
                    moved_start_date = shifted_rows[0]["plan_date"]
                    moved_start_min = shifted_rows[0]["start_min"]
                    moved_end_rows = rows(
                        con.execute(
                            "SELECT plan_date, end_min FROM planning_row WHERE envelope_id IN (" + ",".join("?" for _ in created_envelope_ids) + ") ORDER BY plan_date DESC, end_min DESC, row_id DESC",
                            created_envelope_ids,
                        )
                    ) if created_envelope_ids else []
                    moved_end_date = moved_end_rows[0]["plan_date"] if moved_end_rows else shifted_rows[-1]["plan_date"]
                    moved_end_min = moved_end_rows[0]["end_min"] if moved_end_rows else shifted_rows[-1]["end_min"]
                    try:
                        shifted_ops = replan_downstream(
                            con,
                            primary_row["ps_id"],
                            int(primary_row["seq_no"] or 0) + 1,
                            moved_end_date,
                            moved_end_min,
                            split_piece_id=primary_row.get("split_piece_id"),
                        )
                    except ValueError as e:
                        if clash_mode == "next_free":
                            shifted_ops = replan_downstream_with_floor(
                                con,
                                primary_row["ps_id"],
                                int(primary_row["seq_no"] or 0) + 1,
                                moved_end_date,
                                moved_end_min,
                                split_piece_id=primary_row.get("split_piece_id"),
                            )
                        else:
                            raise
                    if remaining_row_ids:
                        _sync_group_locks_from_row_ids(con, remaining_row_ids)
                    con.execute("RELEASE move_envelope")
                    adjusted = moved_start_date != requested_date or moved_start_min != requested_start
                    message = None
                    if adjusted:
                        message = f"Requested {requested_date} {hhmm(requested_start)} was unavailable; moved to next valid slot at {moved_start_date} {hhmm(moved_start_min)}."
                    elif split_envelopes:
                        message = "Moved envelope across multiple groups to fit the available machine time."
                    return jsonify({
                        "ok": True,
                        "adjusted": adjusted,
                        "split_envelopes": split_envelopes,
                        "requested_date": requested_date,
                        "requested_start_min": requested_start,
                        "actual_date": moved_start_date,
                        "actual_start_min": moved_start_min,
                        "created_envelope_ids": created_envelope_ids,
                        "message": message,
                        "shifted_ops": shifted_ops,
                    })
                except ValueError as e:
                    last_error = str(e)
                    con.execute("ROLLBACK TO move_envelope")
                    con.execute("RELEASE move_envelope")
                    if clash_mode != "next_free":
                        shifted_probe = _shift_rows_to_anchor(move_rows, search_date, search_start)
                        return jsonify({
                            "ok": True,
                            "needs_clash_resolution": True,
                            "requested_date": requested_date,
                            "requested_start_min": requested_start,
                            "clash_date": search_date,
                            "clash_start_min": int(search_start),
                            "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], shifted_probe)[:8]],
                            "message": last_error,
                            "clash_options": [
                                {"label": "Cancel", "value": "cancel"},
                                {"label": "Plan but skip clash", "value": "skip_clash"},
                                {"label": "Next free slot", "value": "next_free"},
                            ],
                        })
                    next_search_date, next_search_start = first_schedulable_start_for_duration(
                        con,
                        machine,
                        search_date,
                        search_start + 1,
                        required_mins,
                        "weekdays",
                        date.fromisoformat(search_date),
                        ignore_envelope_id=source_envelope_id,
                    )
                    if next_search_date == search_date and int(next_search_start) == int(search_start):
                        search_start += 1
                    else:
                        search_date, search_start = next_search_date, int(next_search_start)
                    continue
            return jsonify({
                "ok": True,
                "needs_clash_resolution": True,
                "requested_date": requested_date,
                "requested_start_min": requested_start,
                "clash_date": search_date,
                "clash_start_min": int(search_start),
                "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], _shift_rows_to_anchor(move_rows, search_date, search_start))[:8]],
                "message": last_error or "Could not find a downstream-safe slot",
                "clash_options": [
                    {"label": "Cancel", "value": "cancel"},
                    {"label": "Plan but skip clash", "value": "skip_clash"},
                    {"label": "Next free slot", "value": "next_free"},
                ],
            })
    except Exception as exc:
        return api_error(f"Move failed: {exc}")


def move_row_qty_by_id(data=None):
    try:
        data = data or (request.get_json() or {})
        move_qty = float(data.get("move_qty") or 0)
        source_date = compact_text(data.get("plan_date"))
        target_date = compact_text(data.get("target_date"))
        target_start = int(data.get("target_start_min") or STANDARD_START)
        anchor_row_id = data.get("row_id")

        if move_qty <= 0:
            return api_error("Move quantity must be greater than 0")
        if not source_date:
            return api_error("Source planning date is required")
        if not target_date:
            return api_error("Target date is required")

        with db() as con:
            row_id = data.get("row_id")
            if row_id in (None, ""):
                return api_error("Row id is required")
            anchor_row, source_envelope_rows = _load_row_envelope(con, int(row_id))
            if not anchor_row:
                return api_error("Planning row not found", 404)
            if not source_envelope_rows:
                return api_error("Envelope has no rows", 404)
            if any(int(r.get("locked") or 0) == 1 for r in source_envelope_rows):
                return api_error("Locked envelopes cannot be edited")

            if anchor_row_id not in (None, ""):
                source_rows = [r for r in source_envelope_rows if int(r["row_id"]) == int(anchor_row_id)]
                if not source_rows:
                    source_rows = [r for r in source_envelope_rows if r["plan_date"] == source_date]
            else:
                source_rows = [r for r in source_envelope_rows if r["plan_date"] == source_date]
            source_rows = sorted(
                [dict(r) for r in source_rows],
                key=lambda r: (r["plan_date"], int(r["start_min"] or 0), int(r["end_min"] or 0), int(r["row_id"] or 0)),
            )
            if source_rows:
                selected_index = next((idx for idx, r in enumerate(source_rows) if int(r["row_id"]) == int(anchor_row_id or row_id)), 0)
                source_rows = source_rows[selected_index:]
            source_row_ids = {int(row["row_id"]) for row in source_rows}
            if not source_rows:
                return api_error("Source planning row was not found", 404)

            available_qty = sum(float(row.get("qty") or 0) for row in source_rows)
            if move_qty - available_qty > 1e-9:
                return api_error(f"Only {available_qty:g} units are available on {source_date}")

            step = one(con.execute("SELECT * FROM part_flow_steps WHERE step_id = ?", (int(anchor_row["flow_step_id"] or 0),)))
            machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (int(anchor_row["machine_id"] or 0),)))
            if not step or not machine:
                return api_error("Row route or machine could not be resolved")

            remaining_rows = [sched_row_payload(row) for row in source_envelope_rows if row["plan_date"] != source_date]
            source_work_rows = [sched_row_payload(row) for row in source_rows]
            kept_source_rows, moved_rows = _partition_source_rows_for_move(source_work_rows, move_qty, step["cycle_time"] or 1)
            remaining_rows.extend(kept_source_rows)

            requested_date = target_date
            requested_start = target_start
            setup_time = effective_setup_time_for_new_envelope(
                con,
                anchor_row["ps_id"],
                anchor_row["machine_id"],
                target_date,
                target_start,
                step["setup_time"] or 0,
                step,
                ignore_envelope_id=int(anchor_row.get("envelope_id") or 0) or None,
            )
            source_envelope_id = int(anchor_row.get("envelope_id") or 0) or None
            required_mins = duration_for_units(step["cycle_time"] or 1, move_qty, setup_time)
            clash_mode = compact_text(data.get("clash_mode") or "")
            if not clash_mode:
                clash_date, clash_start = first_unblocked_start_for_move(
                    con,
                    machine,
                    target_date,
                    target_start,
                    "weekdays",
                    date.fromisoformat(target_date),
                    ignore_envelope_id=int(anchor_row.get("envelope_id") or 0) or None,
                    ignore_row_ids={int(anchor_row["row_id"])},
                )
                if clash_date != target_date or int(clash_start) != int(target_start):
                    shifted_probe = _shift_rows_to_anchor(moved_rows, clash_date, clash_start)
                    return jsonify({
                        "ok": True,
                        "needs_clash_resolution": True,
                        "requested_date": requested_date,
                        "requested_start_min": requested_start,
                        "clash_date": clash_date,
                        "clash_start_min": int(clash_start),
                        "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], shifted_probe)[:8]],
                        "message": "Requested slot clashes with existing work. Choose how to resolve it.",
                        "clash_options": [
                            {"label": "Cancel", "value": "cancel"},
                            {"label": "Plan but skip clash", "value": "skip_clash"},
                            {"label": "Next free slot", "value": "next_free"},
                        ],
                    })
            if clash_mode == "next_free":
                target_date, target_start = first_unblocked_start_for_move(
                    con,
                    machine,
                    target_date,
                    target_start,
                    "weekdays",
                    date.fromisoformat(target_date),
                    ignore_row_ids={int(anchor_row["row_id"])},
                )
            search_date = target_date
            search_start = int(target_start)
            last_error = None
            move_split_marker = next((compact_text(r.get("split_group_id")) for r in source_rows if compact_text(r.get("split_group_id"))), "")
            move_piece_id = next((compact_text(r.get("split_piece_id")) for r in source_rows if compact_text(r.get("split_piece_id"))), "") or _split_piece_marker("piece")
            for _ in range(180):
                con.execute("SAVEPOINT move_qty_envelope")
                try:
                    shifted_rows = _planned_move_rows(
                        con,
                        machine,
                        moved_rows,
                        search_date,
                        search_start,
                        move_qty,
                        step["cycle_time"] or 1,
                        setup_time,
                        ignore_envelope_id=source_envelope_id,
                    )
                    for row in shifted_rows:
                        row["split_piece_id"] = move_piece_id
                    blockers = _shifted_rows_blockers(con, machine["machine_id"], shifted_rows) if machine.get("shift_profile") == "24HR" else []
                    if blockers:
                        raise ValueError("Requested slot clashes with existing work")

                    combined_rows = remaining_rows + shifted_rows
                    combined_rows.sort(key=lambda row: (row["plan_date"], int(row["start_min"]), int(row["end_min"])))
                    split_envelopes = any(
                        not rows_share_envelope(con, machine, combined_rows[idx], combined_rows[idx + 1])
                        for idx in range(len(combined_rows) - 1)
                    )

                    if source_row_ids:
                        q_marks = ",".join("?" for _ in source_row_ids)
                        con.execute(f"DELETE FROM planning_row WHERE row_id IN ({q_marks})", list(source_row_ids))
                    created_envelope_ids = create_blocks_from_schedule(
                        con,
                        anchor_row["ps_id"],
                        step,
                        machine,
                        combined_rows,
                        split_group_id=move_split_marker,
                        merge_adjacent_rows=False,
                    )
                    moved_start_date = shifted_rows[0]["plan_date"]
                    moved_start_min = shifted_rows[0]["start_min"]
                    final_rows = rows(
                        con.execute(
                            "SELECT plan_date, end_min FROM planning_row WHERE envelope_id IN (" + ",".join("?" for _ in created_envelope_ids) + ") ORDER BY plan_date DESC, end_min DESC, row_id DESC",
                            created_envelope_ids,
                        )
                    ) if created_envelope_ids else []
                    moved_end_date = final_rows[0]["plan_date"] if final_rows else shifted_rows[-1]["plan_date"]
                    moved_end_min = final_rows[0]["end_min"] if final_rows else shifted_rows[-1]["end_min"]
                    try:
                        shifted_ops = replan_downstream(
                            con,
                            anchor_row["ps_id"],
                            int(anchor_row["seq_no"] or 0) + 1,
                            moved_end_date,
                            moved_end_min,
                            split_piece_id=anchor_row.get("split_piece_id"),
                        )
                    except ValueError as e:
                        if clash_mode == "next_free":
                            shifted_ops = replan_downstream_with_floor(
                                con,
                                anchor_row["ps_id"],
                                int(anchor_row["seq_no"] or 0) + 1,
                                moved_end_date,
                                moved_end_min,
                                split_piece_id=anchor_row.get("split_piece_id"),
                            )
                        else:
                            raise

                    con.execute("RELEASE move_qty_envelope")
                    adjusted = moved_start_date != requested_date or moved_start_min != requested_start
                    message = None
                    if adjusted:
                        message = f"Requested {requested_date} {hhmm(requested_start)} was unavailable; moved to next valid slot at {moved_start_date} {hhmm(moved_start_min)}."
                    elif split_envelopes:
                        message = "Moved quantity across multiple envelopes to fit the available machine time."

                    return jsonify({
                        "ok": True,
                        "adjusted": adjusted,
                        "split_envelopes": split_envelopes,
                        "requested_date": requested_date,
                        "requested_start_min": requested_start,
                        "actual_date": moved_start_date,
                        "actual_start_min": moved_start_min,
                        "created_envelope_ids": created_envelope_ids,
                        "message": message,
                        "shifted_ops": shifted_ops,
                    })
                except ValueError as e:
                    last_error = str(e)
                    con.execute("ROLLBACK TO move_qty_envelope")
                    con.execute("RELEASE move_qty_envelope")
                    if clash_mode != "next_free":
                        shifted_probe = _shift_rows_to_anchor(moved_rows, search_date, search_start)
                        return jsonify({
                            "ok": True,
                            "needs_clash_resolution": True,
                            "requested_date": requested_date,
                            "requested_start_min": requested_start,
                            "clash_date": search_date,
                            "clash_start_min": int(search_start),
                            "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], shifted_probe)[:8]],
                            "message": last_error,
                            "clash_options": [
                                {"label": "Cancel", "value": "cancel"},
                                {"label": "Plan but skip clash", "value": "skip_clash"},
                                {"label": "Next free slot", "value": "next_free"},
                            ],
                        })
                    next_search_date, next_search_start = first_unblocked_start_for_move(
                        con,
                        machine,
                        search_date,
                        search_start + 1,
                        "weekdays",
                        date.fromisoformat(search_date),
                        ignore_row_ids=source_row_id_set,
                    )
                    if next_search_date == search_date and int(next_search_start) == int(search_start):
                        search_start += 1
                    else:
                        search_date, search_start = next_search_date, int(next_search_start)
                    continue
            return jsonify({
                "ok": True,
                "needs_clash_resolution": True,
                "requested_date": requested_date,
                "requested_start_min": requested_start,
                "clash_date": search_date,
                "clash_start_min": int(search_start),
                "blocking_rows": [dict(r) for r in _shifted_rows_blockers(con, machine["machine_id"], _shift_rows_to_anchor(moved_rows, search_date, search_start))[:8]],
                "message": last_error or "Could not find a downstream-safe slot",
                "clash_options": [
                    {"label": "Cancel", "value": "cancel"},
                    {"label": "Plan but skip clash", "value": "skip_clash"},
                    {"label": "Next free slot", "value": "next_free"},
                ],
            })
    except Exception as exc:
        return api_error(f"Move quantity failed: {exc}")


def split_row_by_id(data=None):
    data = data or (request.get_json() or {})
    split_qty = float(data.get("split_qty") or 0)
    target_row_id = int(data.get("row_id") or 0)
    envelope_row_ids = [
        int(v)
        for v in (data.get("envelope_row_ids") or [])
        if str(v).strip().isdigit() and int(v) > 0
    ]
    if split_qty <= 0:
        return api_error("Split quantity must be greater than 0")
    with db() as con:
        if not target_row_id:
            return api_error("Row id is required")
        anchor_row = one(
            con.execute(
                """
                SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                       setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                FROM planning_row
                WHERE row_id = ?
                """,
                (target_row_id,),
            )
        )
        if not anchor_row:
            return api_error("Planning row not found", 404)
        if envelope_row_ids:
            q_marks = ",".join("?" for _ in envelope_row_ids)
            source_rows = rows(
                con.execute(
                f"""
                SELECT row_id, envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                       setup_mins, locked, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                FROM planning_row
                WHERE row_id IN ({q_marks})
                ORDER BY plan_date, start_min, end_min, row_id
                """,
                envelope_row_ids,
                )
            )
        else:
            _, source_rows = _load_row_envelope(con, target_row_id)
        if not source_rows:
            return api_error("Envelope has no rows", 404)
        if any(int(r.get("locked") or 0) == 1 for r in source_rows):
            return api_error("Locked envelopes cannot be split")
        step = one(con.execute("SELECT * FROM part_flow_steps WHERE step_id = ?", (int(anchor_row["flow_step_id"] or 0),)))
        machine = one(con.execute("SELECT * FROM machines WHERE machine_id = ?", (int(anchor_row["machine_id"] or 0),)))
        total_qty = sum(float(row["qty"] or 0) for row in source_rows)
        if not step or not machine or not source_rows or split_qty >= total_qty:
            return api_error("Invalid split quantity")
        split_boundary_date = compact_text(anchor_row["plan_date"])
        preserved_actuals = [
            {
                "plan_date": row["plan_date"],
                "start_min": int(row["start_min"] or 0),
                "end_min": int(row["end_min"] or 0),
                "actual_out": float(row.get("actual_out") or 0),
                "actual_out_set": int(row.get("actual_out_set") or 0),
            }
            for row in source_rows
            if compact_text(row["plan_date"]) < split_boundary_date and int(row.get("actual_out_set") or 0) == 1
        ]
        target_index = next((idx for idx, row in enumerate(source_rows) if int(row["row_id"]) == target_row_id), None) if target_row_id else None
        split_mode = "head"
        remaining_to_split = float(split_qty)
        moved_rows = []
        remaining_row_ids = []
        source_piece_id = compact_text(source_rows[0].get("split_piece_id")) if source_rows else ""
        moved_piece_id = source_piece_id or _split_piece_marker("piece")
        remaining_piece_id = _split_piece_marker("piece")
        for idx, row in enumerate(source_rows):
            if remaining_to_split <= 0:
                remaining_row_ids.extend(int(r["row_id"]) for r in source_rows[idx:])
                break
            row_qty = float(row["qty"] or 0)
            if row_qty <= 0:
                remaining_row_ids.append(int(row["row_id"]))
                continue
            move_qty = min(row_qty, remaining_to_split)
            setup_here = float(row["setup_mins"] or 0) if not moved_rows else 0.0
            move_duration = duration_for_units(step["cycle_time"] or 1, move_qty, setup_here)
            original_start = int(row["start_min"])
            original_end = int(row["end_min"])
            if move_qty >= row_qty - 1e-9:
                con.execute("DELETE FROM planning_row WHERE row_id = ?", (row["row_id"],))
                move_end = (
                    original_start + move_duration
                    if machine["shift_profile"] == "24HR"
                    else standard_day_finish_min(original_start, move_duration)
                )
                moved_rows.append({
                    "plan_date": row["plan_date"],
                    "start_min": original_start,
                    "end_min": move_end,
                    "qty": row_qty,
                    "setup_mins": setup_here,
                    "split_piece_id": moved_piece_id,
                })
            else:
                new_start = min(original_end, max(original_start, original_start + move_duration))
                con.execute(
                    "UPDATE planning_row SET qty = ?, start_min = ?, setup_mins = 0 WHERE row_id = ?",
                    (row_qty - move_qty, new_start, row["row_id"]),
                )
                remaining_row_ids.append(int(row["row_id"]))
                moved_rows.append({
                    "plan_date": row["plan_date"],
                    "start_min": original_start,
                    "end_min": new_start,
                    "qty": move_qty,
                    "setup_mins": setup_here,
                    "split_piece_id": moved_piece_id,
                })
            remaining_to_split -= move_qty

        if remaining_to_split > 1e-9:
            con.rollback()
            return api_error("Could not split the requested quantity")

        split_marker = _split_group_marker("split")
        moved_rows.sort(key=lambda r: (r["plan_date"], r["start_min"]))
        new_envelope_ids = create_blocks_from_schedule(
            con,
            anchor_row["ps_id"],
            step,
            machine,
            moved_rows,
            merge_same_day_same_step=False,
            merge_adjacent_rows=False,
            split_group_id=split_marker,
        )
        if preserved_actuals and new_envelope_ids:
            preserved_actuals.sort(key=lambda r: (r["plan_date"], int(r["start_min"]), int(r["end_min"])))
            rebuilt_rows = rows(
                con.execute(
                    "SELECT row_id, plan_date, start_min, end_min FROM planning_row WHERE envelope_id IN (" + ",".join("?" for _ in new_envelope_ids) + ") ORDER BY plan_date, start_min, end_min, row_id",
                    new_envelope_ids,
                )
            )
            rebuilt_candidates = [
                row for row in rebuilt_rows
                if compact_text(row["plan_date"]) < split_boundary_date
            ]
            for actual_row, rebuilt_row in zip(preserved_actuals, rebuilt_candidates):
                con.execute(
                    """
                    UPDATE planning_row
                    SET actual_out = ?, actual_out_set = ?
                    WHERE row_id = ?
                    """,
                    (actual_row["actual_out"], actual_row["actual_out_set"], int(rebuilt_row["row_id"])),
                )
        tail_envelope_ids = []
        if remaining_row_ids:
            q_marks = ",".join("?" for _ in remaining_row_ids)
            remaining_rows = rows(
                con.execute(
                    f"""
                    SELECT row_id, plan_date, start_min, end_min, qty, setup_mins, seq_no, op_no, machine_id, machine_code, flow_step_id, split_group_id, split_piece_id
                    FROM planning_row
                    WHERE row_id IN ({q_marks})
                    ORDER BY plan_date, start_min, end_min, row_id
                    """,
                    remaining_row_ids,
                )
            )
            if remaining_rows:
                for row in remaining_rows:
                    row["split_piece_id"] = remaining_piece_id
                con.execute(f"DELETE FROM planning_row WHERE row_id IN ({q_marks})", remaining_row_ids)
                tail_envelope_ids = create_blocks_from_schedule(
                    con,
                    anchor_row["ps_id"],
                    step,
                    machine,
                    [{**sched_row_payload(r), "split_piece_id": remaining_piece_id} for r in remaining_rows],
                    merge_same_day_same_step=False,
                    merge_adjacent_rows=False,
                    split_group_id=split_marker,
                )
        return jsonify({
            "ok": True,
            "new_envelope_ids": new_envelope_ids,
            "tail_envelope_ids": tail_envelope_ids,
            "split_debug": {
                "target_row_id": target_row_id or None,
                "target_row_index": target_index,
                "split_mode": split_mode,
                "row_ids_in_envelope": [int(r["row_id"]) for r in source_rows],
                "requested_qty": split_qty,
                "total_qty_before": total_qty,
                "rows_scanned": len(source_rows),
                "moved_qty": sum(float(r["qty"] or 0) for r in moved_rows),
                "moved_rows": [
                    {
                        "plan_date": r["plan_date"],
                        "start_min": r["start_min"],
                        "end_min": r["end_min"],
                        "qty": r["qty"],
                    }
                    for r in moved_rows
                ],
            },
        })


@app.post("/api/rows/<int:source_row_id>/merge-into/<int:target_row_id>")
def merge_row_into_row(source_row_id, target_row_id):
    with db() as con:
        source_row, source_envelope = _load_row_envelope(con, source_row_id)
        target_row, target_envelope = _load_row_envelope(con, target_row_id)
        if not source_row or not target_row:
            return api_error("Row not found", 404)
        if not source_envelope or not target_envelope:
            return api_error("Envelope row not found", 404)
        if compact_text(source_row["ps_id"]) != compact_text(target_row["ps_id"]) or int(source_row["flow_step_id"]) != int(target_row["flow_step_id"]) or int(source_row["machine_id"]) != int(target_row["machine_id"]):
            return api_error("Rows must match on PS, op, and machine")
        source_envelope_id = int(source_row.get("envelope_id") or 0)
        target_envelope_id = int(target_row.get("envelope_id") or 0)
        if source_envelope_id and source_envelope_id == target_envelope_id:
            ordered_rows = sorted(
                [dict(r) for r in source_envelope],
                key=lambda r: (r["plan_date"], int(r["start_min"] or 0), int(r["end_min"] or 0), int(r["row_id"] or 0)),
            )
            pieces = []
            current_piece = []
            for row in ordered_rows:
                current_piece_piece_id = compact_text(current_piece[0].get("split_piece_id")) if current_piece else ""
                row_piece_id = compact_text(row.get("split_piece_id"))
                if current_piece and row_piece_id != current_piece_piece_id:
                    pieces.append(current_piece)
                    current_piece = []
                current_piece.append(row)
            if current_piece:
                pieces.append(current_piece)

            def abs_min(plan_date, minute):
                return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

            def piece_of_row(row_id_value):
                for piece in pieces:
                    if any(int(r["row_id"]) == int(row_id_value) for r in piece):
                        return piece
                return None

            source_piece = piece_of_row(source_row_id)
            target_piece = piece_of_row(target_row_id)
            has_split_marker = any(compact_text(r.get("split_piece_id")) or compact_text(r.get("split_group_id")) for r in ordered_rows)
            if has_split_marker and source_piece and target_piece and source_piece != target_piece:
                source_start_abs = abs_min(source_piece[0]["plan_date"], source_piece[0]["start_min"])
                source_end_abs = abs_min(source_piece[-1]["plan_date"], source_piece[-1]["end_min"])
                target_start_abs = abs_min(target_piece[0]["plan_date"], target_piece[0]["start_min"])
                target_end_abs = abs_min(target_piece[-1]["plan_date"], target_piece[-1]["end_min"])
                if source_start_abs < target_start_abs:
                    earlier_piece, later_piece = source_piece, target_piece
                    earlier_end_abs, later_start_abs = source_end_abs, target_start_abs
                else:
                    earlier_piece, later_piece = target_piece, source_piece
                    earlier_end_abs, later_start_abs = target_end_abs, source_start_abs
                shift_delta = earlier_end_abs - later_start_abs
                if shift_delta != 0:
                    later_row_ids = [int(r["row_id"]) for r in later_piece]
                    q_marks = ",".join("?" for _ in later_row_ids)
                    later_rows = rows(
                        con.execute(
                            f"""
                            SELECT row_id, plan_date, start_min, end_min
                            FROM planning_row
                            WHERE row_id IN ({q_marks})
                            ORDER BY plan_date, start_min, row_id
                            """,
                            later_row_ids,
                        )
                    )
                    for row in later_rows:
                        start_abs = abs_min(row["plan_date"], row["start_min"]) + shift_delta
                        end_abs = abs_min(row["plan_date"], row["end_min"]) + shift_delta
                        new_start_date = date.fromordinal(start_abs // 1440).isoformat()
                        new_start_min = int(start_abs % 1440)
                        new_end_date = date.fromordinal(end_abs // 1440).isoformat()
                        new_end_min = int(end_abs % 1440)
                        con.execute(
                            """
                            UPDATE planning_row
                            SET plan_date = ?, start_min = ?, end_min = ?
                            WHERE row_id = ?
                            """,
                            (new_start_date, new_start_min, new_end_date, new_end_min, int(row["row_id"])),
                        )

                combined_ids = sorted({int(r["row_id"]) for r in source_piece + target_piece})
                if combined_ids:
                    q_marks = ",".join("?" for _ in combined_ids)
                    con.execute(f"UPDATE planning_row SET split_group_id = '' WHERE row_id IN ({q_marks})", combined_ids)
                    con.execute(f"UPDATE planning_row SET split_piece_id = '' WHERE row_id IN ({q_marks})", combined_ids)
                    merge_contiguous_rows_for_rows(con, combined_ids)
                    _sync_group_locks_from_row_ids(con, combined_ids)
                return jsonify({
                    "ok": True,
                    "merged_row_ids": combined_ids,
                    "source_row_id": int(source_row_id),
                    "target_row_id": int(target_row_id),
                    "message": "Rows merged and stitched",
                })
            move_result = move_row_by_id({
                "row_id": source_row_id,
                "new_date": target_row["plan_date"],
                "new_start_min": int(target_row["end_min"] or 0),
                "machine_id": int(target_row.get("machine_id") or 0),
                "clash_mode": "next_free",
            })
            payload = move_result.get_json(silent=True) if hasattr(move_result, "get_json") else None
            if not isinstance(payload, dict):
                payload = {}
            payload.setdefault("ok", True)
            payload.setdefault("message", "Row moved to next available timing")
            payload.setdefault("source_row_id", int(source_row_id))
            payload.setdefault("target_row_id", int(target_row_id))
            payload.setdefault("same_envelope_move", True)
            return jsonify(payload)
        if not target_envelope_id:
            target_envelope_id = source_envelope_id
        if not target_envelope_id:
            return api_error("Target envelope not found", 404)
        source_row_ids = [int(r["row_id"]) for r in source_envelope]
        target_row_ids = [int(r["row_id"]) for r in target_envelope]

        def abs_min(plan_date, minute):
            return date.fromisoformat(plan_date).toordinal() * 1440 + int(minute or 0)

        merged_rows = sorted(
            [dict(r) for r in source_envelope + target_envelope],
            key=lambda r: (abs_min(r["plan_date"], r["start_min"]), int(r["start_min"] or 0), int(r["row_id"] or 0)),
        )
        merge_cutoff_date = compact_text(target_row["plan_date"])
        preserved_actuals = [
            {
                "plan_date": row["plan_date"],
                "start_min": int(row["start_min"] or 0),
                "end_min": int(row["end_min"] or 0),
                "actual_out": float(row.get("actual_out") or 0),
                "actual_out_set": int(row.get("actual_out_set") or 0),
            }
            for row in merged_rows
            if compact_text(row["plan_date"]) < merge_cutoff_date and int(row.get("actual_out_set") or 0) == 1
        ]
        for row in merged_rows:
            row["split_group_id"] = ""
            row["split_piece_id"] = ""
        earliest_row = merged_rows[0]
        step = one(
            con.execute(
                """
                SELECT step_id, seq_no, op_no, cycle_time, setup_time
                FROM part_flow_steps
                WHERE step_id = ?
                """,
                (int(earliest_row["flow_step_id"]),),
            )
        )
        machine = one(
            con.execute(
                """
                SELECT machine_id, machine_code, shift_profile
                FROM machines
                WHERE machine_id = ?
                """,
                (int(earliest_row["machine_id"]),),
            )
        )
        if not step or not machine:
            return api_error("Planning metadata not found", 404)
        combined_qty = sum(float(r.get("qty") or 0) for r in merged_rows)
        setup_here = effective_setup_time_for_new_envelope(
            con,
            compact_text(earliest_row["ps_id"]),
            int(machine["machine_id"]),
            earliest_row["plan_date"],
            int(earliest_row["start_min"] or 0),
            step["setup_time"] or 0,
            step,
        )
        duration_mins = duration_for_units(step["cycle_time"] or 1, combined_qty, setup_here)
        delete_row_ids = sorted({int(r["row_id"]) for r in merged_rows})
        if delete_row_ids:
            q_marks = ",".join("?" for _ in delete_row_ids)
            con.execute(f"DELETE FROM planning_row WHERE row_id IN ({q_marks})", delete_row_ids)
        archive_envelope_ids = sorted({source_envelope_id, target_envelope_id})
        if archive_envelope_ids:
            q_marks = ",".join("?" for _ in archive_envelope_ids)
            con.execute(
                f"""
                UPDATE planning_envelope
                SET archived = 1, split_group_id = '', updated_at = CURRENT_TIMESTAMP
                WHERE envelope_id IN ({q_marks})
                """,
                archive_envelope_ids,
            )
        scheduled_rows = schedule_rows(
            con,
            machine,
            earliest_row["plan_date"],
            int(earliest_row["start_min"] or 0),
            duration_mins,
            combined_qty,
            step["cycle_time"] or 1,
            setup_here,
            carry_mode="weekdays",
            reference_date=date.fromisoformat(earliest_row["plan_date"]),
        )
        created_envelope_ids = create_blocks_from_schedule(
            con,
            compact_text(earliest_row["ps_id"]),
            step,
            machine,
            scheduled_rows,
        )
        if preserved_actuals and created_envelope_ids:
            preserved_actuals.sort(key=lambda r: (r["plan_date"], int(r["start_min"]), int(r["end_min"])))
            rebuilt_rows = rows(
                con.execute(
                    "SELECT row_id, plan_date, start_min, end_min FROM planning_row WHERE envelope_id IN (" + ",".join("?" for _ in created_envelope_ids) + ") ORDER BY plan_date, start_min, end_min, row_id",
                    created_envelope_ids,
                )
            )
            rebuilt_candidates = [
                row for row in rebuilt_rows
                if compact_text(row["plan_date"]) < merge_cutoff_date
            ]
            for actual_row, rebuilt_row in zip(preserved_actuals, rebuilt_candidates):
                con.execute(
                    """
                    UPDATE planning_row
                    SET actual_out = ?, actual_out_set = ?
                    WHERE row_id = ?
                    """,
                    (actual_row["actual_out"], actual_row["actual_out_set"], int(rebuilt_row["row_id"])),
                )
        new_row_ids = []
        if created_envelope_ids:
            q_marks = ",".join("?" for _ in created_envelope_ids)
            new_row_ids = [
                int(r["row_id"])
                for r in rows(
                    con.execute(
                        f"""
                        SELECT row_id
                        FROM planning_row
                        WHERE envelope_id IN ({q_marks})
                        ORDER BY plan_date, start_min, row_id
                        """,
                        created_envelope_ids,
                    )
                )
            ]
        return jsonify({
            "ok": True,
            "merged_row_ids": new_row_ids,
            "merged_from_row_ids": delete_row_ids,
            "source_row_id": int(source_row_id),
            "target_row_id": int(target_row_id),
            "message": "Rows merged and replanned",
        })


def archive_row_by_id(data=None):
    data = data or {}
    with db() as con:
        row_id = data.get("row_id")
        if row_id in (None, ""):
            return api_error("Row id is required")
        row = one(con.execute(
            """
            SELECT row_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set,
                   seq_no, op_no, machine_id, machine_code, flow_step_id, envelope_id
            FROM planning_row
            WHERE row_id = ?
            """,
            (int(row_id),),
        ))
        if not row:
            return api_error("Planning row not found", 404)
        group_key = _group_key_from_row(row)
        con.execute(
            """
            INSERT INTO history_block (ps_id, op_no, seq_no, machine_code, total_qty, actual_out)
            VALUES (?, ?, ?, ?, ?, ?)
            """,
            (row["ps_id"], row["op_no"], row["seq_no"], row["machine_code"], row["qty"], row["actual_out"]),
        )
        con.execute(
            "INSERT INTO history_envelope (ps_id, seq_no, op_no, machine_id, machine_code, flow_step_id, total_qty, actual_out, locked, archived_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, CURRENT_TIMESTAMP)",
            (row["ps_id"], row["seq_no"], row["op_no"], row.get("machine_id") or 0, row["machine_code"], row.get("flow_step_id") or 0, row["qty"], row["actual_out"], row["locked"] or 0),
        )
        hist_envelope_id = con.execute("SELECT last_insert_rowid()").fetchone()[0]
        con.execute(
            "INSERT INTO history_row (hist_envelope_id, ps_id, plan_date, start_min, end_min, qty, actual_out, actual_out_set, seq_no, op_no, machine_id, machine_code, flow_step_id) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
            (hist_envelope_id, row["ps_id"], row["plan_date"], row["start_min"], row["end_min"], row["qty"], row["actual_out"], row["actual_out_set"], row["seq_no"], row["op_no"], row.get("machine_id") or 0, row["machine_code"], row.get("flow_step_id") or 0),
        )
        con.execute("DELETE FROM planning_row WHERE row_id = ?", (int(row_id),))
        remaining_row_ids = [int(r["row_id"]) for r in _rows_for_group_key(con, group_key)]
        if remaining_row_ids:
            _sync_group_locks_from_row_ids(con, remaining_row_ids)
        else:
            con.execute(
                """
                UPDATE planning_envelope
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                """,
                group_key,
            )
            con.execute(
                """
                UPDATE planning_block
                SET archived = 1, updated_at = CURRENT_TIMESTAMP
                WHERE ps_id = ? AND seq_no = ? AND op_no = ? AND machine_id = ? AND machine_code = ? AND flow_step_id = ?
                """,
                group_key,
            )
        return jsonify({"ok": True})


@app.get("/api/history")
def api_history():
    ps_id = request.args.get("ps_id")
    sql = "SELECT *, printf('%02d:%02d', start_min / 60, start_min % 60) start_hhmm, printf('%02d:%02d', end_min / 60, end_min % 60) end_hhmm FROM history"
    params = []
    if ps_id:
        sql += " WHERE ps_id LIKE ?"
        params.append(f"%{ps_id}%")
    sql += " ORDER BY archived_at DESC, hist_id DESC"
    with db() as con:
        return jsonify(rows(con.execute(sql, params)))


ensure_db()

if __name__ == "__main__":
    if os.environ.get("APP_AUTO_RELOAD", "1") == "1":
        _start_dev_reload_watcher()
    app.run(debug=True, port=5000, use_reloader=False)
